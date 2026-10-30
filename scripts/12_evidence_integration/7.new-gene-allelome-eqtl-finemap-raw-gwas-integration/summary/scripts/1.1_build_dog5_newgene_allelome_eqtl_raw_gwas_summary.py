#!/usr/bin/python3.10
# -*- coding: utf-8 -*-


from __future__ import annotations

import argparse
import csv
import fcntl
import gzip
import html
import re
import traceback
from collections import Counter, defaultdict
from contextlib import redirect_stderr, redirect_stdout
from datetime import datetime
from pathlib import Path
from typing import Dict, Iterable, List, Mapping, Sequence, Tuple
from xml.sax.saxutils import escape


from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


SCRIPT_DIR = Path(
    '/path/to/repository/scripts/12_evidence_integration/7.new-gene-allelome-eqtl-finemap-raw-gwas-integration/summary/scripts'
)
LOG_DIR = SCRIPT_DIR / "log"
MASTER_LOG = LOG_DIR / "1.1_build_dog5_newgene_allelome_eqtl_raw_gwas_summary.log"
LOCK_FILE = LOG_DIR / "1.1_build_dog5_newgene_allelome_eqtl_raw_gwas_summary.lock"
SOURCE_RUN_DIR = Path(
    '/path/to/project/analysis_20260817/results/5-dogs/7.new-gene-allelome-eqtl-finemap-raw-gwas-integration/1.1_integrate_dog5_newgene_allelome_eqtl_raw_gwas/20260917_201538'
)
DEFAULT_CANDIDATES = (
    SOURCE_RUN_DIR / "dog5_newgene_eqtl_raw_gwas_all_candidates.tsv"
)
DEFAULT_PCG_VALUES = (
    SOURCE_RUN_DIR / "dog5_newgene_to_pcGene_candidate_links.tsv.gz"
)
DEFAULT_SOURCE_SUMMARY = (
    SOURCE_RUN_DIR
    / "dog5_newgene_allelome_eqtl_raw_gwas_integration_summary.tsv"
)
DEFAULT_SOURCE_MANIFEST = (
    SOURCE_RUN_DIR
    / "dog5_newgene_allelome_eqtl_raw_gwas_integration_manifest.tsv"
)
DEFAULT_OUTPUT_ROOT = Path(
    '/path/to/project/analysis_20260817/results/5-dogs/7.new-gene-allelome-eqtl-finemap-raw-gwas-integration/summary/results'
)
CONCURRENT_TASKS = 1
THREADS_PER_TASK = 1

EXPECTED_DIRECTION = "new_lncRNA_to_pcGene"
EXPECTED_PAIR_ORIENTATION = "lncRNA_pcGene_class_orientation"
EXPECTED_CAUSAL_DIRECTION_SUPPORTED = "no"
EXPECTED_DIRECTION_BASIS = (
    "posthoc_class_orientation_lncRNA_candidate_regulator_"
    "pcGene_putative_target"
)
EXPECTED_INTEGRATION_LEVEL = "block_level_new_gene_allelome_relation"

CANDIDATE_REQUIRED = {
    "candidate_id",
    "block_id",
    "new_gene_id",
    "tissue",
    "chrom",
    "eqtl_gwas_pos",
    "eqtl_gwas_ref",
    "eqtl_gwas_alt",
    "normalized_variant_key",
    "eqtl_finemap_snp",
    "finemap_role",
    "is_credible_set_member",
    "finemap_ppa",
    "marker_rank_by_ppa",
    "block_top_ppa_tie_count",
    "block_top_marker_is_unique",
    "block_credible_set_size",
    "phase_scope",
    "gwas_trait",
    "gwas_marker_id",
    "gwas_p_lrt",
    "gwas_variant_p_tier",
    "interpretation_note",
    "gwas_p_value_status",
    "gwas_modelled_locus_count",
    "gwas_region_match_status",
    "beta_alignment_status",
    "is_p_lt_1e-5",
    "allele_match_status",
    "new_gene_start0",
    "new_gene_end0",
    "new_gene_strand",
    "new_gene_class_codes",
    "new_gene_novelty_categories",
    "new_gene_feelnc_biotypes",
    "eqtl_gwas_snp_inside_new_gene_span",
    "eqtl_gwas_snp_distance_to_new_gene_span_bp",
    "allelome_block_match",
    "pcgene_expanded_link_count",
    "unique_pcgene_count",
}

PCGENE_REQUIRED = {
    "candidate_id",
    "block_id",
    "new_gene_id",
    "tissue",
    "chrom",
    "eqtl_gwas_pos",
    "eqtl_gwas_ref",
    "eqtl_gwas_alt",
    "normalized_variant_key",
    "eqtl_finemap_snp",
    "finemap_role",
    "is_credible_set_member",
    "finemap_ppa",
    "ppa_tier",
    "marker_rank_by_ppa",
    "block_top_ppa_tie_count",
    "block_top_marker_is_unique",
    "block_credible_set_size",
    "phase_scope",
    "gwas_trait",
    "gwas_marker_id",
    "gwas_p_lrt",
    "gwas_variant_p_tier",
    "interpretation_note",
    "gwas_p_value_status",
    "gwas_modelled_locus_count",
    "gwas_region_match_status",
    "beta_alignment_status",
    "is_p_lt_1e-5",
    "gwas_af",
    "gwas_beta",
    "gwas_se",
    "gwas_allele1",
    "gwas_allele0",
    "source_overlap_type",
    "allele_match_status",
    "new_gene_chr",
    "new_gene_start0",
    "new_gene_end0",
    "new_gene_span_bp",
    "new_gene_strand",
    "new_gene_merge_status",
    "new_gene_dataset_support",
    "new_gene_class_codes",
    "new_gene_novelty_categories",
    "new_gene_feelnc_biotypes",
    "eqtl_gwas_snp_inside_new_gene_span",
    "eqtl_gwas_snp_distance_to_new_gene_span_bp",
    "animal_id",
    "sample_tissue",
    "canonical_pair_id",
    "sample_pair_id",
    "evidence_stratum",
    "overlap_bp",
    "either_endpoint_multigene",
    "shared_variant_identity_checked",
    "source_link_file",
    "source_link_line",
    "interval_gap_bp",
    "mechanism",
    "linkage_score",
    "new_gene_endpoint_role",
    "allelome_endpoint_gene_id",
    "allelome_endpoint_start0",
    "allelome_endpoint_end0",
    "allelome_endpoint_strand",
    "allelome_endpoint_chrom",
    "allelome_endpoint_pos",
    "allelome_endpoint_ref",
    "allelome_endpoint_alt",
    "allelome_endpoint_name",
    "eqtl_gwas_snp_equals_allelome_endpoint",
    "eqtl_gwas_snp_distance_to_allelome_endpoint_bp",
    "partner_gene_id",
    "partner_gene_symbol",
    "partner_gene_strand",
    "partner_chrom",
    "partner_start",
    "partner_end",
    "partner_biotype",
    "partner_endpoint_type",
    "candidate_direction",
    "pair_orientation",
    "causal_direction_supported",
    "evidence_tier",
    "direction_basis",
    "integration_level",
}

BLOCK_COLUMNS = [
    "pcgene_evidence_strata",
    "candidate_id",
    "block_id",
    "new_gene_id",
    "new_gene_type",
    "new_gene_class_codes",
    "new_gene_novelty_categories",
    "new_gene_feelnc_biotypes",
    "tissue",
    "eqtl_raw_gwas_variant",
    "finemap_role",
    "finemap_ppa",
    "marker_rank_by_ppa",
    "block_top_ppa_tie_count",
    "block_top_marker_is_unique",
    "block_credible_set_size",
    "gwas_trait",
    "gwas_marker_id",
    "gwas_p_lrt",
    "gwas_variant_p_tier",
    "interpretation_note",
    "gwas_p_value_status",
    "gwas_modelled_locus_count",
    "gwas_region_match_status",
    "beta_alignment_status",
    "is_p_lt_1e-5",
    "eqtl_gwas_snp_inside_new_gene_span",
    "eqtl_gwas_snp_distance_to_new_gene_span_bp",
    "allelome_block_match",
    "evidence_tier",
    "pcgene_expanded_row_count",
    "unique_pcgene_pair_count",
    "interpretation",
]

PAIR_COLUMNS = [
    "pcgene_evidence_strata",
    "candidate_id",
    "block_id",
    "new_gene_id",
    "new_gene_type",
    "new_gene_class_codes",
    "new_gene_novelty_categories",
    "new_gene_feelnc_biotypes",
    "new_gene_interval",
    "new_gene_strand",
    "tissue",
    "eqtl_raw_gwas_variant",
    "finemap_role",
    "finemap_ppa",
    "marker_rank_by_ppa",
    "block_top_ppa_tie_count",
    "block_top_marker_is_unique",
    "block_credible_set_size",
    "gwas_trait",
    "gwas_marker_id",
    "gwas_p_lrt",
    "gwas_variant_p_tier",
    "interpretation_note",
    "gwas_p_value_status",
    "gwas_modelled_locus_count",
    "gwas_region_match_status",
    "beta_alignment_status",
    "is_p_lt_1e-5",
    "eqtl_gwas_snp_inside_new_gene_span",
    "eqtl_gwas_snp_distance_to_new_gene_span_bp",
    "partner_gene_id",
    "partner_gene_symbol",
    "partner_biotype",
    "partner_gene_interval",
    "new_gene_pcgene_overlap",
    "new_gene_pcgene_overlap_bp",
    "overlap_percent_of_new_gene",
    "overlap_percent_of_pcgene",
    "evidence_tier",
    "exact_endpoint_row_count",
    "expanded_evidence_row_count",
    "supporting_dog_count",
    "supporting_dogs",
    "allelome_endpoint_count",
    "allelome_endpoints",
    "mechanism_status",
    "mechanisms",
    "linkage_score_min",
    "linkage_score_max",
    "interval_gap_bp_min",
    "interval_gap_bp_max",
    "pair_orientation",
    "causal_direction_supported",
    "interpretation",
]

FULL_COLUMNS = [
    "candidate_id",
    "gwas_trait",
    "gwas_p_lrt",
    "gwas_variant_p_tier",
    "interpretation_note",
    "gwas_p_value_status",
    "gwas_modelled_locus_count",
    "gwas_region_match_status",
    "beta_alignment_status",
    "is_p_lt_1e-5",
    "gwas_marker_id",
    "gwas_variant",
    "gwas_af",
    "gwas_beta",
    "gwas_se",
    "gwas_allele1",
    "gwas_allele0",
    "eqtl_finemap_variant",
    "finemap_role",
    "finemap_ppa",
    "ppa_tier",
    "marker_rank_by_ppa",
    "block_top_ppa_tie_count",
    "block_top_marker_is_unique",
    "block_credible_set_size",
    "eqtl_finemap_raw_gwas_exact_match",
    "allele_match_status",
    "source_overlap_type",
    "paired_protein_coding_gene",
    "paired_protein_coding_gene_id",
    "paired_protein_coding_gene_symbol",
    "paired_protein_coding_gene_strand",
    "paired_protein_coding_gene_interval",
    "allelome_pattern",
    "allelome_score",
    "allelome_interval_gap_bp",
    "animal_id",
    "tissue",
    "sample_tissue",
    "canonical_pair_id",
    "sample_pair_id",
    "evidence_stratum",
    "overlap_bp",
    "either_endpoint_multigene",
    "shared_variant_identity_checked",
    "source_link_file",
    "source_link_line",
    "matched_new_gene_id",
    "new_gene_interval",
    "new_gene_strand",
    "new_gene_class_code",
    "new_gene_novelty_category",
    "new_gene_feelnc_biotype",
    "new_gene_dataset_support",
    "new_gene_merge_status",
    "new_gene_pcgene_overlap",
    "new_gene_pcgene_overlap_bp",
    "overlap_percent_of_new_gene",
    "overlap_percent_of_pcgene",
    "finemap_gwas_inside_new_gene_span",
    "finemap_gwas_distance_to_new_gene_span_bp",
    "evidence_tier",
    "pair_orientation",
    "causal_direction_supported",
    "interpretation",
]

DISPLAY_COLUMNS = ["gwas_trait","gwas_variant","gwas_p_lrt","gwas_variant_p_tier","is_p_lt_1e-5","finemap_ppa","marker_rank_by_ppa","block_top_ppa_tie_count","matched_new_gene_id","new_gene_interval","paired_protein_coding_gene","allelome_pattern","allelome_score","sample_tissue","evidence_stratum","shared_variant_identity_checked","new_gene_pcgene_overlap","new_gene_pcgene_overlap_bp","finemap_gwas_distance_to_new_gene_span_bp"]

DISPLAY_LABELS = {"gwas_trait":"GWAS Trait","gwas_variant":"Expression–ASE / GWAS Variant","gwas_p_lrt":"GWAS p_lrt","gwas_variant_p_tier":"GWAS variant P tier","is_p_lt_1e-5":"GWAS p<1e-5","finemap_ppa":"Expression–ASE PPA","marker_rank_by_ppa":"PPA Rank (block)","block_top_ppa_tie_count":"Top PPA Tie Count","matched_new_gene_id":"Allelome New Gene","new_gene_interval":"New Gene Interval (1-based)","paired_protein_coding_gene":"Partner pcGene","allelome_pattern":"Allelome Pattern","allelome_score":"Linkage Score","sample_tissue":"Animal / Tissue","evidence_stratum":"Review Stratum","shared_variant_identity_checked":"Shared Variants Checked","new_gene_pcgene_overlap":"Gene Overlap","new_gene_pcgene_overlap_bp":"Overlap (bp)","finemap_gwas_distance_to_new_gene_span_bp":"GWAS–New Gene Distance (bp)"}

DISPLAY_GROUPS = [("GWAS / Expression–ASE", 8, "E8F0EA"), ("Direct gene-endpoint Allelome links", 6, "EEF2F7"), ("Review flags and spatial context", 5, "F7EFE7")]

DISPLAY_COLUMN_SPECS = [["gwas_trait",160,20,"text_left"],["gwas_variant",170,22,"text_left"],["gwas_p_lrt",110,14,"numeric"],["gwas_variant_p_tier",150,22,"text_left"],["is_p_lt_1e-5",95,13,"badge"],["finemap_ppa",100,13,"numeric"],["marker_rank_by_ppa",100,13,"numeric"],["block_top_ppa_tie_count",100,13,"numeric"],["matched_new_gene_id",180,23,"text_left"],["new_gene_interval",185,23,"text_left"],["paired_protein_coding_gene",130,18,"text_left"],["allelome_pattern",100,13,"badge"],["allelome_score",85,11,"numeric"],["sample_tissue",180,23,"text_left"],["evidence_stratum",180,23,"text_left"],["shared_variant_identity_checked",100,15,"badge"],["new_gene_pcgene_overlap",95,14,"badge"],["new_gene_pcgene_overlap_bp",95,13,"numeric"],["finemap_gwas_distance_to_new_gene_span_bp",120,16,"numeric"]]

def yes_no_label(value: str) -> str:
    normalized = str(value).strip().lower()
    if normalized == "yes":
        return "Yes"
    if normalized == "no":
        return "No"
    return str(value).strip() or "NA"


def binary_flag_label(value: str) -> str:
    normalized = str(value).strip().lower()
    if normalized in {"1", "yes", "true"}:
        return "Yes"
    if normalized in {"0", "no", "false"}:
        return "No"
    return str(value).strip() or "NA"


def mechanism_label(value: str) -> str:
    text = str(value).strip()
    return text.capitalize() if text else "NA"


def build_display_rows(
    rows: Sequence[Mapping[str, str]],
) -> List[Dict[str, str]]:
    display_rows: List[Dict[str, str]] = []
    for row in rows:
        display_row = {
            field: str(row.get(field, "")).strip() or "NA"
            for field in DISPLAY_COLUMNS
        }
        display_row["sample_tissue"] = display_row["sample_tissue"].replace(
            "_", " "
        )
        display_row["allelome_pattern"] = mechanism_label(
            display_row["allelome_pattern"]
        )
        display_row["new_gene_pcgene_overlap"] = yes_no_label(
            display_row["new_gene_pcgene_overlap"]
        )
        display_row["shared_variant_identity_checked"] = yes_no_label(
            display_row["shared_variant_identity_checked"]
        )
        display_row["is_p_lt_1e-5"] = binary_flag_label(
            display_row["is_p_lt_1e-5"]
        )
        display_row["_candidate_id"] = str(
            row.get("candidate_id", "")
        ).strip()
        display_rows.append(display_row)
    return display_rows


def wrap_text(value: str, width: int) -> List[str]:
    text = str(value).strip() or "NA"
    pieces = re.split(r"([,; /_])", text)
    lines: List[str] = []
    current = ""
    for piece in pieces:
        if len(current) + len(piece) <= width:
            current += piece
            continue
        if current:
            lines.append(current.strip())
        current = piece.strip()
    if current:
        lines.append(current.strip())
    return lines or [text]


def svg_badge_style(field: str, value: str) -> Tuple[str, str]:
    if field == "is_p_lt_1e-5":
        if value == "Yes":
            return "#e7f3ec", "#1f5a3a"
        return "#f1f3f5", "#5c6773"
    if field == "new_gene_pcgene_overlap":
        if value == "Yes":
            return "#e7f3ec", "#1f5a3a"
        return "#f1f3f5", "#5c6773"
    if field == "allelome_pattern":
        if value == "Enhancing":
            return "#e7f3ec", "#1f5a3a"
        if value == "Repressing":
            return "#fff1e6", "#9a4d00"
    if field == "evidence_tier":
        if value == "EXACT_ASE_ENDPOINT":
            return "#e7f3ec", "#1f5a3a"
        return "#eef2f7", "#34495e"
    return "#f1f3f5", "#5c6773"


def log(message: str) -> None:
    timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    print(f"[{timestamp} ] {message}", flush=True)


def require_file(path: Path, label: str) -> None:
    if not path.is_file() or path.stat().st_size == 0:
        raise FileNotFoundError(f"{label} does not exist or is empty: {path}")


def read_tsv(path: Path) -> Tuple[List[str], List[Dict[str, str]]]:
    opener = gzip.open if path.suffix == ".gz" else Path.open
    if path.suffix == ".gz":
        handle = opener(path, "rt", encoding="utf-8", newline="")
    else:
        handle = opener(path, "r", encoding="utf-8", newline="")
    with handle:
        reader = csv.DictReader(handle, delimiter="\t")
        fields = list(reader.fieldnames or [])
        return fields, list(reader)


def read_metrics(path: Path) -> Dict[str, str]:
    fields, rows = read_tsv(path)
    require_columns(path, fields, {"metric", "value"})
    metrics: Dict[str, str] = {}
    for row in rows:
        key = row["metric"].strip()
        if not key or key in metrics:
            raise ValueError(f"{path} contains empty or duplicate metric: {key!r}")
        metrics[key] = row["value"].strip()
    return metrics


def require_columns(
    path: Path, observed: Sequence[str], required: Iterable[str]
) -> None:
    missing = sorted(set(required) - set(observed))
    if missing:
        raise ValueError(f"{path} missing fields: {','.join(missing)}")


def to_int(value: str, label: str) -> int:
    try:
        return int(float(value))
    except (TypeError, ValueError) as exc:
        raise ValueError(f"{label} is not an integer: {value!r}") from exc


def to_float(value: str, label: str) -> float:
    try:
        return float(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"{label} is not numeric: {value!r}") from exc


def unique_join(values: Iterable[str]) -> str:
    result: List[str] = []
    seen = set()
    for value in values:
        text = str(value).strip()
        if text and text not in {"NA", "."} and text not in seen:
            seen.add(text)
            result.append(text)
    return "|".join(result) if result else "NA"


def variant_text(row: Mapping[str, str]) -> str:
    return (
        f"{row['chrom']}:{row['eqtl_gwas_pos']} "
        f"{row['eqtl_gwas_ref']}/{row['eqtl_gwas_alt']}"
    )


def interval_text(
    chrom: str, start: str, end: str, *, start_is_zero_based: bool
) -> str:
    start_value = to_int(start, "interval_start")
    if start_is_zero_based:
        start_value += 1
    return f"{chrom}:{start_value}-{to_int(end, 'interval_end')}"


def new_gene_pcgene_overlap_fields(
    row: Mapping[str, str],
) -> Dict[str, str]:
    new_start0 = to_int(row["new_gene_start0"], "new_gene_start0")
    new_end0 = to_int(row["new_gene_end0"], "new_gene_end0")
    new_span = to_int(row["new_gene_span_bp"], "new_gene_span_bp")
    partner_start0 = to_int(row["partner_start"], "partner_start")
    partner_end0 = to_int(row["partner_end"], "partner_end")
    if (
        new_start0 < 0
        or new_end0 <= new_start0
        or new_end0 - new_start0 != new_span
    ):
        raise ValueError("new gene coordinates or spans are inconsistent")
    if partner_start0 < 0 or partner_end0 <= partner_start0:
        raise ValueError("pcGene of 0-based half-open coordinates invalid")


    partner_span = partner_end0 - partner_start0
    overlap_bp = 0
    if row["new_gene_chr"] == row["partner_chrom"]:
        overlap_bp = max(
            0,
            min(new_end0, partner_end0)
            - max(new_start0, partner_start0),
        )
    return {
        "new_gene_pcgene_overlap": "yes" if overlap_bp > 0 else "no",
        "new_gene_pcgene_overlap_bp": str(overlap_bp),
        "overlap_percent_of_new_gene": (
            f"{100.0 * overlap_bp / new_span:.4f}"
        ),
        "overlap_percent_of_pcgene": (
            f"{100.0 * overlap_bp / partner_span:.4f}"
        ),
    }


def new_gene_type(row: Mapping[str, str]) -> str:
    values = {
        item.strip()
        for item in row["new_gene_feelnc_biotypes"].replace(";", "|").split("|")
        if item.strip() and item.strip() not in {"NA", "."}
    }
    return "lncRNA" if "lncRNA" in values else unique_join(sorted(values))


def evidence_tier(rows: Sequence[Mapping[str, str]]) -> str:
    return "DIRECT_NEW_GENE_ENDPOINT_BLOCK_SUPPORT"


def mechanism_status(rows: Sequence[Mapping[str, str]]) -> str:
    mechanisms = {
        row["mechanism"].strip()
        for row in rows
        if row["mechanism"].strip() not in {"", "NA", "."}
    }
    if mechanisms == {"enhancing"}:
        return "consistent_enhancing"
    if mechanisms == {"repressing"}:
        return "consistent_repressing"
    return "mixed_or_unknown"


def validate_inputs(
    candidate_path: Path,
    candidate_fields: Sequence[str],
    candidates: Sequence[Mapping[str, str]],
    pcgene_path: Path,
    pcgene_fields: Sequence[str],
    pcgene_rows: Sequence[Mapping[str, str]],
    source_metrics: Mapping[str, str],
) -> None:
    require_columns(candidate_path, candidate_fields, CANDIDATE_REQUIRED)
    require_columns(pcgene_path, pcgene_fields, PCGENE_REQUIRED)
    if source_metrics.get("overall_status") != "PASS":
        raise ValueError("upstream Step 7 summary is not PASS")

    expected_candidates = to_int(
        source_metrics.get("new_gene_candidate_rows", ""),
        "new_gene_candidate_rows",
    )
    expected_unique_candidate_genes = to_int(
        source_metrics.get("new_gene_candidate_unique_genes", ""),
        "new_gene_candidate_unique_genes",
    )
    expected_matched_candidates = to_int(
        source_metrics.get("matched_candidate_rows", ""),
        "matched_candidate_rows",
    )
    expected_unmatched_candidates = to_int(
        source_metrics.get("unmatched_candidate_rows", ""),
        "unmatched_candidate_rows",
    )
    expected_pcgene = to_int(
        source_metrics.get("lncRNA_pcGene_class_oriented_rows", ""),
        "lncRNA_pcGene_class_oriented_rows",
    )
    if len(candidates) != expected_candidates:
        raise ValueError(
            f"candidate block row counts differ: {len(candidates)} != {expected_candidates}"
        )
    if (
        len({row["new_gene_id"] for row in candidates})
        != expected_unique_candidate_genes
    ):
        raise ValueError("candidate table unique novel gene count and upstream summary mismatch")
    matched_count = sum(
        row["allelome_block_match"] == "yes" for row in candidates
    )
    unmatched_count = sum(
        row["allelome_block_match"] == "no" for row in candidates
    )
    if matched_count != expected_matched_candidates:
        raise ValueError("Allelome matching candidate count with upstream summary mismatch")
    if unmatched_count != expected_unmatched_candidates:
        raise ValueError("Allelome unmatched candidate count and upstream summary mismatch")
    if matched_count + unmatched_count != len(candidates):
        raise ValueError("Candidate table contains an allelome_block_match value other than yes or no")
    if len(pcgene_rows) != expected_pcgene:
        raise ValueError(
            f"pcGene expanded rows count mismatch: {len(pcgene_rows)} != {expected_pcgene}"
        )

    candidate_ids = [row["candidate_id"] for row in candidates]
    if len(candidate_ids) != len(set(candidate_ids)):
        raise ValueError("candidate table candidate_id is not unique")
    known_candidates = set(candidate_ids)
    for line_number, row in enumerate(candidates, start=2):
        if row["is_credible_set_member"] != "1":
            raise ValueError(
                f"{candidate_path}:{line_number} does not belong to eQTL credible set"
            )
        if to_float(row["finemap_ppa"], "finemap_ppa") < 0.1:
            raise ValueError(f"{candidate_path}:{line_number} PPA below 0.1")
        if row["phase_scope"] != row["chrom"]:
            raise ValueError(
                f"{candidate_path}:{line_number} is not all chromosome phase_scope"
            )
        if row["allele_match_status"] != "allele_set_match":
            raise ValueError(
                f"{candidate_path}:{line_number} is not unordered allele set matching"
            )
        p_lrt = to_float(row["gwas_p_lrt"], "gwas_p_lrt")
        expected_p_flag = "1" if p_lrt < 1e-5 else "0"
        if row["is_p_lt_1e-5"] != expected_p_flag:
            raise ValueError(
                f"{candidate_path}:{line_number} of p_lrt with P<1e-5 marker mismatch"
            )

    expanded_keys = set()
    for line_number, row in enumerate(pcgene_rows, start=2):
        if row["candidate_id"] not in known_candidates:
            raise ValueError(
                f"{pcgene_path}:{line_number} candidate_id not in candidate table"
            )
        if row["candidate_direction"] != EXPECTED_DIRECTION:
            raise ValueError(
                f"{pcgene_path}:{line_number} candidate_direction does not satisfy the contract"
            )
        if row["pair_orientation"] != EXPECTED_PAIR_ORIENTATION:
            raise ValueError(
                f"{pcgene_path}:{line_number} pair_orientation does not satisfy the contract"
            )
        if (
            row["causal_direction_supported"]
            != EXPECTED_CAUSAL_DIRECTION_SUPPORTED
        ):
            raise ValueError(
                f"{pcgene_path}:{line_number} unexpected claim support causal orientation"
            )
        expected_tier = evidence_tier([row])
        if row["evidence_tier"] != expected_tier:
            raise ValueError(
                f"{pcgene_path}:{line_number} evidence_tier with endpoint evidence mismatch"
            )
        if row["direction_basis"] != EXPECTED_DIRECTION_BASIS:
            raise ValueError(
                f"{pcgene_path}:{line_number} direction_basis does not satisfy the contract"
            )
        if row["integration_level"] != EXPECTED_INTEGRATION_LEVEL:
            raise ValueError(
                f"{pcgene_path}:{line_number} integration_level does not satisfy the contract"
            )
        if row["partner_endpoint_type"] != "known_protein_coding":
            raise ValueError(
                f"{pcgene_path}:{line_number} partner is not pcGene"
            )
        if row["partner_biotype"] != "protein_coding":
            raise ValueError(
                f"{pcgene_path}:{line_number} partner is not protein_coding"
            )
        if not row["source_overlap_type"].startswith("exact_variant_"):
            raise ValueError(
                f"{pcgene_path}:{line_number} is not eQTL fine-map with raw GWAS exact variant overlap"
            )
        p_lrt = to_float(row["gwas_p_lrt"], "gwas_p_lrt")
        expected_p_flag = "1" if p_lrt < 1e-5 else "0"
        if row["is_p_lt_1e-5"] != expected_p_flag:
            raise ValueError(
                f"{pcgene_path}:{line_number} of p_lrt with P<1e-5 marker mismatch"
            )
        key = (
            row["candidate_id"],
            row["sample_pair_id"],
            row["new_gene_endpoint_role"],
        )
        if key in expanded_keys:
            raise ValueError(f"{pcgene_path}:{line_number} duplicate expanded evidence key")
        expanded_keys.add(key)


def validate_source_manifest(
    path: Path,
    *,
    candidates: Path,
    pcgene_links: Path,
    source_summary: Path,
) -> None:
    fields, rows = read_tsv(path)
    require_columns(path, fields, {"kind", "label", "path", "purpose"})
    observed = {
        row["label"]: Path(row["path"])
        for row in rows
        if row["kind"] == "output"
    }
    expected = {
        "all_new_gene_candidates": candidates,
        "new_lncRNA_to_pcGene_links": pcgene_links,
        "integration_summary": source_summary,
    }
    for label, expected_path in expected.items():
        if observed.get(label) != expected_path:
            raise ValueError(
                f"{path} in {label} path with current input mismatch: "
                f"{observed.get(label)} != {expected_path}"
            )


def build_expanded_rows(
    rows: Sequence[Dict[str, str]],
) -> Tuple[List[str], List[Dict[str, str]]]:
    result: List[Dict[str, str]] = []
    for row in rows:
        overlap_fields = new_gene_pcgene_overlap_fields(row)
        interpretation = (
            "Exact variant matches between Expression–ASE and raw GWAS; Allelome directly connects the new-gene interval to a pcGene. "
            "Gene endpoints do not specify a single SNP; this does not establish colocalization or causal direction and must be reviewed with evidence_stratum."
        )
        partner_label = (
            row["partner_gene_symbol"]
            if row["partner_gene_symbol"] not in {"", "NA", "."}
            else row["partner_gene_id"]
        )
        result.append(
            {
                "gwas_trait": row["gwas_trait"],
                "gwas_p_lrt": row["gwas_p_lrt"],
                "gwas_variant_p_tier": row["gwas_variant_p_tier"],
                "interpretation_note": row["interpretation_note"],
                "gwas_p_value_status": row["gwas_p_value_status"],
                "gwas_modelled_locus_count": row["gwas_modelled_locus_count"],
                "gwas_region_match_status": row["gwas_region_match_status"],
                "beta_alignment_status": row["beta_alignment_status"],
                "is_p_lt_1e-5": row[
                    "is_p_lt_1e-5"
                ],
                "gwas_marker_id": row["gwas_marker_id"],
                "gwas_variant": variant_text(row),
                "gwas_af": row["gwas_af"],
                "gwas_beta": row["gwas_beta"],
                "gwas_se": row["gwas_se"],
                "gwas_allele1": row["gwas_allele1"],
                "gwas_allele0": row["gwas_allele0"],
                "eqtl_finemap_variant": row["eqtl_finemap_snp"],
                "finemap_role": row["finemap_role"],
                "finemap_ppa": row["finemap_ppa"],
                "ppa_tier": row["ppa_tier"],
                "marker_rank_by_ppa": row["marker_rank_by_ppa"],
                "block_top_ppa_tie_count": row["block_top_ppa_tie_count"],
                "block_top_marker_is_unique": row["block_top_marker_is_unique"],
                "block_credible_set_size": row["block_credible_set_size"],
                "eqtl_finemap_raw_gwas_exact_match": "yes",
                "allele_match_status": row["allele_match_status"],
                "source_overlap_type": row["source_overlap_type"],
                "paired_protein_coding_gene": partner_label,
                "paired_protein_coding_gene_id": row["partner_gene_id"],
                "paired_protein_coding_gene_symbol": (
                    row["partner_gene_symbol"]
                ),
                "paired_protein_coding_gene_strand": (
                    row["partner_gene_strand"]
                ),
                "paired_protein_coding_gene_interval": interval_text(
                    row["partner_chrom"],
                    row["partner_start"],
                    row["partner_end"],
                    start_is_zero_based=True,
                ),
                "allelome_pattern": row["mechanism"],
                "allelome_score": row["linkage_score"],
                "allelome_interval_gap_bp": row["interval_gap_bp"],
                "animal_id": row["animal_id"],
                "tissue": row["tissue"],
                "sample_tissue": row["sample_tissue"],
                "canonical_pair_id": row["canonical_pair_id"],
                "sample_pair_id": row["sample_pair_id"],
                "evidence_stratum": row["evidence_stratum"],
                "overlap_bp": row["overlap_bp"],
                "either_endpoint_multigene": row["either_endpoint_multigene"],
                "shared_variant_identity_checked": row["shared_variant_identity_checked"],
                "source_link_file": row["source_link_file"],
                "source_link_line": row["source_link_line"],
                "candidate_id": row["candidate_id"],
                "matched_new_gene_id": row["new_gene_id"],
                "new_gene_interval": interval_text(
                    row["new_gene_chr"],
                    row["new_gene_start0"],
                    row["new_gene_end0"],
                    start_is_zero_based=True,
                ),
                "new_gene_strand": row["new_gene_strand"],
                "new_gene_class_code": row["new_gene_class_codes"],
                "new_gene_novelty_category": (
                    row["new_gene_novelty_categories"]
                ),
                "new_gene_feelnc_biotype": (
                    row["new_gene_feelnc_biotypes"]
                ),
                "new_gene_dataset_support": row["new_gene_dataset_support"],
                "new_gene_merge_status": row["new_gene_merge_status"],
                **overlap_fields,
                "finemap_gwas_inside_new_gene_span": (
                    row["eqtl_gwas_snp_inside_new_gene_span"]
                ),
                "finemap_gwas_distance_to_new_gene_span_bp": (
                    row["eqtl_gwas_snp_distance_to_new_gene_span_bp"]
                ),
                "evidence_tier": row["evidence_tier"],
                "pair_orientation": row["pair_orientation"],
                "causal_direction_supported": (
                    row["causal_direction_supported"]
                ),
                "interpretation": interpretation,
            }
        )
    return FULL_COLUMNS, result


def build_pair_rows(
    rows: Sequence[Mapping[str, str]],
) -> List[Dict[str, str]]:
    grouped: Dict[Tuple[str, str], List[Mapping[str, str]]] = defaultdict(list)
    for row in rows:
        grouped[(row["candidate_id"], row["partner_gene_id"])].append(row)

    result: List[Dict[str, str]] = []
    for (candidate_id, partner_gene_id), group in grouped.items():
        first = group[0]
        for row in group[1:]:
            fixed_fields = [
                "block_id",
                "new_gene_id",
                "tissue",
                "normalized_variant_key",
                "finemap_role",
                "finemap_ppa",
                "marker_rank_by_ppa",
                "block_top_ppa_tie_count",
                "block_top_marker_is_unique",
                "block_credible_set_size",
                "gwas_trait",
                "gwas_marker_id",
                "gwas_p_lrt",
                "gwas_variant_p_tier",
                "interpretation_note",
                "gwas_p_value_status",
                "gwas_modelled_locus_count",
                "gwas_region_match_status",
                "beta_alignment_status",
                "is_p_lt_1e-5",
                "partner_gene_symbol",
                "new_gene_chr",
                "new_gene_start0",
                "new_gene_end0",
                "new_gene_span_bp",
                "partner_chrom",
                "partner_start",
                "partner_end",
            ]
            for field in fixed_fields:
                if row[field] != first[field]:
                    raise ValueError(
                        f"{candidate_id}/{partner_gene_id} of {field} mismatch"
                    )

        exact_count = "NA_gene_endpoint"
        scores = [
            to_float(row["linkage_score"], "linkage_score") for row in group
        ]
        distances = [
            to_int(row["interval_gap_bp"], "interval_gap_bp")
            for row in group
        ]
        overlap_fields = new_gene_pcgene_overlap_fields(first)
        result.append(
            {
                "candidate_id": candidate_id,
                "pcgene_evidence_strata": unique_join(sorted({r["evidence_stratum"] for r in group})),
                "block_id": first["block_id"],
                "new_gene_id": first["new_gene_id"],
                "new_gene_type": new_gene_type(first),
                "new_gene_class_codes": first["new_gene_class_codes"],
                "new_gene_novelty_categories": (
                    first["new_gene_novelty_categories"]
                ),
                "new_gene_feelnc_biotypes": (
                    first["new_gene_feelnc_biotypes"]
                ),
                "new_gene_interval": (
                    f"{first['new_gene_chr']}:"
                    f"{to_int(first['new_gene_start0'], 'new_gene_start0') + 1}"
                    f"-{first['new_gene_end0']}"
                ),
                "new_gene_strand": first["new_gene_strand"],
                "tissue": first["tissue"],
                "eqtl_raw_gwas_variant": variant_text(first),
                "finemap_role": first["finemap_role"],
                "finemap_ppa": first["finemap_ppa"],
                "marker_rank_by_ppa": first["marker_rank_by_ppa"],
                "block_top_ppa_tie_count": first["block_top_ppa_tie_count"],
                "block_top_marker_is_unique": first["block_top_marker_is_unique"],
                "block_credible_set_size": first["block_credible_set_size"],
                "gwas_trait": first["gwas_trait"],
                "gwas_marker_id": first["gwas_marker_id"],
                "gwas_p_lrt": first["gwas_p_lrt"],
                "gwas_variant_p_tier": first["gwas_variant_p_tier"],
                "interpretation_note": first["interpretation_note"],
                "gwas_p_value_status": first["gwas_p_value_status"],
                "gwas_modelled_locus_count": first["gwas_modelled_locus_count"],
                "gwas_region_match_status": first["gwas_region_match_status"],
                "beta_alignment_status": first["beta_alignment_status"],
                "is_p_lt_1e-5": first[
                    "is_p_lt_1e-5"
                ],
                "eqtl_gwas_snp_inside_new_gene_span": (
                    first["eqtl_gwas_snp_inside_new_gene_span"]
                ),
                "eqtl_gwas_snp_distance_to_new_gene_span_bp": (
                    first["eqtl_gwas_snp_distance_to_new_gene_span_bp"]
                ),
                "partner_gene_id": partner_gene_id,
                "partner_gene_symbol": first["partner_gene_symbol"],
                "partner_biotype": first["partner_biotype"],
                "partner_gene_interval": interval_text(
                    first["partner_chrom"],
                    first["partner_start"],
                    first["partner_end"],
                    start_is_zero_based=True,
                ),
                **overlap_fields,
                "evidence_tier": evidence_tier(group),
                "exact_endpoint_row_count": str(exact_count),
                "expanded_evidence_row_count": str(len(group)),
                "supporting_dog_count": str(
                    len({row["animal_id"] for row in group})
                ),
                "supporting_dogs": unique_join(
                    sorted({row["animal_id"] for row in group})
                ),
                "allelome_endpoint_count": str(
                    len({row["allelome_endpoint_gene_id"] for row in group})
                ),
                "allelome_endpoints": unique_join(
                    sorted(
                        {row["allelome_endpoint_gene_id"] for row in group}
                    )
                ),
                "mechanism_status": mechanism_status(group),
                "mechanisms": unique_join(
                    sorted({row["mechanism"] for row in group})
                ),
                "linkage_score_min": f"{min(scores):.6g}",
                "linkage_score_max": f"{max(scores):.6g}",
                "interval_gap_bp_min": str(min(distances)),
                "interval_gap_bp_max": str(max(distances)),
                "pair_orientation": first["pair_orientation"],
                "causal_direction_supported": first[
                    "causal_direction_supported"
                ],
                "interpretation": (
                    "Convergence of multiple layers of candidate evidence; not proof of colocalization, causality, or regulatory direction."
                ),
            }
        )

    return sorted(
        result,
        key=lambda row: (
            row["new_gene_id"],
            row["tissue"],
            row["partner_gene_symbol"],
            row["partner_gene_id"],
        ),
    )


def build_block_rows(
    candidates: Sequence[Mapping[str, str]],
    pcgene_rows: Sequence[Mapping[str, str]],
) -> List[Dict[str, str]]:
    by_candidate: Dict[str, List[Mapping[str, str]]] = defaultdict(list)
    for row in pcgene_rows:
        by_candidate[row["candidate_id"]].append(row)

    result: List[Dict[str, str]] = []
    for row in candidates:
        evidence = by_candidate.get(row["candidate_id"], [])
        source_expanded_count = to_int(
            row["pcgene_expanded_link_count"],
            "pcgene_expanded_link_count",
        )
        source_pair_count = to_int(
            row["unique_pcgene_count"], "unique_pcgene_count"
        )
        if evidence:
            tier = evidence_tier(evidence)
            pair_count = len(
                {item["partner_gene_id"] for item in evidence}
            )
            interpretation = "This block has a new-gene-interval–pcGene link; it is not a three-layer exact match at one SNP, and risk flags require review."
        else:
            tier = "NO_PCGENE_ALLELOME_MATCH"
            pair_count = 0
            interpretation = "This candidate has no relationship oriented by lncRNA–pcGene endpoint categories."
        if source_expanded_count != len(evidence):
            raise ValueError(
                f"{row['candidate_id']} of pcGene expanded rows count and candidate table mismatch"
            )
        if source_pair_count != pair_count:
            raise ValueError(
                f"{row['candidate_id']} unique pcGene count and candidate table mismatch"
            )
        result.append(
            {
                "candidate_id": row["candidate_id"],
                "block_id": row["block_id"],
                "pcgene_evidence_strata": unique_join(sorted({r["evidence_stratum"] for r in evidence})),
                "new_gene_id": row["new_gene_id"],
                "new_gene_type": new_gene_type(row),
                "new_gene_class_codes": row["new_gene_class_codes"],
                "new_gene_novelty_categories": (
                    row["new_gene_novelty_categories"]
                ),
                "new_gene_feelnc_biotypes": (
                    row["new_gene_feelnc_biotypes"]
                ),
                "tissue": row["tissue"],
                "eqtl_raw_gwas_variant": variant_text(row),
                "finemap_role": row["finemap_role"],
                "finemap_ppa": row["finemap_ppa"],
                "marker_rank_by_ppa": row["marker_rank_by_ppa"],
                "block_top_ppa_tie_count": row["block_top_ppa_tie_count"],
                "block_top_marker_is_unique": row["block_top_marker_is_unique"],
                "block_credible_set_size": row["block_credible_set_size"],
                "gwas_trait": row["gwas_trait"],
                "gwas_marker_id": row["gwas_marker_id"],
                "gwas_p_lrt": row["gwas_p_lrt"],
                "gwas_variant_p_tier": row["gwas_variant_p_tier"],
                "interpretation_note": row["interpretation_note"],
                "gwas_p_value_status": row["gwas_p_value_status"],
                "gwas_modelled_locus_count": row["gwas_modelled_locus_count"],
                "gwas_region_match_status": row["gwas_region_match_status"],
                "beta_alignment_status": row["beta_alignment_status"],
                "is_p_lt_1e-5": row[
                    "is_p_lt_1e-5"
                ],
                "eqtl_gwas_snp_inside_new_gene_span": (
                    row["eqtl_gwas_snp_inside_new_gene_span"]
                ),
                "eqtl_gwas_snp_distance_to_new_gene_span_bp": (
                    row["eqtl_gwas_snp_distance_to_new_gene_span_bp"]
                ),
                "allelome_block_match": row["allelome_block_match"],
                "evidence_tier": tier,
                "pcgene_expanded_row_count": str(source_expanded_count),
                "unique_pcgene_pair_count": str(source_pair_count),
                "interpretation": interpretation,
            }
        )
    return result


def write_tsv(
    path: Path, fields: Sequence[str], rows: Sequence[Mapping[str, str]]
) -> None:
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(
            handle, fieldnames=list(fields), delimiter="\t"
        )
        writer.writeheader()
        writer.writerows(
            {field: row.get(field, "") for field in fields} for row in rows
        )


def style_worksheet(worksheet) -> None:
    header_fill = PatternFill("solid", fgColor="D9EAD3")
    for cell in worksheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    for column_cells in worksheet.columns:
        values = [str(cell.value or "") for cell in column_cells[:200]]
        width = min(max(max(map(len, values), default=0) + 2, 10), 45)
        worksheet.column_dimensions[
            get_column_letter(column_cells[0].column)
        ].width = width
    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def append_sheet(
    workbook: Workbook,
    title: str,
    fields: Sequence[str],
    rows: Sequence[Mapping[str, str]],
) -> None:
    worksheet = workbook.create_sheet(title)
    worksheet.append(list(fields))
    for row in rows:
        worksheet.append([row.get(field, "") for field in fields])
    style_worksheet(worksheet)


def append_focus_sheet(
    workbook: Workbook,
    rows: Sequence[Mapping[str, str]],
) -> None:
    display_rows = build_display_rows(rows)
    worksheet = workbook.create_sheet("display_all_pc_links")
    start_column = 1
    for label, width, color in DISPLAY_GROUPS:
        end_column = start_column + width - 1
        worksheet.merge_cells(
            start_row=1,
            start_column=start_column,
            end_row=1,
            end_column=end_column,
        )
        cell = worksheet.cell(row=1, column=start_column, value=label)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=color)
        cell.font = Font(bold=True, color="222222")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        for column in range(start_column + 1, end_column + 1):
            worksheet.cell(row=1, column=column).fill = PatternFill(
                "solid", fgColor=color
            )
        start_column = end_column + 1

    worksheet.append([DISPLAY_LABELS[field] for field in DISPLAY_COLUMNS])
    for row in display_rows:
        worksheet.append([row.get(field, "") for field in DISPLAY_COLUMNS])

    header_fill = PatternFill("solid", fgColor="D9EAD3")
    for cell in worksheet[2]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
    worksheet.freeze_panes = "A3"
    worksheet.auto_filter.ref = (
        f"A2:{get_column_letter(len(DISPLAY_COLUMNS))}"
        f"{worksheet.max_row}"
    )
    width_by_field = {
        field: max(12, min(35, round(pixel_width / 7)))
        for field, pixel_width, _wrap_width, _align in DISPLAY_COLUMN_SPECS
    }
    for column_index, field in enumerate(DISPLAY_COLUMNS, start=1):
        worksheet.column_dimensions[
            get_column_letter(column_index)
        ].width = width_by_field[field]
    for row in worksheet.iter_rows(min_row=3):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def write_workbook(
    path: Path,
    summary_rows: Sequence[Mapping[str, str]],
    block_rows: Sequence[Mapping[str, str]],
    pair_rows: Sequence[Mapping[str, str]],
    expanded_fields: Sequence[str],
    expanded_rows: Sequence[Mapping[str, str]],
) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    append_sheet(workbook, "overview", ["metric", "value"], summary_rows)
    append_sheet(workbook, "candidate_blocks", BLOCK_COLUMNS, block_rows)
    append_sheet(workbook, "candidate_pc_pairs", PAIR_COLUMNS, pair_rows)
    append_focus_sheet(workbook, expanded_rows)
    append_sheet(workbook, "variant_p_lt_1e-5", expanded_fields,
                 [r for r in expanded_rows if r["is_p_lt_1e-5"] == "1"])
    append_sheet(
        workbook,
        "full_results",
        expanded_fields,
        expanded_rows,
    )
    append_sheet(
        workbook,
        "field_notes",
        ["field_group", "meaning"],
        [
            {
                "field_group": "GWAS/eQTL fine-map",
                "meaning": (
                    "Exact variant matching requires the same coordinate and unordered REF/ALT set; "
                    "PPA, credible sets, and marker_rank_by_ppa are Expression–ASE molecular metrics, not GWAS PIP/credible sets; "
                    "is_p_lt_1e-5 marks only single-variant p_lrt<1×10^-5, not regional primary membership."
                    "gwas_variant_p_tier follows upstream categories: genomewide P<2.97e-6; suggestive 2.97e-6≤P<1e-5; "
                    "exploratory_point 1e-5≤P<5e-5；extended_exploratory_point 5e-5≤P≤1e-4。"
                    "interpretation_note retains the molecular analysis limitations of no association/FDR screening or null-causal model; "
                    "gwas_region_match_status annotates only modeled membership, and beta_alignment_status records effect-direction alignment status."
                ),
            },
            {
                "field_group": "Counting and review",
                "meaning": (
                    "candidate_pc_pairs aggregates by candidate_id+partner_gene_id; "
                    "Unique gene pairs across candidates are separately deduplicated by new_gene_id+partner_gene_id."
                    "display_all_pc_links retains the full pcGene subset, whereas variant_p_lt_1e-5 applies only a strict p-value filter; "
                    "Records requiring review are not automatically excluded; evidence_stratum and shared_variant_identity_checked must be read separately."
                    "The number of top-PPA ties is not the number of ties involving the current SNP; this step does not read GWAS fine-map credible sets."
                ),
            },
            {
                "field_group": "Original Allelome",
                "meaning": (
                    "Retain new-gene interval endpoints, paired protein-coding genes, review flags, "
                    "enhancing/repressing, linkage scores, and sample tissues."
                ),
            },
            {
                "field_group": "Manual new gene",
                "meaning": (
                    "Five-dog candidate new genes used directly by Allelome, with their class code, "
                    "novelty category, and FEELnc biotype."
                ),
            },
            {
                "field_group": "New gene / pcGene spatial relationship",
                "meaning": (
                    "Calculate overlap in bp using harmonized 0-based half-open coordinates; "
                    "Report overlap as a percentage of the new-gene and pcGene lengths separately."
                ),
            },
        ],
    )
    workbook.save(path)


def write_html(
    path: Path,
    expanded_rows: Sequence[Mapping[str, str]],
    summary: Mapping[str, str],
) -> None:
    display_rows = build_display_rows(expanded_rows)
    display_alignments = {
        field: align
        for field, _width, _wrap_width, align in DISPLAY_COLUMN_SPECS
    }
    header = "".join(
        f"<th>{html.escape(DISPLAY_LABELS[field])}</th>"
        for field in DISPLAY_COLUMNS
    )
    body = []
    previous_candidate = None
    visual_group_index = -1
    for row in display_rows:
        candidate_id = row["_candidate_id"]
        block_start = candidate_id != previous_candidate
        if block_start:
            visual_group_index += 1
            previous_candidate = candidate_id
        classes = []
        if block_start:
            classes.append("block-start")
        if visual_group_index % 2:
            classes.append("alternate-block")
        cells = []
        for field in DISPLAY_COLUMNS:
            value = html.escape(str(row.get(field, "")))
            if display_alignments[field] == "badge":
                badge_class = re.sub(
                    r"[^A-Za-z0-9_-]+",
                    "-",
                    str(row.get(field, "")),
                )
                cells.append(
                    '<td class="badge-cell">'
                    f'<span class="badge {badge_class}">{value}</span></td>'
                )
            else:
                cells.append(f"<td>{value}</td>")
        class_text = f' class="{" ".join(classes)}"' if classes else ""
        body.append(f"<tr{class_text}>{''.join(cells)}</tr>")
    group_header = "".join(
        (
            f'<th colspan="{width}" style="background:#{color}">'
            f"{html.escape(label)}</th>"
        )
        for label, width, color in DISPLAY_GROUPS
    )
    document = f"""<!doctype html>
<html lang="en">
<head>
<meta charset="utf-8">
<title>Five-dog step 7 multilayer candidate evidence summary</title>
<style>
body {{ font-family: Arial, sans-serif; margin: 24px; background: #f7f7f4; color: #222; }}
h1 {{ font-size: 22px; margin-bottom: 8px; }}
p.meta {{ max-width: 1200px; margin: 0 0 16px; color: #555; font-size: 13px; line-height: 1.55; }}
.table-wrap {{ max-width: 100%; overflow-x: auto; border: 1px solid #cfcfcf; background: #fff; }}
table {{ border-collapse: collapse; min-width: 2900px; width: 100%; font-size: 12px; background: #fff; }}
th, td {{ border: 1px solid #cfcfcf; padding: 6px 8px; text-align: left; vertical-align: top; }}
th {{ background: #e6efe8; font-weight: 700; position: sticky; top: 29px; z-index: 1; }}
.group th {{ color: #222; text-align: center; top: 0; z-index: 2; }}
tr.alternate-block {{ background: #fafafa; }}
tr.block-start td {{ border-top: 2px solid #777; }}
.badge-cell {{ text-align: center; }}
.badge {{ display: inline-block; border-radius: 10px; padding: 2px 8px; background: #f1f3f5; color: #5c6773; font-weight: 700; white-space: nowrap; }}
.badge.Yes, .badge.Enhancing, .badge.EXACT_ASE_ENDPOINT {{ background: #e7f3ec; color: #1f5a3a; }}
.badge.Repressing {{ background: #fff1e6; color: #9a4d00; }}
.badge.BLOCK_LEVEL_ONLY {{ background: #eef2f7; color: #34495e; }}
</style>
</head>
<body>
<h1>Five-dog step 7: new gene × Allelome × eQTL fine-map × raw GWAS</h1>
<p class="meta">All candidate blocks: {html.escape(summary["candidate_block_rows"])};
Unique candidate new genes: {html.escape(summary["candidate_unique_new_gene_count"])};
Candidate blocks with p_lrt&lt;1×10^-5: {html.escape(summary["p_lt_1e_minus5_candidate_block_rows"])};
Candidate blocks with Allelome block support: {html.escape(summary["matched_candidate_block_rows"])};
Within-candidate new-gene–pcGene combinations: {html.escape(summary["unique_pair_rows"])}; unique gene pairs across candidates: {html.escape(summary["unique_gene_pairs_across_candidates"])};
Expanded evidence rows: {html.escape(summary["expanded_pcgene_rows"])};
Unique combinations with interval overlap: {html.escape(summary["overlapping_unique_pair_rows"])};
Single-SNP endpoint comparison: not determinable (Allelome endpoints are gene intervals).
Expression–ASE and raw GWAS are matched by exact variant identity; Allelome uses new-gene interval endpoints, so three-layer exact matching at a single SNP cannot be established. The focus table reports
{html.escape(summary["expanded_pcgene_rows"])} expanded new-gene–pcGene relationships oriented by endpoint categories,
rather than all upstream Allelome links ({html.escape(summary["all_allelome_expanded_rows"])} in total).
A single-variant P-value subset does not imply regional primary status or a passed review. PPA and credible sets belong to the Expression–ASE molecular analysis, not GWAS PIP/credible sets.
These molecular candidates were not screened by association/FDR or evaluated against a null-causal model; effect-allele directions have not been checked, so expression and behavioral effect directions cannot be linked.
Single-variant categories: genomewide P&lt;2.97×10^-6; suggestive 2.97×10^-6≤P&lt;10^-5;
exploratory_point 10^-5≤P&lt;5×10^-5；extended_exploratory_point 5×10^-5≤P≤10^-4。
Neither gene-interval overlap nor an Allelome
pattern establishes causal direction.</p>
<div class="table-wrap"><table><thead>
<tr class="group">{group_header}</tr>
<tr>{header}</tr></thead><tbody>{''.join(body)}</tbody></table></div>
</body>
</html>
"""
    path.write_text(document, encoding="utf-8")


def write_svg(
    path: Path,
    expanded_rows: Sequence[Mapping[str, str]],
    summary: Mapping[str, str],
) -> None:
    display_rows = build_display_rows(expanded_rows)
    left_margin = 44
    top_margin = 36
    title_height = 70
    group_header_height = 24
    line_height = 11

    header_lines = [
        wrap_text(DISPLAY_LABELS[field], wrap_width)
        for field, _width, wrap_width, _align in DISPLAY_COLUMN_SPECS
    ]
    header_height = max(38, 14 + line_height * max(map(len, header_lines), default=1))
    row_padding = 10
    width = left_margin * 2 + sum(
        column_width
        for _field, column_width, _wrap_width, _align
        in DISPLAY_COLUMN_SPECS
    )

    row_line_map: List[List[List[str]]] = []
    row_heights: List[int] = []
    for row in display_rows:
        wrapped_cells: List[List[str]] = []
        max_lines = 1
        for field, _column_width, wrap_width, align in DISPLAY_COLUMN_SPECS:
            lines = (
                [row.get(field, "NA")]
                if align == "badge"
                else wrap_text(row.get(field, "NA"), wrap_width)
            )
            wrapped_cells.append(lines)
            max_lines = max(max_lines, len(lines))
        row_line_map.append(wrapped_cells)
        row_heights.append(
            max(header_height, row_padding + line_height * max_lines)
        )

    height = (
        top_margin
        + title_height
        + group_header_height
        + header_height
        + sum(row_heights)
        + 22
    )
    x_positions: List[int] = []
    current_x = left_margin
    for _field, column_width, _wrap_width, _align in DISPLAY_COLUMN_SPECS:
        x_positions.append(current_x)
        current_x += column_width

    svg = [
        f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}">',
        '<rect width="100%" height="100%" fill="#ffffff"/>',
        (
            f'<text x="{width / 2:.1f}" y="{top_margin}" font-size="16" '
            'font-family="Helvetica, Arial, sans-serif" font-weight="400" '
            'fill="#111111" text-anchor="middle" dominant-baseline="middle">'
            "Five-dog new-gene multilayer candidate summary</text>"
        ),
        (
            f'<text x="{width / 2:.1f}" y="{top_margin + 18}" font-size="11.5" '
            'font-family="Helvetica, Arial, sans-serif" fill="#111111" '
            'text-anchor="middle" dominant-baseline="middle">'
            f"candidate blocks={escape(summary['candidate_block_rows'])}; "
            f"unique genes={escape(summary['candidate_unique_new_gene_count'])}; "
            f"p_lrt&lt;1e-5 blocks="
            f"{escape(summary['p_lt_1e_minus5_candidate_block_rows'])}; "
            f"Allelome-supported blocks="
            f"{escape(summary['matched_candidate_block_rows'])}; "
            f"candidate pairs={escape(summary['unique_pair_rows'])}; "
            f"unique gene pairs={escape(summary['unique_gene_pairs_across_candidates'])}; "
            f"expanded rows={escape(summary['expanded_pcgene_rows'])}; "
            "SNP endpoint comparison=not evaluable (gene endpoints)"
            "</text>"
        ),
        (
            f'<text x="{width / 2:.1f}" y="{top_margin + 36}" font-size="11.5" '
            'font-family="Helvetica, Arial, sans-serif" fill="#111111" '
            'text-anchor="middle" dominant-baseline="middle">'
            "PPA/CS refer to Expression–ASE; variant P tiers are not region scopes. No colocalization or causal direction is established."
            "</text>"
        ),
    ]

    header_y = top_margin + title_height
    table_width = width - left_margin * 2
    table_bottom = height - 22
    svg.append(
        f'<line x1="{left_margin}" y1="{header_y}" '
        f'x2="{left_margin + table_width}" y2="{header_y}" '
        'stroke="#111111" stroke-width="1.1"/>'
    )

    group_start = 0
    group_boundaries: List[int] = []
    for label, column_count, color in DISPLAY_GROUPS:
        group_end = group_start + column_count
        group_x = x_positions[group_start]
        group_width = sum(
            DISPLAY_COLUMN_SPECS[index][1]
            for index in range(group_start, group_end)
        )
        svg.append(
            f'<rect x="{group_x}" y="{header_y}" width="{group_width}" '
            f'height="{group_header_height}" fill="#{color}"/>'
        )
        svg.append(
            f'<text x="{group_x + group_width / 2:.1f}" '
            f'y="{header_y + 16}" font-size="10.5" '
            'font-family="Helvetica, Arial, sans-serif" font-weight="700" '
            f'fill="#111111" text-anchor="middle">{escape(label)}</text>'
        )
        if group_end < len(DISPLAY_COLUMN_SPECS):
            group_boundaries.append(x_positions[group_end])
        group_start = group_end

    column_header_y = header_y + group_header_height
    svg.append(
        f'<line x1="{left_margin}" y1="{column_header_y}" '
        f'x2="{left_margin + table_width}" y2="{column_header_y}" '
        'stroke="#aaaaaa" stroke-width="0.7"/>'
    )
    svg.append(
        f'<line x1="{left_margin}" '
        f'y1="{column_header_y + header_height}" '
        f'x2="{left_margin + table_width}" '
        f'y2="{column_header_y + header_height}" '
        'stroke="#111111" stroke-width="0.8"/>'
    )
    for index, (field, column_width, wrap_width, _align) in enumerate(
        DISPLAY_COLUMN_SPECS
    ):
        x = x_positions[index]
        lines = header_lines[index]
        for line_index, line in enumerate(lines):
            svg.append(
                f'<text x="{x + column_width / 2:.1f}" '
                f'y="{column_header_y + 14 + line_index * line_height}" '
                'font-size="10" font-family="Helvetica, Arial, sans-serif" '
                'font-weight="700" fill="#111111" text-anchor="middle">'
                f"{escape(line)}</text>"
            )

    current_y = column_header_y + header_height
    previous_candidate = None
    visual_group_index = -1
    for row_index, (row, wrapped_cells, row_height) in enumerate(
        zip(display_rows, row_line_map, row_heights),
        start=1,
    ):
        y = current_y
        candidate_id = row["_candidate_id"]
        if candidate_id != previous_candidate:
            visual_group_index += 1
            previous_candidate = candidate_id
            if row_index > 1:
                svg.append(
                    f'<line x1="{left_margin}" y1="{y}" '
                    f'x2="{left_margin + table_width}" y2="{y}" '
                    'stroke="#777777" stroke-width="1.1"/>'
                )
        if visual_group_index % 2:
            svg.append(
                f'<rect x="{left_margin}" y="{y}" width="{table_width}" '
                f'height="{row_height}" fill="#fafafa"/>'
            )
        for column_index, (
            field,
            column_width,
            _wrap_width,
            align,
        ) in enumerate(DISPLAY_COLUMN_SPECS):
            x = x_positions[column_index]
            lines = wrapped_cells[column_index]
            if align == "badge":
                value = row.get(field, "NA")
                badge_fill, badge_text = svg_badge_style(field, value)
                badge_width = min(
                    column_width - 16,
                    max(74, len(value) * 5.8 + 18),
                )
                badge_x = x + (column_width - badge_width) / 2
                badge_y = y + 8
                svg.append(
                    f'<rect x="{badge_x:.1f}" y="{badge_y}" '
                    f'width="{badge_width:.1f}" height="18" rx="9" '
                    f'fill="{badge_fill}"/>'
                )
                svg.append(
                    f'<text x="{badge_x + badge_width / 2:.1f}" '
                    f'y="{badge_y + 12}" font-size="9" font-family="Arial" '
                    f'font-weight="700" fill="{badge_text}" '
                    f'text-anchor="middle">{escape(value)}</text>'
                )
                continue
            text_x = x + (8 if align == "text_left" else column_width / 2)
            text_start_y = (
                y
                + (row_height - line_height * len(lines)) / 2
                + line_height * 0.75
            )
            for line_index, line in enumerate(lines):
                text_y = text_start_y + line_index * line_height
                font = (
                    "Menlo,Consolas,monospace"
                    if align == "numeric"
                    else "Helvetica,Arial,sans-serif"
                )
                anchor = "start" if align == "text_left" else "middle"
                svg.append(
                    f'<text x="{text_x:.1f}" y="{text_y:.1f}" '
                    f'font-size="9.5" font-family="{font}" fill="#111111" '
                    f'text-anchor="{anchor}" dominant-baseline="middle">'
                    f"{escape(line)}</text>"
                )
        current_y += row_height
        svg.append(
            f'<line x1="{left_margin}" y1="{current_y}" '
            f'x2="{left_margin + table_width}" y2="{current_y}" '
            'stroke="#dddddd" stroke-width="0.7"/>'
        )

    for boundary_x in group_boundaries:
        svg.append(
            f'<line x1="{boundary_x}" y1="{header_y}" '
            f'x2="{boundary_x}" y2="{table_bottom}" '
            'stroke="#888888" stroke-width="0.8"/>'
        )
    svg.append(
        f'<line x1="{left_margin}" y1="{table_bottom}" '
        f'x2="{left_margin + table_width}" y2="{table_bottom}" '
        'stroke="#111111" stroke-width="1.0"/>'
    )
    svg.append("</svg>")
    path.write_text("\n".join(svg), encoding="utf-8")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Summarize step 7 five-dog new-gene-Allelome-eQTL-fine-mapping-raw-GWAS results."
        )
    )
    parser.add_argument(
        "--candidates", type=Path, default=DEFAULT_CANDIDATES
    )
    parser.add_argument(
        "--pcgene-links", type=Path, default=DEFAULT_PCG_VALUES
    )
    parser.add_argument(
        "--source-summary", type=Path, default=DEFAULT_SOURCE_SUMMARY
    )
    parser.add_argument(
        "--source-manifest", type=Path, default=DEFAULT_SOURCE_MANIFEST
    )
    parser.add_argument(
        "--output-root", type=Path, default=DEFAULT_OUTPUT_ROOT
    )
    parser.add_argument(
        "--run-id", default=datetime.now().strftime("%Y%m%d_%H%M%S")
    )
    return parser.parse_args()


def run_workflow(args: argparse.Namespace) -> int:
    log("[stage 1/5] check Step 7 main input")
    if not re.fullmatch(r"[0-9]{8}_[0-9]{6}", args.run_id):
        raise ValueError("run_id must use YYYYMMDD_HHMMSS format")
    for path in [
        args.candidates,
        args.pcgene_links,
        args.source_summary,
        args.source_manifest,
        args.output_root,
    ]:
        if not path.is_absolute():
            raise ValueError(f"input output path must be an absolute path: {path}")
    log(f"source_run_dir={args.candidates.parent}")
    log(f"output_root={args.output_root}")
    log(
        f"concurrent_tasks={CONCURRENT_TASKS}; "
        f"threads_per_task={THREADS_PER_TASK}; "
        "resource_baseline=preserved_serial_script"
    )
    for path, label in [
        (args.candidates, "Candidate block table"),
        (args.pcgene_links, "Expanded lncRNA–pcGene table"),
        (args.source_summary, "Upstream summary"),
        (args.source_manifest, "Upstream manifest"),
    ]:
        require_file(path, label)
        log(f"{label}: {path}")

    candidate_fields, candidates = read_tsv(args.candidates)
    pcgene_fields, pcgene_rows = read_tsv(args.pcgene_links)
    source_metrics = read_metrics(args.source_summary)
    validate_source_manifest(
        args.source_manifest,
        candidates=args.candidates,
        pcgene_links=args.pcgene_links,
        source_summary=args.source_summary,
    )
    validate_inputs(
        args.candidates,
        candidate_fields,
        candidates,
        args.pcgene_links,
        pcgene_fields,
        pcgene_rows,
        source_metrics,
    )
    log(
        f"candidate block={len(candidates)}; "
        f"pcGene expand evidence={len(pcgene_rows)}"
    )

    log("[stage 2/5] build candidate block, unique pair and expand evidence")
    expanded_fields, expanded_rows = build_expanded_rows(pcgene_rows)
    pair_rows = build_pair_rows(pcgene_rows)
    block_rows = build_block_rows(candidates, pcgene_rows)
    exact_rows = "NA_gene_endpoints_not_single_variants"

    unique_gene_pairs = {(r["new_gene_id"], r["partner_gene_id"]) for r in pcgene_rows}
    log(
        f"within candidate new gene–pcGene combination={len(pair_rows)}; "
        f"across candidate unique gene for={len(unique_gene_pairs)}; "
        f"single SNP endpoint compare={exact_rows}"
    )

    candidate_gene_ids = {row["new_gene_id"] for row in block_rows}
    matched_blocks = [
        row for row in block_rows if row["allelome_block_match"] == "yes"
    ]
    unmatched_blocks = [
        row for row in block_rows if row["allelome_block_match"] == "no"
    ]
    p_lt_1e_minus5_blocks = [
        row
        for row in block_rows
        if row["is_p_lt_1e-5"] == "1"
    ]
    summary_map = {
        "source_run_id": source_metrics.get("run_id", "NA"),
        "overall_status": "PASS",
        "candidate_finemap_ppa_threshold": "0.1",
        "candidate_gwas_p_lrt_threshold": "1e-4",
        "p_lt_1e_minus5_p_lrt_threshold": "<1e-5",
        "candidate_block_rows": str(len(block_rows)),
        "candidate_variant_p_tier_counts": str(dict(Counter(r["gwas_variant_p_tier"] for r in candidates))),
        "gwas_filter_scope": source_metrics["gwas_filter_scope"],
        "variant_interpretation": source_metrics["variant_interpretation"],
        "molecular_interpretation_notes": unique_join(r["interpretation_note"] for r in candidates),
        "beta_alignment_statuses": unique_join(r["beta_alignment_status"] for r in candidates),
        "candidate_unique_new_gene_count": str(len(candidate_gene_ids)),
        "matched_candidate_block_rows": str(len(matched_blocks)),
        "matched_unique_new_gene_count": str(
            len({row["new_gene_id"] for row in matched_blocks})
        ),
        "unmatched_candidate_block_rows": str(len(unmatched_blocks)),
        "p_lt_1e_minus5_candidate_block_rows": str(len(p_lt_1e_minus5_blocks)),
        "p_lt_1e_minus5_unique_new_gene_count": str(
            len({row["new_gene_id"] for row in p_lt_1e_minus5_blocks})
        ),
        "p_lt_1e_minus5_allelome_matched_block_rows": str(
            sum(row["allelome_block_match"] == "yes" for row in p_lt_1e_minus5_blocks)
        ),
        "expanded_pcgene_rows": str(len(expanded_rows)),
        "unique_pair_rows": str(len(pair_rows)),
        "unique_gene_pairs_across_candidates": str(len(unique_gene_pairs)),
        "pair_row_unit": "candidate_id+partner_gene_id",
        "all_allelome_expanded_rows": source_metrics["expanded_allelome_link_rows"],
        "pcgene_review_stratum_counts": str(dict(Counter(r["evidence_stratum"] for r in pcgene_rows))),
        "pcgene_supported_unique_new_gene_count": str(
            len({row["new_gene_id"] for row in pair_rows})
        ),
        "unique_partner_pcgene_count": str(
            len({row["partner_gene_id"] for row in pair_rows})
        ),
        "overlapping_expanded_rows": str(
            sum(
                row["new_gene_pcgene_overlap"] == "yes"
                for row in expanded_rows
            )
        ),
        "overlapping_unique_pair_rows": str(
            sum(
                row["new_gene_pcgene_overlap"] == "yes"
                for row in pair_rows
            )
        ),
        "exact_endpoint_rows": str(exact_rows),
        "eqtl_finemap_raw_gwas_exact_candidate_block_rows": str(
            len(block_rows)
        ),
        "eqtl_finemap_raw_gwas_exact_unique_variant_count": str(
            len({row["normalized_variant_key"] for row in candidates})
        ),
        "finemap_gwas_inside_new_gene_span_candidate_rows": str(
            sum(
                row["eqtl_gwas_snp_inside_new_gene_span"] == "yes"
                for row in block_rows
            )
        ),
        "direct_gene_endpoint_candidate_pair_rows": str(
            sum(row["evidence_tier"] == "DIRECT_NEW_GENE_ENDPOINT_BLOCK_SUPPORT" for row in pair_rows)
        ),
        "orientation_rule": EXPECTED_PAIR_ORIENTATION,
        "causal_direction_supported": (
            EXPECTED_CAUSAL_DIRECTION_SUPPORTED
        ),
        "interpretation_boundary": (
            "multilayer_candidate_convergence_not_colocalization_or_causal_proof"
        ),
    }
    summary_rows = [
        {"metric": key, "value": value} for key, value in summary_map.items()
    ]

    log("[stage 3/5] create timestamp result directory")
    run_dir = args.output_root / args.run_id
    if run_dir.exists():
        raise FileExistsError(f"refuse to overwrite existing result directory: {run_dir}")
    run_dir.mkdir(parents=True, exist_ok=False)

    block_path = run_dir / "dog5_candidate_blocks.tsv"
    pair_path = run_dir / "dog5_unique_newgene_pcGene_pairs.tsv"
    expanded_path = run_dir / "dog5_expanded_newgene_pcGene_evidence.tsv"
    summary_path = run_dir / "dog5_summary_metrics.tsv"
    manifest_path = run_dir / "dog5_summary_manifest.tsv"
    workbook_path = run_dir / "dog5_newgene_multilayer_summary.xlsx"
    html_path = run_dir / "dog5_newgene_multilayer_full_results.html"
    svg_path = run_dir / "dog5_newgene_multilayer_full_results.svg"
    marker_path = run_dir / "execution_completion.marker"

    log("[stage 4/5] write TSV, Excel, HTML and SVG")
    write_tsv(block_path, BLOCK_COLUMNS, block_rows)
    write_tsv(pair_path, PAIR_COLUMNS, pair_rows)
    write_tsv(expanded_path, expanded_fields, expanded_rows)
    write_tsv(summary_path, ["metric", "value"], summary_rows)
    write_workbook(
        workbook_path,
        summary_rows,
        block_rows,
        pair_rows,
        expanded_fields,
        expanded_rows,
    )
    write_html(html_path, expanded_rows, summary_map)
    write_svg(svg_path, expanded_rows, summary_map)

    manifest_rows = [
        {
            "kind": "input",
            "label": "candidate_blocks",
            "path": str(args.candidates),
            "purpose": "All new-gene candidate blocks from step 7",
        },
        {
            "kind": "input",
            "label": "pcgene_links",
            "path": str(args.pcgene_links),
            "purpose": "Step 7 expanded lncRNA–pcGene evidence oriented by post hoc endpoint categories",
        },
        {
            "kind": "input",
            "label": "source_summary",
            "path": str(args.source_summary),
            "purpose": "Step 7 PASS status and input/output count requirements",
        },
        {
            "kind": "input",
            "label": "source_manifest",
            "path": str(args.source_manifest),
            "purpose": "Step 7 input provenance",
        },
    ]
    for label, path, purpose in [
        (
            "candidate_blocks",
            block_path,
            f"{len(block_rows)} candidate blocks and their evidence tiers",
        ),
        ("candidate_pc_pairs", pair_path, "candidate_id+partner_gene_id combinations; see summary for the number of unique gene pairs across candidates."),
        (
            "expanded_evidence",
            expanded_path,
            "Complete results organized by GWAS, Allelome, new gene, and evidence limits",
        ),
        ("summary_metrics", summary_path, "Deduplicated counts and interpretation limits"),
        (
            "workbook",
            workbook_path,
            "Multi-sheet Excel summary including new-gene–pcGene interval overlap proportions",
        ),
        (
            "html",
            html_path,
            f"Interactive table of {len(expanded_rows)} evidence records and gene-interval overlaps",
        ),
        (
            "svg",
            svg_path,
            f"Static overview of {len(expanded_rows)} evidence records and gene-interval overlaps",
        ),
        (
            "execution_completion_marker",
            marker_path,
            "Indicates only normal completion of the summary program, not a pass for result quality.",
        ),
    ]:
        manifest_rows.append(
            {
                "kind": "output",
                "label": label,
                "path": str(path),
                "purpose": purpose,
            }
        )
    write_tsv(
        manifest_path,
        ["kind", "label", "path", "purpose"],
        manifest_rows,
    )

    log("[stage 5/5] record write count and write completion marker")
    log(
        f"candidate blocks={len(block_rows)}; "
        f"unique pairs={len(pair_rows)}; "
        f"expanded evidence={len(expanded_rows)}; "
        "overlapping pairs="
        f"{summary_map['overlapping_unique_pair_rows']}"
    )
    marker_path.write_text(
        "execution_status=success\n"
        f"run_id={args.run_id}\n"
        f"completed_at={datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n",
        encoding="utf-8",
    )
    log(f"result directory: {run_dir}")
    log("task completed / done")
    log("execution_status=success")
    log(
        "completion only indicates that the main program exited normally; output completeness, format, counts and"
        "biological quality still requires separately check"
    )
    return 0


def main() -> int:
    args = parse_args()
    LOG_DIR.mkdir(parents=True, exist_ok=True)
    lock_handle = LOCK_FILE.open("a+", encoding="utf-8")
    lock_acquired = False
    try:
        try:
            fcntl.flock(
                lock_handle.fileno(), fcntl.LOCK_EX | fcntl.LOCK_NB
            )
            lock_acquired = True
        except BlockingIOError:
            print(
                f"execution_status=failed; existing the same step task holds run lock: "
                f"{LOCK_FILE}",
            )
            return 75

        with MASTER_LOG.open("w", encoding="utf-8") as log_handle:
            with redirect_stdout(log_handle), redirect_stderr(log_handle):
                try:
                    log("start five dogs Step 7 result summary")
                    log(f"canonical_master_log={MASTER_LOG}")
                    return run_workflow(args)
                except Exception:
                    log("execution_status=failed")
                    traceback.print_exc()
                    return 1
    finally:
        if lock_acquired:
            fcntl.flock(lock_handle.fileno(), fcntl.LOCK_UN)
        lock_handle.close()


if __name__ == "__main__":
    raise SystemExit(main())
