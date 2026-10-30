#!/usr/bin/python3

from __future__ import annotations
import csv
import gzip
import json
import os
import signal
import traceback
from collections import defaultdict
from contextlib import ExitStack, redirect_stdout, redirect_stderr
from datetime import datetime
from html import escape
from pathlib import Path
from typing import Mapping, Sequence

ROOT = Path('/path/to/project/analysis_20260817')
STEP = '1.1_export_dog5_allelome_results'
SCRIPT_DIR = ROOT / 'results/5-dogs/1.allelome/scripts' / STEP
OUTPUT_ROOT = ROOT / 'results/5-dogs/1.allelome' / STEP
INPUT_RUN = ROOT / '5-dogs/7.Allelome.LINK/1.4_annotate_and_merge_dog5_allelome_links/20260907_213429'
INPUT_LINKS = INPUT_RUN / 'dog5_allelome_links.annotated.tsv'
INPUT_PAIRS = INPUT_RUN / 'dog5_allelome_pairs.summary.tsv'
INPUT_SUMMARY = INPUT_RUN / 'dog5_allelome_annotation.summary.tsv'
INPUT_CONTRACT = ROOT / '5-dogs/7.Allelome.LINK/1.1_validate_dog5_allelome_inputs/20260907_195800/dog5_allelome_parameter_contract.tsv'
INPUT_TASKS = INPUT_RUN / 'dog5_allelome_annotation.task_manifest.tsv'
INPUT_MARKER = INPUT_RUN / 'execution_completion.marker'

def log(message):
    print(f'[{datetime.now():%Y-%m-%d %H:%M:%S} ] {message}', flush=True)

def read_rows(path):
    with path.open(encoding='utf-8-sig', newline='') as handle:
        yield from csv.DictReader(handle, delimiter='\t')

def write_csv(path, fields, records):
    with path.open('x', encoding='utf-8-sig', newline='') as handle:
        out = csv.DictWriter(handle, fieldnames=fields)
        out.writeheader()
        out.writerows(records)

def new_stats():
    return dict(total_links=0, enhancing_links=0, repressing_links=0,
                pairs=set(), animals=set(), tissues=set(), samples=set())

def add(stats, row, key):
    stats['total_links'] += 1
    stats[row['mechanism'] + '_links'] += 1
    stats['pairs'].add(key)
    stats['animals'].add(row['animal_id'])
    stats['tissues'].add(row['tissue'])
    stats['samples'].add((row['animal_id'], row['tissue']))

def summarize(stats):
    return dict(total_links=stats['total_links'], enhancing_links=stats['enhancing_links'],
                repressing_links=stats['repressing_links'], unique_endpoint_pairs=len(stats['pairs']),
                animal_count=len(stats['animals']), tissue_count=len(stats['tissues']),
                sample_tissue_count=len(stats['samples']))

def write_svg(path: Path, pair_rows: Sequence[Mapping[str, object]]) -> None:
    width = 1450
    left = 550
    right = 80
    top = 84
    row_height = 34
    bar_width = width - left - right
    height = top + row_height * len(pair_rows) + 80
    max_total = max((int(row["total_links"]) for row in pair_rows), default=1)
    svg = [
        '<?xml version="1.0" encoding="UTF-8"?>',
        (
            f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" '
            f'height="{height}" viewBox="0 0 {width} {height}">'
        ),
        '<rect width="100%" height="100%" fill="#ffffff"/>',
        (
            '<text x="40" y="38" font-family="Arial" font-size="22" '
            'font-weight="bold" fill="#111111">'
            'BarkBase Allelome.LINK: expanded sample-level links</text>'
        ),
        (
            '<text x="40" y="64" font-family="Arial" font-size="13" '
            'fill="#444444">Green: same-haplotype bias (enhancing); '
            'red: opposite-haplotype bias (repressing)</text>'
        ),
    ]
    for index, row in enumerate(pair_rows):
        y = top + index * row_height
        total = int(row["total_links"])
        enhancing = int(row["enhancing_links"])
        repressing = int(row["repressing_links"])
        total_width = bar_width * total / max_total
        enhancing_width = total_width * enhancing / total
        repressing_width = total_width - enhancing_width
        label = escape(str(row["pair_type"]))
        svg.extend(
            [
                (
                    f'<text x="{left - 12}" y="{y + 18}" font-family="Arial" '
                    f'font-size="12" text-anchor="end" fill="#222222">{label}</text>'
                ),
                (
                    f'<rect x="{left}" y="{y + 5}" width="{enhancing_width:.2f}" '
                    'height="18" fill="#478577"/>'
                ),
                (
                    f'<rect x="{left + enhancing_width:.2f}" y="{y + 5}" '
                    f'width="{repressing_width:.2f}" height="18" fill="#BC220D"/>'
                ),
                (
                    f'<text x="{left + total_width + 8:.2f}" y="{y + 19}" '
                    f'font-family="Arial" font-size="11" fill="#222222">{total:,}</text>'
                ),
            ]
        )
    svg.append("</svg>")
    temp_path = path.with_suffix(".svg.tmp")
    temp_path.write_text("\n".join(svg), encoding="utf-8")
    os.replace(temp_path, path)


def export(run):

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    log('[1/4] read required version input; serial=1 threads=1(retain the original serial settings)')
    for path in (INPUT_LINKS, INPUT_PAIRS, INPUT_SUMMARY, INPUT_CONTRACT, INPUT_TASKS, INPUT_MARKER):
        if not path.is_file():
            raise FileNotFoundError(path)
        log(f'input={path}')
    marker = dict(line.split('=', 1) for line in INPUT_MARKER.read_text().splitlines() if '=' in line)
    if marker.get('execution_status') != 'success':
        raise ValueError('the input step-four run did not finish normally')
    categories, samples, risks = (defaultdict(new_stats) for _ in range(3))
    total = new_stats()
    log('[2/4] export all original field, add undirected endpoint-pair identifiers; no filtering of links')
    with INPUT_LINKS.open(encoding='utf-8-sig', newline='') as source, gzip.open(
            run / 'five_dogs_allelome_links_detailed.csv.gz', 'xt', encoding='utf-8', newline='') as dest:
        reader = csv.DictReader(source, delimiter='\t')
        extra = ['canonical_pair_id', 'sample_pair_id', 'sample_tissue', 'mechanism_interpretation']
        out = csv.DictWriter(dest, fieldnames=list(reader.fieldnames) + extra)
        out.writeheader()
        for i, row in enumerate(reader, 1):

            key = (row['chr_base'], row['endpoint1_gene_id'], row['endpoint2_gene_id'])
            row['canonical_pair_id'] = json.dumps(key, ensure_ascii=False, separators=(',', ':'))
            row['sample_pair_id'] = json.dumps((row['animal_id'], row['tissue'], *key), ensure_ascii=False, separators=(',', ':'))
            row['sample_tissue'] = row['animal_id'] + '|' + row['tissue']
            row['mechanism_interpretation'] = {'enhancing': 'same_haplotype_bias', 'repressing': 'opposite_haplotype_bias'}[row['mechanism']]
            out.writerow(row)
            add(total, row, key)
            add(categories[row['pair_type']], row, key)
            add(samples[(row['animal_id'], row['tissue'])], row, key)
            add(risks[row['evidence_stratum']], row, key)
            if i % 50000 == 0:
                log(f'exported {i} records expand link')
    category_rows = [dict(pair_type=k, **summarize(v)) for k, v in categories.items()]
    category_rows.sort(key=lambda r: (-r['total_links'], r['pair_type']))
    sample_rows = [dict(animal_id=k[0], tissue=k[1], **summarize(v)) for k, v in sorted(samples.items())]
    risk_rows = [dict(evidence_stratum=k, **summarize(v)) for k, v in sorted(risks.items())]
    metrics = summarize(total)
    metrics.update(detail_input=str(INPUT_LINKS), pair_input=str(INPUT_PAIRS),
                   parameter_input=str(INPUT_CONTRACT), output_directory=str(run),
                   link_directionality='unordered', filtering='none; all input links retained',
                   coordinate_convention='0-based half-open',
                   pair_definition='chromosome + unordered GeneID pair',
                   risk_pair_note='risk groups are disjoint for rows; a pair may occur in multiple risk groups',
                   score_note='LINK score is not a p-value or posterior probability',
                   count_note='site allele support sum; not deduplicated molecule counts')
    overall_rows = [dict(metric=k, value=v) for k, v in metrics.items()]
    log('[3/4] write classification, sample, risk and unique endpoint for; Excel summarize only, do not load all detailed records')
    common = list(summarize(total))
    tables = [
        ('pair_class_summary', ['pair_type'] + common, category_rows),
        ('sample_tissue_summary', ['animal_id', 'tissue'] + common, sample_rows),
        ('risk_summary', ['evidence_stratum'] + common, risk_rows),
        ('overall', ['metric', 'value'], overall_rows),
        ('upstream_summary', ['metric', 'value'], list(read_rows(INPUT_SUMMARY))),
        ('parameter_contract', ['parameter', 'value'], list(read_rows(INPUT_CONTRACT))),
    ]
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, fields, records in tables:
        write_csv(run / f'five_dogs_allelome_{name}.csv', fields, records)
        sheet = workbook.create_sheet(name)
        sheet.append(fields)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill('solid', fgColor='D9EAD3')
        for row in records:
            sheet.append([row[f] for f in fields])
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions

    with INPUT_PAIRS.open(encoding='utf-8-sig', newline='') as handle:
        reader = csv.DictReader(handle, delimiter='\t')
        write_csv(run / 'five_dogs_allelome_unique_pairs.csv', reader.fieldnames, reader)
    with INPUT_TASKS.open(encoding='utf-8-sig', newline='') as handle:
        reader = csv.DictReader(handle, delimiter='\t')
        write_csv(run / 'five_dogs_allelome_source_tasks.csv', reader.fieldnames, reader)
    workbook.save(run / 'five_dogs_allelome_summary.xlsx')
    log('[4/4] write classification SVG; plot counts represent links expanded by sample')
    write_svg(run / 'five_dogs_allelome_pair_class_summary.svg', category_rows)
    log(f'exported_rows={total["total_links"]}; unique_pairs={len(total["pairs"])}')

def stop(signum, frame):
    raise InterruptedError(f'received_signal={signum}')

def main():
    stamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    log_dir = SCRIPT_DIR / 'log'
    log_dir.mkdir(parents=True, exist_ok=True)
    run = OUTPUT_ROOT / stamp
    for key in ('OMP_NUM_THREADS', 'OPENBLAS_NUM_THREADS', 'MKL_NUM_THREADS', 'NUMEXPR_NUM_THREADS'):
        os.environ[key] = '1'
    with ExitStack() as stack:
        master = stack.enter_context((log_dir / f'{STEP}.{stamp}.log').open('x', encoding='utf-8', buffering=1))
        stack.enter_context(redirect_stdout(master))
        stack.enter_context(redirect_stderr(master))
        handles = {}
        writers = {}
        fields = ['run_id', 'execution_status', 'output_path', 'reason']
        for state in ('success', 'failed'):
            handles[state] = stack.enter_context((log_dir / f'{STEP}.{stamp}.{state}.tsv').open('x', encoding='utf-8', newline='', buffering=1))
            writers[state] = csv.DictWriter(handles[state], fieldnames=fields, delimiter='\t')
            writers[state].writeheader()
        header_end = handles['success'].tell()
        marker_owned = False
        signal.signal(signal.SIGTERM, stop)
        signal.signal(signal.SIGINT, stop)
        try:
            run.mkdir(parents=True, exist_ok=False)
            log(f'run_id={stamp}; output={run}')
            export(run)
            with (run / 'execution_completion.marker').open('x', encoding='utf-8') as handle:
                marker_owned = True
                handle.write(f'execution_status=success\ninput_run={INPUT_RUN}\nrun_id={stamp}\n')
            writers['success'].writerow(dict(run_id=stamp, execution_status='success', output_path=str(run), reason='normal_end'))
            log('task completed / done; execution_status=success')
            log('execution completed does not imply biological QC passed; please separately check count, format and evidence interpretation.')
        except BaseException as error:
            if marker_owned:
                (run / 'execution_completion.marker').unlink(missing_ok=True)
            handles['success'].seek(header_end)
            handles['success'].truncate()
            writers['failed'].writerow(dict(run_id=stamp, execution_status='failed', output_path=str(run), reason=str(error)))
            traceback.print_exc()
            log('execution_status=failed; output generation for this run did not finish normally.')
            raise

if __name__ == '__main__':
    main()
