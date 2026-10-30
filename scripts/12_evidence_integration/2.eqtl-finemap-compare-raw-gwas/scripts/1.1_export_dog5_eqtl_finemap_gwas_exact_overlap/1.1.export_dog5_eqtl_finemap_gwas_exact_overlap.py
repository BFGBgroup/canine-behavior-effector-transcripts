#!/usr/bin/python3.10


from __future__ import annotations

import csv
import math
import os
import tempfile
from collections import defaultdict
from datetime import datetime
from html import escape
from pathlib import Path
from typing import Dict, Iterable, List, Mapping, Sequence

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


DEFAULT_INPUT_RUN_DIR = Path(
    '/path/to/project/analysis_20260817/5-dogs/10.finemaping/formal/9.1_gwas_overlap_expression_ase_catalogue/20260917_191751'
)
DEFAULT_OUTPUT_ROOT = Path(
    '/path/to/project/analysis_20260817/results/5-dogs/2.eqtl-finemap-compare-raw-gwas/1.1_export_dog5_eqtl_finemap_gwas_exact_overlap'
)
RUN_ID = os.environ.get("RUN_ID", datetime.now().strftime("%Y%m%d_%H%M%S"))
INPUT_RUN_DIR = Path(os.environ.get("INPUT_RUN_DIR", str(DEFAULT_INPUT_RUN_DIR)))
INPUT_EXACT_OVERLAPS = (
    INPUT_RUN_DIR / "gwas_catalogue_exact_variant_overlaps_p_le_1e-4.tsv"
)
INPUT_RUN_SUMMARY = INPUT_RUN_DIR / "gwas_catalogue_overlap_run_summary.tsv"
INPUT_COMPLETION_MARKER = INPUT_RUN_DIR / "execution_completion.marker"

OUTPUT_DIR = Path(
    os.environ.get("RUN_DIR", str(DEFAULT_OUTPUT_ROOT / RUN_ID))
)
DETAIL_OUTPUT = OUTPUT_DIR / "five_dogs_eqtl_finemap_gwas_exact_overlaps.csv"
TRAIT_SUMMARY_OUTPUT = (
    OUTPUT_DIR / "five_dogs_eqtl_finemap_gwas_exact_overlap_trait_summary.csv"
)
OVERALL_SUMMARY_OUTPUT = (
    OUTPUT_DIR / "five_dogs_eqtl_finemap_gwas_exact_overlap_overall_summary.csv"
)
XLSX_OUTPUT = OUTPUT_DIR / "five_dogs_eqtl_finemap_gwas_exact_overlap_summary.xlsx"
SVG_OUTPUT = OUTPUT_DIR / "five_dogs_eqtl_finemap_gwas_exact_overlap_trait_summary.svg"

FOCUS_OUTPUT = OUTPUT_DIR / "five_dogs_eqtl_finemap_gwas_exact_overlaps_p_lt_1e-5.csv"


P_TIERS = ("genomewide", "suggestive", "exploratory_point", "extended_exploratory_point")
TIER_SUMMARY_COLUMNS = ("gwas_variant_p_tier", "exact_overlap_rows", "unique_variants",
                        "unique_trait_variants", "unique_features", "new_gene_rows")
TIER_OUTPUT = OUTPUT_DIR / "five_dogs_eqtl_finemap_gwas_variant_p_tier_summary.csv"

FINAL_OUTPUTS = (
    TIER_OUTPUT,
    FOCUS_OUTPUT,
    DETAIL_OUTPUT,
    TRAIT_SUMMARY_OUTPUT,
    OVERALL_SUMMARY_OUTPUT,
    XLSX_OUTPUT,
    SVG_OUTPUT,
)

REQUIRED_EXACT_COLUMNS = (
    "catalogue_role",
    "feature_origin",
    "expression_matrix_scope",
    "interpretation_note",
    "is_top_variant",
    "is_credible_set_member",
    "ppa_tier",
    "block_max_ppa",
    "block_top_ppa_tie_count",
    "block_top_marker_is_unique",
    "block_credible_set_size",
    "block_credible_set_ppa_sum",
    "block_id",
    "tissue",
    "feature_id",
    "feature_primary_id",
    "feature_label",
    "feature_class",
    "expression_gene_id",
    "expression_source",
    "selected_sample_count",
    "common_marker_count",
    "total_reads",
    "chrom",
    "pos",
    "marker_id",
    "ref",
    "alt",
    "phase_scope",
    "local_ps_by_sample",
    "ppa",
    "marker_rank_by_ppa",
    "result_status",
    "result_interpretable",
    "trait",
    "gwas_marker_id",
    "gwas_chrom",
    "gwas_pos",
    "gwas_allele1",
    "gwas_allele0",
    "gwas_af",
    "gwas_beta",
    "gwas_se",
    "gwas_p_lrt",
    "gwas_p_value_column",
    "gwas_p_value",
    "gwas_variant_p_tier",
    "gwas_p_value_status",
    "gwas_modelled_locus_count",
    "gwas_region_match_status",
    "normalized_variant_key",
    "overlap_type",
    "allele_match_status",
    "gwas_allele1_catalogue_relation",
    "gwas_allele0_catalogue_relation",
    "allele_order_status",
    "beta_alignment_status",
    "distance_to_gwas_marker_bp",
)

DETAIL_COLUMNS = (
    "eqtl_finemap_snp",
    "normalized_variant_key",
    "chrom",
    "pos",
    "ref",
    "alt",
    "feature_id",
    "feature_primary_id",
    "feature_label",
    "feature_class",
    "feature_origin",
    "expression_matrix_scope",
    "interpretation_note",
    "expression_gene_id",
    "expression_source",
    "tissue",
    "block_id",
    "finemap_role",
    "is_top_variant",
    "is_credible_set_member",
    "finemap_ppa",
    "ppa_tier",
    "marker_rank_by_ppa",
    "block_max_ppa",
    "block_top_ppa_tie_count",
    "block_top_marker_is_unique",
    "block_credible_set_size",
    "block_credible_set_ppa_sum",
    "selected_sample_count",
    "common_marker_count",
    "total_reads",
    "phase_scope",
    "local_ps_by_sample",
    "gwas_trait",
    "gwas_marker_id",
    "gwas_p_lrt",
    "gwas_variant_p_tier",
    "gwas_p_value_status",
    "gwas_modelled_locus_count",
    "gwas_region_match_status",
    "is_p_lt_1e-5",
    "gwas_af",
    "gwas_beta",
    "gwas_se",
    "gwas_allele1",
    "gwas_allele0",
    "gwas_allele1_catalogue_relation",
    "gwas_allele0_catalogue_relation",
    "allele_order_status",
    "allele_match_status",
    "beta_alignment_status",
    "overlap_type",
)

TRAIT_SUMMARY_COLUMNS = (
    "gwas_trait",
    "exact_overlap_rows",
    "unique_finemap_variants",
    "unique_blocks",
    "unique_features",
    "unique_tissues",
    "reference_annotation_rows",
    "dog5_manual_new_gene_rows",
    "cumulative_p_lt_1e-5_rows",
    "cumulative_p_lt_1e-5_unique_variants",
    "top_variant_rows",
    "credible_set_rows",
    "min_gwas_p_lrt",
    "max_finemap_ppa",
)


def log(message: str) -> None:
    print(
        "[{0} ] {1}".format(datetime.now().strftime("%Y-%m-%d %H:%M:%S"), message),
        flush=True,
    )


def clean(value: object) -> str:
    return "" if value is None else str(value).strip()


def safe_int(value: object) -> int:
    try:
        return int(float(clean(value)))
    except (TypeError, ValueError):
        raise ValueError("cannot parse integer value: {0!r}".format(value))


def safe_float(value: object) -> float:
    try:
        parsed = float(clean(value))
    except (TypeError, ValueError):
        raise ValueError("cannot parse floating point value: {0!r}".format(value))
    if not math.isfinite(parsed):
        raise ValueError("numeric value is not finite value: {0!r}".format(value))
    return parsed


def format_number(value: float) -> str:
    return "{0:.12g}".format(value) if math.isfinite(value) else "NA"


def read_tsv(path: Path) -> tuple[List[str], List[Dict[str, str]]]:
    if not path.is_file() or path.stat().st_size == 0:
        raise FileNotFoundError("input file does not exist or is empty: {0}".format(path))
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle, delimiter="\t")
        rows = [dict(row) for row in reader]
        return list(reader.fieldnames or []), rows


def require_columns(
    path: Path, observed: Sequence[str], required: Sequence[str]
) -> None:
    missing = sorted(set(required) - set(observed))
    if missing:
        raise ValueError(
            "{0} missing required fields: {1}".format(path, ", ".join(missing))
        )


def load_key_value_tsv(path: Path) -> Dict[str, str]:
    fields, rows = read_tsv(path)
    require_columns(path, fields, ("key", "value"))
    values: Dict[str, str] = {}
    for row in rows:
        key = clean(row.get("key"))
        if not key:
            raise ValueError("{0} contains empty key".format(path))
        if key in values:
            raise ValueError("{0} contains duplicate key: {1}".format(path, key))
        values[key] = clean(row.get("value"))
    return values


def load_completion_marker(path: Path) -> Dict[str, str]:
    if not path.is_file() or path.stat().st_size == 0:
        raise FileNotFoundError("upstream execution completion marker does not exist or is empty: {0}".format(path))
    values: Dict[str, str] = {}
    with path.open("r", encoding="utf-8") as handle:
        for line_number, raw_line in enumerate(handle, start=1):
            line = raw_line.rstrip("\n")
            if "=" not in line:
                raise ValueError(
                    "upstream execution completion marker row {0} does not use key=value format".format(
                        line_number
                    )
                )
            key, value = line.split("=", 1)
            key = key.strip()
            if not key or key in values:
                raise ValueError("upstream execution completion marker contains empty or duplicate key: {0!r}".format(key))
            values[key] = value.strip()
    return values


def expected_feature_origin(feature_id: str) -> str:
    if feature_id.startswith("gene-"):
        return "reference_annotation_gene"
    if feature_id.startswith("DOG5_NEWGENE_"):
        return "dog5_manual_new_gene"
    raise ValueError("cannot based on feature_id determine feature_origin: {0}".format(feature_id))


def validate_inputs(
    rows: Sequence[Mapping[str, str]], run_summary: Mapping[str, str]
) -> None:
    required_summary_values = {
        "run_id": INPUT_RUN_DIR.name,
        "expression_matrix_scope": "reference_annotation_plus_dog5_manual_new_gene",
        "feature_origin_rule": (
            "gene-=>reference_annotation_gene;"
            "DOG5_NEWGENE_=>dog5_manual_new_gene"
        ),
        "no_gwas_final_qc_execution_status": "success",
        "no_gwas_final_qc_status": "PASS",
        "no_gwas_final_qc_issue_count": "0",
        "p_value_column": "p_lrt",
        "proximity_analysis_applied": "0",
        "primary_result_type": "exact_variant_allele_set_match",
        "allele_matching_rule": "CHROM_POS_unordered_allele_set",
    }
    for key, expected in required_summary_values.items():
        observed = clean(run_summary.get(key))
        if observed != expected:
            raise ValueError(
                "9.1 run summary field {0} does not satisfy expected: observed={1!r}, expected={2!r}".format(
                    key, observed, expected
                )
            )

    expected_rows = safe_int(run_summary.get("exact_overlap_rows_p_le_1e-4"))
    if len(rows) != expected_rows:
        raise ValueError(
            "Exact-intersection row count differs from the run summary: observed={0}, expected={1}".format(
                len(rows), expected_rows
            )
        )

    threshold = safe_float(run_summary.get("gwas_candidate_p_threshold"))
    strong_threshold = safe_float(
        run_summary.get("suggestive_p_threshold")
    )
    if not math.isclose(threshold, 1e-4, rel_tol=0.0, abs_tol=1e-15):
        raise ValueError("9.1 candidate threshold is not filename convention of 1e-4")
    if not math.isclose(strong_threshold, 1e-5, rel_tol=0.0, abs_tol=1e-15):
        raise ValueError("9.1 cumulative display threshold is not convention of 1e-5")

    for key, expected in (("genomewide_p_threshold", 2.97e-6),
                          ("exploratory_point_p_threshold", 5e-5)):
        if not math.isclose(safe_float(run_summary.get(key)), expected, rel_tol=0, abs_tol=1e-15):
            raise ValueError("9.1 single site tier threshold mismatch: " + key)
    expression_matrix_scope = clean(run_summary.get("expression_matrix_scope"))
    seen_keys = set()
    for index, row in enumerate(rows, start=2):
        ppa = safe_float(row.get("ppa"))
        p_value = safe_float(row.get("gwas_p_value"))
        p_lrt = safe_float(row.get("gwas_p_lrt"))
        if not 0.1 <= ppa <= 1.0 + 1e-8:
            raise ValueError("row {0} PPA falls outside the expected range [0.1,1]".format(index))
        if not 0 <= p_lrt <= threshold:
            raise ValueError(
                "row {0} GWAS P value exceeds threshold {1}".format(index, threshold)
            )
        expected_tier = ("genomewide" if p_lrt < 2.97e-6 else
                         "suggestive" if p_lrt < 1e-5 else
                         "exploratory_point" if p_lrt < 5e-5 else
                         "extended_exploratory_point")
        if clean(row.get("gwas_p_value_status")) != "valid" or clean(row.get("gwas_variant_p_tier")) != expected_tier:
            raise ValueError("row {0} P value status or single-site tier is inconsistent with the original p_lrt".format(index))
        if clean(row.get("gwas_p_value_column")) != "p_lrt":
            raise ValueError("row {0} gwas_p_value_column is not p_lrt".format(index))
        if not math.isclose(p_value, p_lrt, rel_tol=1e-7, abs_tol=0.0):
            raise ValueError("row {0} gwas_p_value differs from gwas_p_lrt".format(index))
        if clean(row.get("result_status")) != "valid":
            raise ValueError("row {0} has an invalid result_status".format(index))
        if clean(row.get("result_interpretable")) != "1":
            raise ValueError("row {0} does not represent an interpretable result".format(index))
        if clean(row.get("allele_match_status")) != "allele_set_match":
            raise ValueError("row {0} fails allele-set matching".format(index))
        if safe_int(row.get("distance_to_gwas_marker_bp")) != 0:
            raise ValueError("row {0} does not represent an exact intersection at the same position".format(index))
        if clean(row.get("chrom")) != clean(row.get("gwas_chrom")):
            raise ValueError("row {0} catalogue and GWAS chromosomes differ".format(index))
        if safe_int(row.get("pos")) != safe_int(row.get("gwas_pos")):
            raise ValueError("row {0} catalogue and GWAS coordinates differ".format(index))

        catalogue_alleles = {
            clean(row.get("ref")).upper(),
            clean(row.get("alt")).upper(),
        }
        gwas_alleles = {
            clean(row.get("gwas_allele1")).upper(),
            clean(row.get("gwas_allele0")).upper(),
        }
        if catalogue_alleles != gwas_alleles:
            raise ValueError("row {0} has mismatched unordered allele sets".format(index))

        feature_id = clean(row.get("feature_id"))
        feature_origin = clean(row.get("feature_origin"))
        if feature_origin != expected_feature_origin(feature_id):
            raise ValueError("row {0} feature_origin is inconsistent with feature_id".format(index))
        if clean(row.get("expression_matrix_scope")) != expression_matrix_scope:
            raise ValueError("row {0} has an unexpected expression_matrix_scope".format(index))

        row_key = (
            clean(row.get("block_id")),
            clean(row.get("marker_id")),
            clean(row.get("trait")),
            clean(row.get("gwas_marker_id")),
        )
        if row_key in seen_keys:
            raise ValueError("exact intersection input contains duplicate workflow key: {0}".format(row_key))
        seen_keys.add(row_key)


def build_detail_rows(
    rows: Sequence[Mapping[str, str]], p_lt_1e5_threshold: float
) -> List[Dict[str, object]]:
    output_rows: List[Dict[str, object]] = []
    for row in rows:
        output_rows.append(
            {
                "eqtl_finemap_snp": clean(row.get("marker_id")),
                "normalized_variant_key": clean(row.get("normalized_variant_key")),
                "chrom": clean(row.get("chrom")),
                "pos": safe_int(row.get("pos")),
                "ref": clean(row.get("ref")),
                "alt": clean(row.get("alt")),
                "feature_id": clean(row.get("feature_id")),
                "feature_primary_id": clean(row.get("feature_primary_id")),
                "feature_label": clean(row.get("feature_label")),
                "feature_class": clean(row.get("feature_class")),
                "feature_origin": clean(row.get("feature_origin")),
                "expression_matrix_scope": clean(
                    row.get("expression_matrix_scope")
                ),
                "expression_gene_id": clean(row.get("expression_gene_id")),
                "expression_source": clean(row.get("expression_source")),

                "interpretation_note": clean(row.get("interpretation_note")),
                "tissue": clean(row.get("tissue")),
                "block_id": clean(row.get("block_id")),
                "finemap_role": clean(row.get("catalogue_role")),
                "is_top_variant": safe_int(row.get("is_top_variant")),
                "is_credible_set_member": safe_int(
                    row.get("is_credible_set_member")
                ),
                "finemap_ppa": safe_float(row.get("ppa")),
                "ppa_tier": clean(row.get("ppa_tier")),
                "marker_rank_by_ppa": safe_int(row.get("marker_rank_by_ppa")),
                "block_max_ppa": safe_float(row.get("block_max_ppa")),
                "block_top_ppa_tie_count": safe_int(
                    row.get("block_top_ppa_tie_count")
                ),
                "block_top_marker_is_unique": safe_int(
                    row.get("block_top_marker_is_unique")
                ),
                "block_credible_set_size": safe_int(
                    row.get("block_credible_set_size")
                ),
                "block_credible_set_ppa_sum": safe_float(
                    row.get("block_credible_set_ppa_sum")
                ),
                "selected_sample_count": safe_int(
                    row.get("selected_sample_count")
                ),
                "common_marker_count": safe_int(row.get("common_marker_count")),
                "total_reads": safe_int(row.get("total_reads")),
                "phase_scope": clean(row.get("phase_scope")),
                "local_ps_by_sample": clean(row.get("local_ps_by_sample")),
                "gwas_trait": clean(row.get("trait")),
                "gwas_marker_id": clean(row.get("gwas_marker_id")),
                "gwas_p_lrt": safe_float(row.get("gwas_p_lrt")),
                "gwas_variant_p_tier": clean(row.get("gwas_variant_p_tier")),
                "gwas_p_value_status": clean(row.get("gwas_p_value_status")),
                "gwas_modelled_locus_count": safe_int(row.get("gwas_modelled_locus_count")),
                "gwas_region_match_status": clean(row.get("gwas_region_match_status")),
                "is_p_lt_1e-5": int(
                    safe_float(row.get("gwas_p_lrt"))
                    < p_lt_1e5_threshold
                ),
                "gwas_af": safe_float(row.get("gwas_af")),
                "gwas_beta": safe_float(row.get("gwas_beta")),
                "gwas_se": safe_float(row.get("gwas_se")),
                "gwas_allele1": clean(row.get("gwas_allele1")),
                "gwas_allele0": clean(row.get("gwas_allele0")),
                "gwas_allele1_catalogue_relation": clean(
                    row.get("gwas_allele1_catalogue_relation")
                ),
                "gwas_allele0_catalogue_relation": clean(
                    row.get("gwas_allele0_catalogue_relation")
                ),
                "allele_order_status": clean(row.get("allele_order_status")),
                "allele_match_status": clean(row.get("allele_match_status")),
                "beta_alignment_status": clean(row.get("beta_alignment_status")),
                "overlap_type": clean(row.get("overlap_type")),
            }
        )
    return sorted(
        output_rows,
        key=lambda row: (
            clean(row["gwas_trait"]),
            float(row["gwas_p_lrt"]),
            clean(row["feature_label"]),
            clean(row["tissue"]),
            clean(row["eqtl_finemap_snp"]),
        ),
    )


def build_trait_summary(
    detail_rows: Sequence[Mapping[str, object]]
) -> List[Dict[str, object]]:
    grouped: Dict[str, List[Mapping[str, object]]] = defaultdict(list)
    for row in detail_rows:
        grouped[clean(row.get("gwas_trait"))].append(row)

    summary_rows: List[Dict[str, object]] = []
    for trait in sorted(grouped):
        rows = grouped[trait]
        summary_rows.append(
            {
                "gwas_trait": trait,
                "exact_overlap_rows": len(rows),
                "unique_finemap_variants": len(
                    {clean(row.get("normalized_variant_key")) for row in rows}
                ),
                "unique_blocks": len(
                    {clean(row.get("block_id")) for row in rows}
                ),
                "unique_features": len(
                    {clean(row.get("feature_id")) for row in rows}
                ),
                "unique_tissues": len(
                    {clean(row.get("tissue")) for row in rows}
                ),
                "reference_annotation_rows": sum(
                    clean(row.get("feature_origin"))
                    == "reference_annotation_gene"
                    for row in rows
                ),
                "dog5_manual_new_gene_rows": sum(
                    clean(row.get("feature_origin")) == "dog5_manual_new_gene"
                    for row in rows
                ),
                "cumulative_p_lt_1e-5_rows": sum(
                    int(row.get("is_p_lt_1e-5", 0))
                    for row in rows
                ),
                "cumulative_p_lt_1e-5_unique_variants": len(
                    {
                        clean(row.get("normalized_variant_key"))
                        for row in rows
                        if int(row.get("is_p_lt_1e-5", 0))
                        == 1
                    }
                ),
                "top_variant_rows": sum(
                    int(row.get("is_top_variant", 0)) for row in rows
                ),
                "credible_set_rows": sum(
                    int(row.get("is_credible_set_member", 0)) for row in rows
                ),
                "min_gwas_p_lrt": min(
                    float(row["gwas_p_lrt"]) for row in rows
                ),
                "max_finemap_ppa": max(
                    float(row["finemap_ppa"]) for row in rows
                ),
            }
        )
    return summary_rows


def build_overall_summary(
    detail_rows: Sequence[Mapping[str, object]],
    run_summary: Mapping[str, str],
) -> List[Dict[str, object]]:
    return [
        {"metric": "run_id", "value": RUN_ID},
        {"metric": "genomewide_p_threshold", "value": clean(run_summary.get("genomewide_p_threshold"))},
        {"metric": "exploratory_point_p_threshold", "value": clean(run_summary.get("exploratory_point_p_threshold"))},
        {"metric": "variant_p_tier_definition", "value": "SNP p_lrt; independent of region analysis_scope; four exclusive tiers within p<=1e-4"},
        {"metric": "focus_filter", "value": "gwas_p_lrt < 1e-5; input pool remains <=1e-4"},
        {"metric": "credible_set_scope", "value": "Expression–ASE only; no GWAS fine-map input"},
        {"metric": "source_9.1_run_id", "value": clean(run_summary.get("run_id"))},
        {"metric": "source_exact_overlap_table", "value": str(INPUT_EXACT_OVERLAPS)},
        {
            "metric": "expression_matrix_scope",
            "value": clean(run_summary.get("expression_matrix_scope")),
        },
        {
            "metric": "feature_origin_rule",
            "value": clean(run_summary.get("feature_origin_rule")),
        },
        {"metric": "gwas_p_value_column", "value": clean(run_summary.get("p_value_column"))},
        {
            "metric": "gwas_candidate_p_threshold",
            "value": clean(run_summary.get("gwas_candidate_p_threshold")),
        },
        {
            "metric": "suggestive_p_threshold",
            "value": clean(run_summary.get("suggestive_p_threshold")),
        },
        {
            "metric": "allele_matching_rule",
            "value": clean(run_summary.get("allele_matching_rule")),
        },
        {
            "metric": "proximity_analysis_applied",
            "value": clean(run_summary.get("proximity_analysis_applied")),
        },
        {"metric": "exact_overlap_rows", "value": len(detail_rows)},
        {
            "metric": "unique_finemap_variants",
            "value": len(
                {clean(row.get("normalized_variant_key")) for row in detail_rows}
            ),
        },
        {
            "metric": "unique_trait_variant_pairs",
            "value": len(
                {
                    (
                        clean(row.get("gwas_trait")),
                        clean(row.get("normalized_variant_key")),
                    )
                    for row in detail_rows
                }
            ),
        },
        {
            "metric": "unique_blocks",
            "value": len({clean(row.get("block_id")) for row in detail_rows}),
        },
        {
            "metric": "unique_features",
            "value": len({clean(row.get("feature_id")) for row in detail_rows}),
        },
        {
            "metric": "unique_tissues",
            "value": len({clean(row.get("tissue")) for row in detail_rows}),
        },
        {
            "metric": "unique_gwas_traits",
            "value": len({clean(row.get("gwas_trait")) for row in detail_rows}),
        },
        {
            "metric": "top_variant_rows",
            "value": sum(int(row.get("is_top_variant", 0)) for row in detail_rows),
        },
        {
            "metric": "credible_set_rows",
            "value": sum(
                int(row.get("is_credible_set_member", 0)) for row in detail_rows
            ),
        },
        {
            "metric": "reference_annotation_rows",
            "value": sum(
                clean(row.get("feature_origin")) == "reference_annotation_gene"
                for row in detail_rows
            ),
        },
        {
            "metric": "dog5_manual_new_gene_rows",
            "value": sum(
                clean(row.get("feature_origin")) == "dog5_manual_new_gene"
                for row in detail_rows
            ),
        },
        {
            "metric": "cumulative_p_lt_1e-5_rows",
            "value": sum(
                int(row.get("is_p_lt_1e-5", 0))
                for row in detail_rows
            ),
        },
        {
            "metric": "cumulative_p_lt_1e-5_unique_variants",
            "value": len(
                {
                    clean(row.get("normalized_variant_key"))
                    for row in detail_rows
                    if int(row.get("is_p_lt_1e-5", 0)) == 1
                }
            ),
        },
        {
            "metric": "min_gwas_p_lrt",
            "value": format_number(
                min((float(row["gwas_p_lrt"]) for row in detail_rows), default=float("nan"))
            ),
        },
        {
            "metric": "max_finemap_ppa",
            "value": format_number(
                max((float(row["finemap_ppa"]) for row in detail_rows), default=float("nan"))
            ),
        },
        {
            "metric": "interpretation_boundary",
            "value": (
                "Exact overlap at the same coordinate and with the same unordered allele set, with p_lrt<=1e-4; "
                "p_lrt<1e-5 defines only a priority tier, not genome-wide GWAS significance, "
                "LD colocalization, or proof of causality; molecular fine-mapping is exploratory expression+ASE analysis without association/FDR screening or comparison against a null-causal model."
            ),
        },
    ]


def build_tier_summary(detail_rows: Sequence[Mapping[str, object]]) -> List[Dict[str, object]]:

    result = []
    for tier in P_TIERS:
        rows = [row for row in detail_rows if row["gwas_variant_p_tier"] == tier]
        result.append({
            "gwas_variant_p_tier": tier,
            "exact_overlap_rows": len(rows),
            "unique_variants": len({row["normalized_variant_key"] for row in rows}),
            "unique_trait_variants": len({(row["gwas_trait"], row["normalized_variant_key"]) for row in rows}),
            "unique_features": len({row["feature_id"] for row in rows}),
            "new_gene_rows": sum(row["feature_origin"] == "dog5_manual_new_gene" for row in rows),
        })
    return result


def write_csv(
    path: Path,
    fieldnames: Sequence[str],
    rows: Iterable[Mapping[str, object]],
) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames, lineterminator="\n")
        writer.writeheader()
        for row in rows:
            writer.writerow({field: row.get(field, "") for field in fieldnames})


def append_sheet(
    workbook: Workbook,
    title: str,
    fieldnames: Sequence[str],
    rows: Sequence[Mapping[str, object]],
) -> None:
    sheet = workbook.create_sheet(title=title)
    sheet.append(list(fieldnames))
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="365F91")
    for row in rows:
        sheet.append([row.get(field, "") for field in fieldnames])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column_cells in sheet.columns:
        width = min(
            60,
            max(10, max(len(clean(cell.value)) for cell in column_cells) + 2),
        )
        sheet.column_dimensions[column_cells[0].column_letter].width = width


def write_xlsx(
    path: Path,
    detail_rows: Sequence[Mapping[str, object]],
    trait_rows: Sequence[Mapping[str, object]],
    overall_rows: Sequence[Mapping[str, object]],
) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    append_sheet(workbook, "exact_overlaps", DETAIL_COLUMNS, detail_rows)
    append_sheet(workbook, "variant_p_tier_summary", TIER_SUMMARY_COLUMNS, build_tier_summary(detail_rows))
    append_sheet(workbook, "trait_summary", TRAIT_SUMMARY_COLUMNS, trait_rows)
    append_sheet(workbook, "overall_summary", ("metric", "value"), overall_rows)
    append_sheet(workbook, "focus_p_lt_1e-5", DETAIL_COLUMNS,
                 [row for row in detail_rows if row["is_p_lt_1e-5"] == 1])
    append_sheet(workbook, "field_definitions", ("field", "meaning"), [
        {"field": "marker_rank_by_ppa", "meaning": "PPA rank among all variants in the Expression–ASE block, not a GWAS rank or a within-credible-set rank; the upstream rank is retained without reranking within the GWAS-overlap subset."},
        {"field": "block_top_ppa_tie_count", "meaning": "Number of variants tied for the highest PPA in the entire Expression–ASE block; it describes ties only when the current variant has the highest PPA, not ties at any other current rank."},
        {"field": "is_credible_set_member", "meaning": "Expression–ASE credible-set member; the raw GWAS side has no credible set."},
        {"field": "block_credible_set_size", "meaning": "Number of variants in the Expression–ASE block credible set."},
        {"field": "is_p_lt_1e-5", "meaning": "Cumulative P<1e-5 flag derived in this step, including genomewide and suggestive variants; it does not indicate regional primary status or change upstream single-variant categories."},
        {"field": "gwas_variant_p_tier", "meaning": "Classification by the SNP's own p_lrt: <2.97e-6; [2.97e-6,1e-5); [1e-5,5e-5); [5e-5,1e-4]; not a GWAS region classification."},
        {"field": "gwas_region_match_status", "meaning": "Copies only the main modeled-variant match status from 9.1; not used for filtering in this step and does not indicate GWAS credible-set membership."},
        {"field": "interpretation_note", "meaning": "Retains the molecular interpretation limits: exploratory joint expression+ASE fine-mapping without association/FDR screening or a null-causal model; PPA or CS membership does not establish significance."},
        {"field": "finemap_ppa", "meaning": "Expression–ASE model posterior priority, not a GWAS PIP or association p value."},
    ])
    workbook.save(path)


def write_svg(
    path: Path, trait_rows: Sequence[Mapping[str, object]]
) -> None:
    width = 1000
    left = 360
    top = 105
    row_height = 58
    chart_width = 300
    height = top + row_height * len(trait_rows) + 100
    maximum = max((int(row["unique_finemap_variants"]) for row in trait_rows), default=0)

    elements = [
        '<svg xmlns="http://www.w3.org/2000/svg" width="{0}" height="{1}" '
        'viewBox="0 0 {0} {1}">'.format(width, height),
        '<rect width="100%" height="100%" fill="#ffffff"/>',
        '<text x="40" y="42" font-family="Arial" font-size="24" '
        'font-weight="bold" fill="#1f2937">Dog5 expression–ASE fine-map × raw GWAS</text>',
        '<text x="40" y="70" font-family="Arial" font-size="14" fill="#4b5563">'
        "Exact CanFam3.1 position + unordered allele set; GWAS p_lrt ≤ 1e-4"
        "</text>",
    ]
    if not trait_rows:
        elements.append('<text x="40" y="130" font-family="Arial" font-size="16">No exact overlaps at the specified thresholds.</text>')
    for index, row in enumerate(trait_rows):
        y = top + index * row_height
        count = int(row["unique_finemap_variants"])
        expanded_rows = int(row["exact_overlap_rows"])
        strong_count = int(row["cumulative_p_lt_1e-5_unique_variants"])
        bar_width = 0 if maximum == 0 else int(chart_width * count / maximum)
        label = escape(clean(row["gwas_trait"]))
        elements.extend(
            [
                '<text x="40" y="{0}" font-family="Arial" font-size="14" '
                'fill="#111827">{1}</text>'.format(y + 21, label),
                '<rect x="{0}" y="{1}" width="{2}" height="26" rx="2" '
                'fill="#3b82f6"/>'.format(left, y, bar_width),
                '<text x="{0}" y="{1}" font-family="Arial" font-size="14" '
                'font-weight="bold" fill="#111827">{2} variants; {3} rows; '
                '{4} at p&lt;1e-5</text>'.format(
                    left + bar_width + 10,
                    y + 19,
                    count,
                    expanded_rows,
                    strong_count,
                ),
            ]
        )
    elements.extend(
        [
            '<text x="40" y="{0}" font-family="Arial" font-size="12" '
            'fill="#6b7280">Bar length is unique normalized variants; rows retain '
            "feature and tissue expansions.</text>".format(height - 35),
            "</svg>",
        ]
    )
    path.write_text("\n".join(elements) + "\n", encoding="utf-8")


def main() -> None:
    if not INPUT_RUN_DIR.is_absolute() or not OUTPUT_DIR.is_absolute():
        raise ValueError("input run directory and output directory must use absolute path")
    if OUTPUT_DIR.name != RUN_ID:
        raise ValueError("output directory end level name must be with RUN_ID consistent")

    log("[1/5] start export five dogs expression+ASE fine-map × raw GWAS exact intersection")
    log("[1/5] exact_overlap_input={0}".format(INPUT_EXACT_OVERLAPS))
    log("[1/5] run_summary_input={0}".format(INPUT_RUN_SUMMARY))
    log("[1/5] completion_marker_input={0}".format(INPUT_COMPLETION_MARKER))
    log("[1/5] output_dir={0}".format(OUTPUT_DIR))

    log("[2/5] read and check fixed of 9.1 input")
    fields, input_rows = read_tsv(INPUT_EXACT_OVERLAPS)
    require_columns(INPUT_EXACT_OVERLAPS, fields, REQUIRED_EXACT_COLUMNS)
    run_summary = load_key_value_tsv(INPUT_RUN_SUMMARY)
    completion_marker = load_completion_marker(INPUT_COMPLETION_MARKER)
    if completion_marker.get("run_id") != INPUT_RUN_DIR.name:
        raise ValueError("run_id in the Step 9.1 execution completion marker does not match the input directory")
    if completion_marker.get("execution_status") != "success":
        raise ValueError("9.1 execution completion marker is not success")
    validate_inputs(input_rows, run_summary)
    log("[2/5] validated_exact_overlap_rows={0}".format(len(input_rows)))

    log("[3/5] build display details and trait summary")
    p_lt_1e5_threshold = safe_float(
        run_summary.get("suggestive_p_threshold")
    )
    detail_rows = build_detail_rows(input_rows, p_lt_1e5_threshold)
    trait_rows = build_trait_summary(detail_rows)
    overall_rows = build_overall_summary(detail_rows, run_summary)
    log(
        "[3/5] detail_rows={0}; unique_variants={1}; traits={2}; "
        "cumulative_p_lt_1e-5_rows={3}".format(
            len(detail_rows),
            len(
                {
                    clean(row.get("normalized_variant_key"))
                    for row in detail_rows
                }
            ),
            len(trait_rows),
            sum(
                int(row.get("is_p_lt_1e-5", 0))
                for row in detail_rows
            ),
        )
    )

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    existing = [str(path) for path in FINAL_OUTPUTS if path.exists()]
    if existing:
        raise FileExistsError(
            "using under main output already exists, refuse to overwrite: {0}".format(", ".join(existing))
        )

    log("[4/5] write CSV, XLSX and SVG")
    with tempfile.TemporaryDirectory(
        prefix=".1.1_export_staging.", dir=OUTPUT_DIR
    ) as staging_dir_name:
        staging_dir = Path(staging_dir_name)
        staged = {path: staging_dir / path.name for path in FINAL_OUTPUTS}
        write_csv(staged[DETAIL_OUTPUT], DETAIL_COLUMNS, detail_rows)
        write_csv(staged[TIER_OUTPUT], TIER_SUMMARY_COLUMNS, build_tier_summary(detail_rows))
        write_csv(staged[FOCUS_OUTPUT], DETAIL_COLUMNS,
                  [row for row in detail_rows if row["is_p_lt_1e-5"] == 1])
        write_csv(
            staged[TRAIT_SUMMARY_OUTPUT],
            TRAIT_SUMMARY_COLUMNS,
            trait_rows,
        )
        write_csv(
            staged[OVERALL_SUMMARY_OUTPUT],
            ("metric", "value"),
            overall_rows,
        )
        write_xlsx(staged[XLSX_OUTPUT], detail_rows, trait_rows, overall_rows)
        write_svg(staged[SVG_OUTPUT], trait_rows)
        for final_path in FINAL_OUTPUTS:
            os.replace(staged[final_path], final_path)

    log("[5/5] export completed")
    for path in FINAL_OUTPUTS:
        log("[output] {0}".format(path))
    log("[5/5] Python main program completed, waiting for the launcher to write the execution-completion marker")


if __name__ == "__main__":
    main()
