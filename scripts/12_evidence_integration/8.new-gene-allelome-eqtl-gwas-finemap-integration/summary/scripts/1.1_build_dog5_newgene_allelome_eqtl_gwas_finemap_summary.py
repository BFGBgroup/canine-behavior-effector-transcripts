#!/usr/bin/python3
# -*- coding: utf-8 -*-


from __future__ import annotations

import argparse
import csv
import gzip
import html
import importlib.util
import math
import os
import re
import sys
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Dict, Iterable, List, Mapping, Sequence, Tuple
from xml.sax.saxutils import escape


OPENPYXL_PYTHON = Path("/usr/bin/python3")
if importlib.util.find_spec("openpyxl") is None:
    current_python = Path(sys.executable).resolve()
    fallback_python = OPENPYXL_PYTHON.resolve()
    if current_python == fallback_python or not fallback_python.is_file():
        raise ModuleNotFoundError(
            "missing openpyxl, and cannot switch to already verify of/usr/bin/python3"
        )
    os.execv(
        str(fallback_python),
        [str(fallback_python), str(Path(__file__).resolve()), *sys.argv[1:]],
    )

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


SOURCE_RUN_DIR = Path(
    '/path/to/project/analysis_20260817/results/5-dogs/8.new-gene-allelome-eqtl-gwas-finemap-integration/1.1_integrate_dog5_newgene_allelome_eqtl_gwas_finemap/20260910_195332'
)
DEFAULT_CANDIDATES = (
    SOURCE_RUN_DIR / "dog5_newgene_eqtl_gwas_finemap_all_candidates.tsv"
)
DEFAULT_PCG_LINKS = (
    SOURCE_RUN_DIR / "dog5_integrated_newgene_pcGene_links.tsv.gz"
)
DEFAULT_PAIRS = (
    SOURCE_RUN_DIR / "dog5_candidate_newgene_pcGene_mechanism_pairs.tsv"
)
DEFAULT_SOURCE_SUMMARY = (
    SOURCE_RUN_DIR
    / "dog5_newgene_allelome_eqtl_gwas_finemap_summary.tsv"
)
DEFAULT_SOURCE_MANIFEST = (
    SOURCE_RUN_DIR
    / "dog5_newgene_allelome_eqtl_gwas_finemap_manifest.tsv"
)
DEFAULT_FIVE_ONLY_NEW_GENE_MAP = Path(
    '/path/to/project/analysis_20260817/reference/Manual-annotation/1.1_build_five_only_new_gene_gtf/20260818_204722/five_only_new_gene_loci.transcript_map.tsv'
)
DEFAULT_OUTPUT_ROOT = Path(
    '/path/to/project/analysis_20260817/results/5-dogs/8.new-gene-allelome-eqtl-gwas-finemap-integration/summary/results'
)

FOCUS_TIER = "A_BOTH_95PCT_CREDIBLE_PLUS_ALLELOME"
GWAS_BRANCH = "main"
FIVE_ONLY_SOURCE_DATASET = "five-dogs"
FIVE_ONLY_SELECTION_STAGE = "2.1_transcript_first_preASE_candidate_assessment"
FIVE_ONLY_MERGE_STATUS = "five_only_not_cross_resource_merged"

GWAS_MEMBER_REQUIRED = {
    "locus_id",
    "trait",
    "analysis_scope",
    "branch",
    "cs_id",
    "chrom",
    "pos",
    "ld_variant_id",
    "ld_a1",
    "ld_a2",
    "pip",
    "p_lrt",
}

TIER_LABELS = {
    "A_BOTH_95PCT_CREDIBLE_PLUS_ALLELOME": (
        "Tier A: shared eQTL and GWAS 95% credible-set variants with Allelome in the same block"
    ),
    "B_GWAS_NONCREDIBLE_PLUS_ALLELOME": (
        "Tier B: overlap between eQTL and GWAS-locus non-credible-set SNPs with Allelome in the same block"
    ),
    "C_BOTH_95PCT_CREDIBLE_NO_ALLELOME": (
        "Tier C: shared eQTL and GWAS 95% credible-set variants without Allelome in the same block"
    ),
    "D_GWAS_NONCREDIBLE_NO_ALLELOME": (
        "Tier D: diagnostic GWAS non-credible-set overlap without Allelome in the same block"
    ),
}

CANDIDATE_REQUIRED = {
    "allelome_evidence_strata",
    "allelome_multigene_link_rows",
    "allelome_overlap_link_rows",
    "candidate_id",
    "new_gene_id",
    "result_layer",
    "normalized_variant_key",
    "eqtl_block_id",
    "eqtl_tissue",
    "eqtl_feature_id",
    "eqtl_feature_primary_id",
    "eqtl_feature_origin",
    "eqtl_expression_gene_id",
    "eqtl_expression_source",
    "eqtl_expression_matrix_scope",
    "eqtl_ppa",
    "eqtl_phase_scope",
    "gwas_analysis_scope",
    "gwas_branch",
    "gwas_locus_id",
    "gwas_trait",
    "allele_match_status",
    "gwas_overlap_snp_id",
    "gwas_overlap_snp_chrom",
    "gwas_overlap_snp_pos",
    "gwas_overlap_snp_p_lrt",
    "gwas_overlap_snp_pip",
    "gwas_overlap_snp_pip_rank",
    "gwas_overlap_snp_in_credible_set",
    "gwas_credible_sets",
    "gwas_variant_count",
    "gwas_locus_lead_snp_id",
    "gwas_locus_lead_chrom",
    "gwas_locus_lead_pos",
    "gwas_locus_lead_allele1",
    "gwas_locus_lead_allele0",
    "gwas_locus_lead_p_lrt",
    "gwas_locus_lead_pip",
    "gwas_locus_lead_pip_rank",
    "gwas_locus_lead_credible_sets",
    "gwas_locus_lead_in_credible_set",
    "overlap_snp_equals_lead",
    "gwas_locus_top_pip_snp_id",
    "gwas_locus_top_pip_chrom",
    "gwas_locus_top_pip_pos",
    "gwas_locus_top_pip_allele1",
    "gwas_locus_top_pip_allele0",
    "gwas_locus_top_pip_p_lrt",
    "gwas_locus_top_pip",
    "gwas_locus_top_pip_rank",
    "gwas_locus_top_pip_credible_sets",
    "gwas_locus_top_pip_in_credible_set",
    "overlap_snp_equals_top_pip",
    "allelome_block_match",
    "multilayer_integration_scope",
    "shared_snp_equals_any_allelome_endpoint",
    "shared_snp_min_distance_to_allelome_endpoint_bp",
    "supporting_dog_count",
    "supporting_dogs",
    "allelome_link_count",
    "pcgene_link_count",
    "partner_gene_count",
    "endpoint_manual_new_gene_strands",
    "endpoint_manual_new_gene_class_codes",
    "endpoint_manual_new_gene_novelty_categories",
    "endpoint_manual_new_gene_candidate_categories",
    "endpoint_manual_new_gene_dataset_support",
    "endpoint_manual_new_gene_merge_statuses",
    "evidence_tier",
    "main_result",
    "evidence_interpretation",
}

PAIR_REQUIRED = {
    "allelome_evidence_strata",
    "allelome_multigene_link_rows",
    "allelome_overlap_link_rows",
    "evidence_tier",
    "main_result",
    "candidate_id",
    "new_gene_id",
    "eqtl_block_id",
    "eqtl_tissue",
    "normalized_variant_key",
    "gwas_trait",
    "gwas_locus_id",
    "gwas_overlap_snp_id",
    "gwas_overlap_snp_chrom",
    "gwas_overlap_snp_pos",
    "gwas_overlap_snp_p_lrt",
    "gwas_overlap_snp_pip",
    "gwas_overlap_snp_pip_rank",
    "gwas_overlap_snp_in_credible_set",
    "gwas_locus_lead_snp_id",
    "gwas_locus_lead_chrom",
    "gwas_locus_lead_pos",
    "gwas_locus_lead_allele1",
    "gwas_locus_lead_allele0",
    "gwas_locus_lead_p_lrt",
    "gwas_locus_lead_pip",
    "gwas_locus_lead_pip_rank",
    "gwas_locus_lead_credible_sets",
    "gwas_locus_lead_in_credible_set",
    "overlap_snp_equals_lead",
    "gwas_locus_top_pip_snp_id",
    "gwas_locus_top_pip_chrom",
    "gwas_locus_top_pip_pos",
    "gwas_locus_top_pip_allele1",
    "gwas_locus_top_pip_allele0",
    "gwas_locus_top_pip_p_lrt",
    "gwas_locus_top_pip",
    "gwas_locus_top_pip_rank",
    "gwas_locus_top_pip_credible_sets",
    "gwas_locus_top_pip_in_credible_set",
    "overlap_snp_equals_top_pip",
    "multilayer_integration_scope",
    "shared_snp_equals_any_allelome_endpoint",
    "shared_snp_min_distance_to_allelome_endpoint_bp",
    "partner_gene_id",
    "partner_gene_symbol",
    "mechanism",
    "expanded_link_rows",
    "supporting_dog_count",
    "supporting_dogs",
    "allelome_endpoint_count",
    "allelome_endpoints",
    "linkage_score_min",
    "linkage_score_max",
    "interval_gap_bp_min",
    "interval_gap_bp_max",
    "candidate_orientation",
    "direction_interpretation",
}

LINK_REQUIRED = {
    "sample_pair_id",
    "evidence_stratum",
    "either_endpoint_multigene",
    "shared_variant_identity_checked",
    "source_link_file",
    "source_link_line",
    "candidate_id",
    "new_gene_id",
    "evidence_tier",
    "main_result",
    "normalized_variant_key",
    "eqtl_block_id",
    "eqtl_tissue",
    "eqtl_ppa",
    "eqtl_phase_scope",
    "gwas_trait",
    "gwas_locus_id",
    "gwas_overlap_snp_id",
    "gwas_overlap_snp_chrom",
    "gwas_overlap_snp_pos",
    "gwas_overlap_snp_p_lrt",
    "gwas_overlap_snp_pip",
    "gwas_overlap_snp_pip_rank",
    "gwas_overlap_snp_in_credible_set",
    "gwas_credible_sets",
    "gwas_variant_count",
    "gwas_locus_lead_snp_id",
    "gwas_locus_lead_chrom",
    "gwas_locus_lead_pos",
    "gwas_locus_lead_allele1",
    "gwas_locus_lead_allele0",
    "gwas_locus_lead_p_lrt",
    "gwas_locus_lead_pip",
    "gwas_locus_lead_pip_rank",
    "gwas_locus_lead_credible_sets",
    "gwas_locus_lead_in_credible_set",
    "overlap_snp_equals_lead",
    "gwas_locus_top_pip_snp_id",
    "gwas_locus_top_pip_chrom",
    "gwas_locus_top_pip_pos",
    "gwas_locus_top_pip_allele1",
    "gwas_locus_top_pip_allele0",
    "gwas_locus_top_pip_p_lrt",
    "gwas_locus_top_pip",
    "gwas_locus_top_pip_rank",
    "gwas_locus_top_pip_credible_sets",
    "gwas_locus_top_pip_in_credible_set",
    "overlap_snp_equals_top_pip",
    "multilayer_integration_scope",
    "animal_id",
    "sample_tissue",
    "canonical_pair_id",
    "interval_gap_bp",
    "mechanism",
    "linkage_score",
    "new_gene_endpoint_manual_ids",
    "shared_finemap_snp_equals_allelome_endpoint",
    "shared_finemap_snp_min_distance_to_allelome_endpoint_bp",
    "partner_gene_id",
    "partner_gene_symbol",
    "partner_gene_strand",
    "partner_chrom",
    "partner_start",
    "partner_end",
    "partner_biotype",
    "partner_endpoint_type",
    "candidate_orientation",
    "direction_basis",
}

FIVE_ONLY_TRANSCRIPT_MAP_REQUIRED = {
    "final_gene_id",
    "final_transcript_id",
    "source_dataset",
    "selection_stage",
    "chr",
    "strand",
    "gene_start_0based",
    "gene_end_0based",
    "class_code",
    "novelty_category",
    "FEELnc_biotype",
}

BLOCK_COLUMNS = [
    "allelome_evidence_strata",
    "allelome_multigene_link_rows",
    "allelome_overlap_link_rows",
    "evidence_tier",
    "tier_meaning",
    "main_result",
    "candidate_id",
    "new_gene_id",
    "new_gene_strand",
    "new_gene_class_code",
    "new_gene_novelty_category",
    "new_gene_candidate_category",
    "new_gene_feelnc_biotype",
    "new_gene_coding_class",
    "new_gene_interval",
    "new_gene_span_bp",
    "new_gene_dataset_support",
    "new_gene_merge_status",
    "eqtl_block_id",
    "eqtl_tissue",
    "eqtl_finemap_variant",
    "eqtl_ppa",
    "gwas_trait",
    "gwas_locus_id",
    "gwas_overlap_snp",
    "gwas_overlap_snp_p_lrt",
    "gwas_overlap_snp_pip",
    "gwas_overlap_snp_pip_rank",
    "gwas_overlap_snp_in_credible_set",
    "gwas_credible_sets",
    "gwas_variant_count",
    "gwas_locus_lead_snp",
    "gwas_locus_lead_alleles",
    "gwas_locus_lead_p_lrt",
    "gwas_locus_lead_pip",
    "gwas_locus_lead_pip_rank",
    "gwas_locus_lead_credible_sets",
    "gwas_locus_lead_in_credible_set",
    "overlap_snp_equals_lead",
    "gwas_locus_top_pip_snp",
    "gwas_locus_top_pip_alleles",
    "gwas_locus_top_pip_p_lrt",
    "gwas_locus_top_pip",
    "gwas_locus_top_pip_rank",
    "gwas_locus_top_pip_credible_sets",
    "gwas_locus_top_pip_in_credible_set",
    "overlap_snp_equals_top_pip",
    "allelome_block_match",
    "multilayer_integration_scope",
    "shared_snp_equals_any_allelome_endpoint",
    "shared_snp_min_distance_to_allelome_endpoint_bp",
    "supporting_dog_count",
    "supporting_dogs",
    "allelome_link_count",
    "pcgene_link_count",
    "partner_gene_count",
    "interpretation",
]

PAIR_COLUMNS = [
    "allelome_evidence_strata",
    "allelome_multigene_link_rows",
    "allelome_overlap_link_rows",
    "evidence_tier",
    "tier_meaning",
    "main_result",
    "candidate_id",
    "new_gene_id",
    "new_gene_class_code",
    "new_gene_novelty_category",
    "new_gene_candidate_category",
    "new_gene_feelnc_biotype",
    "new_gene_coding_class",
    "new_gene_interval",
    "new_gene_strand",
    "eqtl_tissue",
    "eqtl_finemap_variant",
    "eqtl_ppa",
    "gwas_trait",
    "gwas_overlap_snp",
    "gwas_overlap_snp_p_lrt",
    "gwas_overlap_snp_pip",
    "gwas_overlap_snp_pip_rank",
    "gwas_overlap_snp_in_credible_set",
    "gwas_locus_lead_snp",
    "gwas_locus_lead_p_lrt",
    "gwas_locus_lead_pip",
    "gwas_locus_lead_pip_rank",
    "gwas_locus_lead_credible_sets",
    "gwas_locus_lead_in_credible_set",
    "overlap_snp_equals_lead",
    "gwas_locus_top_pip_snp",
    "gwas_locus_top_pip_alleles",
    "gwas_locus_top_pip_p_lrt",
    "gwas_locus_top_pip",
    "gwas_locus_top_pip_rank",
    "gwas_locus_top_pip_credible_sets",
    "gwas_locus_top_pip_in_credible_set",
    "overlap_snp_equals_top_pip",
    "partner_gene_id",
    "partner_gene_symbol",
    "partner_biotype",
    "partner_gene_interval",
    "new_gene_pcgene_overlap",
    "new_gene_pcgene_overlap_bp",
    "overlap_percent_of_new_gene",
    "overlap_percent_of_pcgene",
    "mechanism",
    "expanded_link_rows",
    "supporting_dog_count",
    "supporting_dogs",
    "allelome_endpoint_count",
    "allelome_endpoints",
    "shared_snp_equals_any_allelome_endpoint",
    "shared_snp_min_distance_to_allelome_endpoint_bp",
    "linkage_score_min",
    "linkage_score_max",
    "interval_gap_bp_min",
    "interval_gap_bp_max",
    "candidate_orientation",
    "direction_interpretation",
]

DETAIL_COLUMNS = [
    "sample_pair_id",
    "evidence_stratum",
    "either_endpoint_multigene",
    "shared_variant_identity_checked",
    "source_link_file",
    "source_link_line",
    "evidence_tier",
    "main_result",
    "candidate_id",
    "new_gene_id",
    "new_gene_class_code",
    "new_gene_novelty_category",
    "new_gene_candidate_category",
    "new_gene_feelnc_biotype",
    "new_gene_coding_class",
    "new_gene_interval",
    "new_gene_strand",
    "eqtl_block_id",
    "eqtl_tissue",
    "eqtl_finemap_variant",
    "eqtl_ppa",
    "gwas_trait",
    "gwas_locus_id",
    "gwas_overlap_snp",
    "gwas_overlap_snp_p_lrt",
    "gwas_overlap_snp_pip",
    "gwas_overlap_snp_pip_rank",
    "gwas_overlap_snp_in_credible_set",
    "gwas_credible_sets",
    "gwas_variant_count",
    "gwas_locus_lead_snp",
    "gwas_locus_lead_p_lrt",
    "gwas_locus_lead_pip",
    "gwas_locus_lead_pip_rank",
    "gwas_locus_lead_credible_sets",
    "gwas_locus_lead_in_credible_set",
    "overlap_snp_equals_lead",
    "gwas_locus_top_pip_snp",
    "gwas_locus_top_pip_alleles",
    "gwas_locus_top_pip_p_lrt",
    "gwas_locus_top_pip",
    "gwas_locus_top_pip_rank",
    "gwas_locus_top_pip_credible_sets",
    "gwas_locus_top_pip_in_credible_set",
    "overlap_snp_equals_top_pip",
    "animal_id",
    "sample_tissue",
    "canonical_pair_id",
    "allelome_endpoint_gene_ids",
    "shared_snp_equals_allelome_endpoint",
    "shared_snp_min_distance_to_allelome_endpoint_bp",
    "partner_gene_id",
    "partner_gene_symbol",
    "partner_gene_strand",
    "partner_gene_interval",
    "partner_biotype",
    "new_gene_pcgene_overlap",
    "new_gene_pcgene_overlap_bp",
    "overlap_percent_of_new_gene",
    "overlap_percent_of_pcgene",
    "mechanism",
    "linkage_score",
    "interval_gap_bp",
    "candidate_orientation",
    "direction_interpretation",
]

DISPLAY_COLUMNS = [
    "eqtl_finemap_variant",
    "eqtl_ppa",
    "evidence_tier",
    "gwas_trait",
    "gwas_variant_display",
    "gwas_overlap_snp_p_lrt",
    "gwas_overlap_snp_pip",
    "gwas_pip_rank_display",
    "gwas_credible_set_pip_rank_display",
    "gwas_credible_sets",
    "gwas_credible_set_min_p_snp_display",
    "gwas_credible_set_min_p_lrt",
    "allelome_endpoint_gene_ids",
    "partner_gene_symbol",
    "mechanism",
    "linkage_score",
    "sample_tissue",
    "evidence_stratum",
    "new_gene_id",
    "new_gene_interval",
    "new_gene_class_code",
    "new_gene_feelnc_biotype",
    "new_gene_coding_class",
    "partner_gene_interval",
    "new_gene_pcgene_overlap",
    "new_gene_pcgene_overlap_bp",
    "overlap_percent_of_new_gene",
    "overlap_percent_of_pcgene",
]

DISPLAY_LABELS = {
    "eqtl_finemap_variant": "eQTL Fine-map SNP",
    "eqtl_ppa": "eQTL PPA",
    "evidence_tier": "Evidence Tier",
    "gwas_trait": "GWAS Trait",
    "gwas_variant_display": "Matched GWAS SNP (ID | chr:pos:alleles)",
    "gwas_overlap_snp_p_lrt": "Matched SNP P LRT",
    "gwas_overlap_snp_pip": "Matched SNP PIP",
    "gwas_pip_rank_display": "Locus PIP Rank (rank/total)",
    "gwas_credible_set_pip_rank_display": (
        "Credible-set PIP Rank (set:rank/size)"
    ),
    "gwas_credible_sets": "Credible Set",
    "gwas_credible_set_min_p_snp_display": (
        "Credible-set Min-p SNP (ID | chr:pos:alleles)"
    ),
    "gwas_credible_set_min_p_lrt": "Credible-set Min P LRT",
    "allelome_endpoint_gene_ids": "Allelome Gene Endpoint",
    "partner_gene_symbol": "Paired Protein-Coding Gene",
    "mechanism": "Allelome Pattern",
    "linkage_score": "Allelome Score",
    "sample_tissue": "Sample Tissue",
    "evidence_stratum": (
        "Allelome Review Stratum"
    ),
    "new_gene_id": "Matched Manual New Gene",
    "new_gene_interval": "New Gene Interval",
    "new_gene_class_code": "New Gene Class Code",
    "new_gene_feelnc_biotype": "New Gene FEELnc Biotype",
    "new_gene_coding_class": "New Gene Coding Class",
    "partner_gene_interval": "Protein-Coding Gene Interval",
    "new_gene_pcgene_overlap": "New Gene-pcGene Overlap",
    "new_gene_pcgene_overlap_bp": "Overlap (bp)",
    "overlap_percent_of_new_gene": "Overlap / New Gene (%)",
    "overlap_percent_of_pcgene": "Overlap / pcGene (%)",
}

DISPLAY_GROUPS = [
    ("eQTL / integration evidence", 3, "DDE9DF"),
    ("GWAS fine-mapping evidence", 9, "E8F0EA"),
    ("Original Allelome association", 6, "EEF2F7"),
    ("Manual new-gene evidence", 5, "F7EFE7"),
    ("New-gene / pcGene spatial relationship", 5, "F1F3F5"),
]

DISPLAY_COLUMN_SPECS = [
    ("eqtl_finemap_variant", 175, 20),
    ("eqtl_ppa", 105, 13),
    ("evidence_tier", 175, 21),
    ("gwas_trait", 170, 20),
    ("gwas_variant_display", 285, 44),
    ("gwas_overlap_snp_p_lrt", 110, 14),
    ("gwas_overlap_snp_pip", 110, 14),
    ("gwas_pip_rank_display", 100, 14),
    ("gwas_credible_set_pip_rank_display", 135, 18),
    ("gwas_credible_sets", 80, 10),
    ("gwas_credible_set_min_p_snp_display", 285, 44),
    ("gwas_credible_set_min_p_lrt", 110, 14),
    ("allelome_endpoint_gene_ids", 190, 22),
    ("partner_gene_symbol", 155, 19),
    ("mechanism", 110, 15),
    ("linkage_score", 95, 12),
    ("sample_tissue", 185, 22),
    ("evidence_stratum", 185, 22),
    ("new_gene_id", 205, 24),
    ("new_gene_interval", 180, 21),
    ("new_gene_class_code", 115, 15),
    ("new_gene_feelnc_biotype", 145, 18),
    ("new_gene_coding_class", 155, 19),
    ("partner_gene_interval", 190, 22),
    ("new_gene_pcgene_overlap", 140, 17),
    ("new_gene_pcgene_overlap_bp", 100, 12),
    ("overlap_percent_of_new_gene", 125, 15),
    ("overlap_percent_of_pcgene", 125, 15),
]


NUMERIC_DISPLAY_COLUMNS = {
    "eqtl_ppa",
    "gwas_overlap_snp_p_lrt",
    "gwas_overlap_snp_pip",
    "gwas_pip_rank_display",
    "gwas_credible_set_min_p_lrt",
    "linkage_score",
    "shared_snp_min_distance_to_allelome_endpoint_bp",
    "new_gene_pcgene_overlap_bp",
    "overlap_percent_of_new_gene",
    "overlap_percent_of_pcgene",
}


def log(message: str) -> None:
    timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    print(f"[{timestamp} ] {message}", flush=True)


def require_file(path: Path, label: str) -> None:
    if not path.is_file() or path.stat().st_size == 0:
        raise FileNotFoundError(f"{label} does not exist or is empty: {path}")


def open_text(path: Path):
    if path.suffix == ".gz":
        return gzip.open(path, "rt", encoding="utf-8", newline="")
    return path.open("r", encoding="utf-8", newline="")


def read_tsv(path: Path) -> Tuple[List[str], List[Dict[str, str]]]:
    with open_text(path) as handle:
        reader = csv.DictReader(handle, delimiter="\t")
        fields = list(reader.fieldnames or [])
        return fields, list(reader)


def require_columns(
    path: Path, observed: Sequence[str], required: Iterable[str]
) -> None:
    missing = sorted(set(required) - set(observed))
    if missing:
        raise ValueError(f"{path} missing fields: {','.join(missing)}")


def read_metrics(path: Path) -> Dict[str, str]:
    fields, rows = read_tsv(path)
    require_columns(path, fields, {"metric", "value"})
    metrics: Dict[str, str] = {}
    for row in rows:
        key = row["metric"].strip()
        if not key or key in metrics:
            raise ValueError(f"{path} contains empty or duplicate metric: {key!r}")
        metrics[key] = row["value"].strip()
    return metrics


def source_manifest_input_path(path: Path, label: str) -> Path:

    fields, rows = read_tsv(path)
    require_columns(path, fields, {"kind", "label", "path"})
    matches = [
        Path(row["path"].strip())
        for row in rows
        if row["kind"].strip() == "input" and row["label"].strip() == label
    ]
    if len(matches) != 1:
        raise ValueError(
            f"{path} in input/{label} must be unique, observed {len(matches)} records"
        )
    resolved = matches[0]
    if not resolved.is_absolute():
        raise ValueError(f"manifest in {label} is not an absolute path: {resolved}")
    require_file(resolved, f"manifest input {label}")
    return resolved


def format_gwas_coordinate(row: Mapping[str, str]) -> str:

    chrom = row.get("chrom", "").strip()
    pos = row.get("pos", "").strip()
    allele1 = row.get("ld_a1", "").strip().upper()
    allele2 = row.get("ld_a2", "").strip().upper()
    if not chrom or not pos or not allele1 or not allele2:
        return "NA"
    chrom = f"chr{chrom.removeprefix('chr')}"
    alleles = sorted((allele1, allele2))
    return f"{chrom}:{pos}:{alleles[0]}/{alleles[1]}"


def read_gwas_credible_members(
    manifest_path: Path,
) -> Tuple[List[Dict[str, str]], List[Dict[str, str]]]:

    fields, rows = read_tsv(manifest_path)
    required_manifest = {
        "locus_id",
        "trait",
        "analysis_scope",
        "branch",
        "status",
        "model_status",
        "credible_set_count",
        "credible_set_members",
    }
    require_columns(manifest_path, fields, required_manifest)
    main_rows = [row for row in rows if row["branch"].strip() == GWAS_BRANCH]
    if not main_rows:
        raise ValueError(f"{manifest_path} contains no main branch locus")
    members: List[Dict[str, str]] = []
    for manifest_row in main_rows:
        if (
            manifest_row["status"].strip() != "PASS"
            or manifest_row["model_status"].strip() != "CONVERGED"
        ):
            raise ValueError(
                "GWAS main manifest contains non-PASS/CONVERGED locus: "
                f"{manifest_row['locus_id']}"
            )
        member_path = Path(manifest_row["credible_set_members"].strip())
        if not member_path.is_absolute():
            raise ValueError(f"credible set member table is not an absolute path: {member_path}")
        require_file(member_path, "GWAS credible set member table")
        member_fields, member_rows = read_tsv(member_path)
        require_columns(member_path, member_fields, GWAS_MEMBER_REQUIRED)
        for member in member_rows:
            for key in ("locus_id", "trait", "analysis_scope", "branch"):
                if member[key].strip() != manifest_row[key].strip():
                    raise ValueError(
                        f"{member_path} of {key} with GWAS manifest mismatch"
                    )
            p_lrt = to_float(member["p_lrt"], f"{member_path} p_lrt")
            if not math.isfinite(p_lrt):
                raise ValueError(f"{member_path} contains non-finite p_lrt")
            members.append(dict(member))
    return main_rows, members


def summarize_gwas_credible_members(
    manifest_rows: Sequence[Mapping[str, str]],
    members: Sequence[Mapping[str, str]],
) -> Dict[str, str]:

    unique_keys = {
        (
            member["chrom"].strip().removeprefix("chr"),
            member["pos"].strip(),
            member["ld_variant_id"].strip(),
        )
        for member in members
    }
    if not unique_keys:
        raise ValueError("GWAS main credible set no member, cannot generate credible set statistics")

    minimum = min(
        members, key=lambda item: float(item["p_lrt"])
    )
    low_p_unique_keys = {
        (
            member["chrom"].strip().removeprefix("chr"),
            member["pos"].strip(),
            member["ld_variant_id"].strip(),
        )
        for member in members
        if float(member["p_lrt"]) < 1e-5
    }
    minimum_display = (
        f"{minimum['ld_variant_id']} | {format_gwas_coordinate(minimum)}"
    )
    return {
        "gwas_main_locus_count": str(len(manifest_rows)),
        "gwas_main_credible_set_count": str(
            sum(
                to_int(row["credible_set_count"], "GWAS credible_set_count")
                for row in manifest_rows
            )
        ),
        "gwas_main_credible_set_membership_rows": str(len(members)),
        "gwas_main_unique_credible_set_snps": str(len(unique_keys)),
        "gwas_main_duplicate_membership_rows": str(
            len(members) - len(unique_keys)
        ),
        "gwas_main_p_lrt_lt_1e-5_unique_snps": str(
            len(low_p_unique_keys)
        ),
        "gwas_main_min_p_lrt": minimum["p_lrt"],
        "gwas_main_min_p_lrt_snp": minimum["ld_variant_id"],
        "gwas_main_min_p_lrt_snp_coordinate": format_gwas_coordinate(minimum),
        "gwas_main_min_p_lrt_snp_display": minimum_display,
    }


def enrich_focus_gwas_layout(
    rows: Sequence[Mapping[str, str]],
    members: Sequence[Mapping[str, str]],
) -> List[Dict[str, str]]:

    member_index: Dict[Tuple[str, str], List[Mapping[str, str]]] = {}
    for member in members:
        key = (member["locus_id"].strip(), member["cs_id"].strip())
        member_index.setdefault(key, []).append(member)

    enriched_rows: List[Dict[str, str]] = []
    for row in rows:
        set_ids = [
            item.strip()
            for item in row.get("gwas_credible_sets", "")
            .replace(",", ";")
            .split(";")
            if item.strip()
        ]
        candidates: List[Mapping[str, str]] = []
        credible_set_rank_displays: List[str] = []
        matched_variant_id = row["gwas_overlap_snp"].split(" ", 1)[0]
        for set_id in set_ids:
            set_members = member_index.get(
                (row["gwas_locus_id"], set_id), []
            )
            if not set_members:
                raise ValueError(
                    "Tier A record no match its GWAS credible set member: "
                    f"locus={row['gwas_locus_id']}; set={set_id}"
                )
            set_matched = [
                member
                for member in set_members
                if member["ld_variant_id"] == matched_variant_id
            ]
            if len(set_matched) != 1:
                raise ValueError(
                    "shared eQTL-GWAS SNP in marker credible set is not in unique member: "
                    f"locus={row['gwas_locus_id']}; set={set_id}; "
                    f"SNP={matched_variant_id}"
                )
            matched_pip = float(set_matched[0]["pip"])

            set_pip_rank = 1 + sum(
                float(member["pip"]) > matched_pip
                for member in set_members
            )
            credible_set_rank_displays.append(
                f"{set_id}:{set_pip_rank}/{len(set_members)}"
            )
            candidates.extend(set_members)
        if not candidates:
            raise ValueError(
                "Tier A record no match its GWAS credible set member: "
                f"locus={row['gwas_locus_id']}; sets={set_ids}"
            )
        matched = [
            member
            for member in candidates
            if member["ld_variant_id"] == matched_variant_id
        ]
        if not matched:
            raise ValueError(
                "shared eQTL-GWAS SNP not in its marker of credible set in: "
                f"{row['gwas_overlap_snp']}"
            )
        expected_coordinate = row["gwas_overlap_snp"].split(" | ", 1)[-1]
        if not any(
            format_gwas_coordinate(member) == expected_coordinate
            and member["p_lrt"] == row["gwas_overlap_snp_p_lrt"]
            and member["pip"] == row["gwas_overlap_snp_pip"]
            for member in matched
        ):
            raise ValueError(
                "shared SNP with original credible set member coordinates, p_lrt or PIP mismatch: "
                f"{row['gwas_overlap_snp']}"
            )
        pip_rank = to_int(
            row["gwas_overlap_snp_pip_rank"], "shared SNP PIP rank"
        )
        variant_count = to_int(
            row["gwas_variant_count"], "GWAS locus variant count"
        )
        if pip_rank < 1 or variant_count < 1 or pip_rank > variant_count:
            raise ValueError(
                f"PIP rank/total invalid: {pip_rank}/{variant_count}"
            )
        minimum = min(candidates, key=lambda item: float(item["p_lrt"]))
        enriched = dict(row)
        enriched["gwas_variant_display"] = row["gwas_overlap_snp"]
        enriched["gwas_pip_rank_display"] = (
            f"{pip_rank}/{variant_count}"
        )
        enriched["gwas_credible_set_pip_rank_display"] = (
            "; ".join(credible_set_rank_displays)
        )
        enriched["gwas_credible_set_min_p_snp_display"] = (
            f"{minimum['ld_variant_id']} | {format_gwas_coordinate(minimum)}"
        )
        enriched["gwas_credible_set_min_p_lrt"] = minimum["p_lrt"]
        enriched_rows.append(enriched)
    return enriched_rows


def to_int(value: str, label: str) -> int:
    try:
        return int(float(value))
    except (TypeError, ValueError) as exc:
        raise ValueError(f"{label} is not an integer: {value!r}") from exc


def to_float(value: str, label: str) -> float:
    try:
        return float(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"{label} is not numeric: {value!r}") from exc


def shared_finemap_snp_text(row: Mapping[str, str]) -> str:
    snp_id = row.get("gwas_overlap_snp_id", "").strip()
    chrom = row.get("gwas_overlap_snp_chrom", "").strip()
    pos = row.get("gwas_overlap_snp_pos", "").strip()
    normalized_variant_key = row.get("normalized_variant_key", "").strip()
    if not snp_id:
        return "NA"

    if normalized_variant_key:
        return f"{snp_id} | {normalized_variant_key}"
    return f"{snp_id} | {chrom}:{pos}"


def lead_snp_text(row: Mapping[str, str]) -> str:
    snp_id = row.get("gwas_locus_lead_snp_id", "").strip()
    chrom = row.get("gwas_locus_lead_chrom", "").strip()
    pos = row.get("gwas_locus_lead_pos", "").strip()
    if not snp_id:
        return "NA"
    return f"{snp_id} ({chrom}:{pos})"


def lead_alleles(row: Mapping[str, str]) -> str:
    allele1 = row.get("gwas_locus_lead_allele1", "").strip()
    allele0 = row.get("gwas_locus_lead_allele0", "").strip()
    return f"{allele1}/{allele0}" if allele1 and allele0 else "NA"


def top_pip_snp_text(row: Mapping[str, str]) -> str:
    snp_id = row.get("gwas_locus_top_pip_snp_id", "").strip()
    chrom = row.get("gwas_locus_top_pip_chrom", "").strip()
    pos = row.get("gwas_locus_top_pip_pos", "").strip()
    if not snp_id:
        return "NA"
    return f"{snp_id} ({chrom}:{pos})"


def top_pip_alleles(row: Mapping[str, str]) -> str:
    allele1 = row.get("gwas_locus_top_pip_allele1", "").strip()
    allele0 = row.get("gwas_locus_top_pip_allele0", "").strip()
    return f"{allele1}/{allele0}" if allele1 and allele0 else "NA"


def gene_interval(row: Mapping[str, str]) -> str:
    chrom = row.get("partner_chrom", "").strip()
    start0 = row.get("partner_start", "").strip()
    end0 = row.get("partner_end", "").strip()
    if not chrom or not start0 or not end0:
        return "NA"
    return f"{chrom}:{to_int(start0, 'partner_start') + 1}-{end0}"


def unique_values(*values: str) -> str:
    observed: List[str] = []
    seen = set()
    for value in values:
        for item in str(value).replace(";", "|").split("|"):
            item = item.strip()
            if item and item not in {"NA", ".", "-"} and item not in seen:
                seen.add(item)
                observed.append(item)
    return "|".join(observed) if observed else "NA"


def coding_class(feelnc_biotype: str) -> str:
    values = set(feelnc_biotype.split("|"))
    if values == {"lncRNA"}:
        return "noncoding_lncRNA"
    if values == {"mRNA"}:
        return "coding_mRNA"
    if values == {"lncRNA", "mRNA"}:
        return "mixed_lncRNA_mRNA"
    return "undetermined"


def candidate_category(novelty: str) -> str:
    mapping = {
        "intergenic_novel_transcript": "intergenic_new_gene_candidate",
        "antisense_novel_transcript": "antisense_candidate",
        "known_gene_novel_isoform": "known_gene_novel_isoform",
        "retained_intron_or_fragment": "retained_intron_or_fragment",
    }
    return unique_values(
        *(mapping.get(item, item) for item in novelty.split("|"))
    )


def build_five_only_gene_index(
    path: Path,
    fields: Sequence[str],
    rows: Sequence[Mapping[str, str]],
) -> Dict[str, Dict[str, str]]:

    require_columns(path, fields, FIVE_ONLY_TRANSCRIPT_MAP_REQUIRED)
    transcript_ids = [row["final_transcript_id"].strip() for row in rows]
    if len(transcript_ids) != len(set(transcript_ids)):
        raise ValueError("five-dog-specific transcript map of final_transcript_id is not unique key")

    grouped: Dict[str, List[Mapping[str, str]]] = {}
    for row in rows:
        gene_id = row["final_gene_id"].strip()
        if not re.fullmatch(r"DOG5_NEWGENE_[0-9]+", gene_id):
            raise ValueError(f"five-dog-specific gene ID incorrect format: {gene_id!r}")
        grouped.setdefault(gene_id, []).append(row)

    gene_index: Dict[str, Dict[str, str]] = {}
    invariant_fields = (
        "chr",
        "strand",
        "gene_start_0based",
        "gene_end_0based",
        "source_dataset",
        "selection_stage",
    )
    for gene_id, gene_rows in grouped.items():
        invariant_values: Dict[str, str] = {}
        for field in invariant_fields:
            values = {row[field].strip() for row in gene_rows}
            if len(values) != 1:
                raise ValueError(
                    f"{gene_id} of gene level field in transcript between mismatch: {field}"
                )
            invariant_values[field] = next(iter(values))
        if invariant_values["source_dataset"] != FIVE_ONLY_SOURCE_DATASET:
            raise ValueError(f"{gene_id} is not five-dogs source")
        if invariant_values["selection_stage"] != FIVE_ONLY_SELECTION_STAGE:
            raise ValueError(f"{gene_id} is not five dogs Step 2.1 filter result")

        start0 = to_int(
            invariant_values["gene_start_0based"], "gene_start_0based"
        )
        end0 = to_int(
            invariant_values["gene_end_0based"], "gene_end_0based"
        )
        if start0 < 0 or end0 <= start0:
            raise ValueError(f"{gene_id} of gene span invalid")
        gene_index[gene_id] = {
            "final_gene_id": gene_id,
            "chr": invariant_values["chr"],
            "strand": invariant_values["strand"],
            "gene_start_0based": str(start0),
            "gene_end_0based": str(end0),
            "gene_span_bp": str(end0 - start0),
            "source_dataset": invariant_values["source_dataset"],
            "selection_stage": invariant_values["selection_stage"],
            "class_code_list": unique_values(
                *(row["class_code"] for row in gene_rows)
            ),
            "novelty_category_list": unique_values(
                *(row["novelty_category"] for row in gene_rows)
            ),
            "FEELnc_biotype_list": unique_values(
                *(row["FEELnc_biotype"] for row in gene_rows)
            ),
        }
    return gene_index


def five_only_annotation(row: Mapping[str, str]) -> Dict[str, str]:
    start0 = to_int(row["gene_start_0based"], "gene_start_0based")
    end0 = to_int(row["gene_end_0based"], "gene_end_0based")
    span = to_int(row["gene_span_bp"], "gene_span_bp")
    if start0 < 0 or end0 <= start0 or end0 - start0 != span:
        raise ValueError(f"{row['final_gene_id']} novel gene coordinates or spans are inconsistent")
    class_code = row["class_code_list"]
    novelty = row["novelty_category_list"]
    feelnc = row["FEELnc_biotype_list"]
    return {
        "new_gene_chrom": row["chr"].strip(),
        "new_gene_start0": str(start0),
        "new_gene_end0": str(end0),
        "new_gene_span_bp": str(span),
        "new_gene_interval": f"{row['chr']}:{start0 + 1}-{end0}",
        "new_gene_strand": row["strand"].strip(),
        "new_gene_class_code": class_code,
        "new_gene_novelty_category": novelty,
        "new_gene_candidate_category": candidate_category(novelty),
        "new_gene_feelnc_biotype": feelnc,
        "new_gene_coding_class": coding_class(feelnc),
        "new_gene_dataset_support": row["source_dataset"].strip(),

        "new_gene_merge_status": FIVE_ONLY_MERGE_STATUS,
    }


def overlap_fields(
    annotation: Mapping[str, str],
    link: Mapping[str, str],
) -> Dict[str, str]:
    new_start0 = to_int(annotation["new_gene_start0"], "new_gene_start0")
    new_end0 = to_int(annotation["new_gene_end0"], "new_gene_end0")
    new_span = to_int(annotation["new_gene_span_bp"], "new_gene_span_bp")
    partner_start0 = to_int(link["partner_start"], "partner_start")
    partner_end0 = to_int(link["partner_end"], "partner_end")
    if partner_start0 < 0 or partner_end0 <= partner_start0:
        raise ValueError("pcGene of 0-based half-open coordinates invalid")


    partner_span = partner_end0 - partner_start0
    overlap_bp = 0
    if annotation["new_gene_chrom"] == link["partner_chrom"]:
        overlap_bp = max(
            0,
            min(new_end0, partner_end0)
            - max(new_start0, partner_start0),
        )
    return {
        "new_gene_pcgene_overlap": "yes" if overlap_bp else "no",
        "new_gene_pcgene_overlap_bp": str(overlap_bp),
        "overlap_percent_of_new_gene": (
            f"{100.0 * overlap_bp / new_span:.4f}"
        ),
        "overlap_percent_of_pcgene": (
            f"{100.0 * overlap_bp / partner_span:.4f}"
        ),
    }


def validate_manifest(
    path: Path,
    expected_outputs: Mapping[str, Path],
) -> None:
    fields, rows = read_tsv(path)
    require_columns(path, fields, {"kind", "label", "path", "bytes", "purpose"})
    observed = {
        row["label"].strip(): Path(row["path"].strip())
        for row in rows
        if row["kind"].strip() == "output"
    }
    for label, expected_path in expected_outputs.items():
        if observed.get(label) != expected_path:
            raise ValueError(
                f"manifest in {label} path does not match: "
                f"{observed.get(label)} != {expected_path}"
            )


def validate_inputs(
    candidate_path: Path,
    candidate_fields: Sequence[str],
    candidates: Sequence[Mapping[str, str]],
    pair_path: Path,
    pair_fields: Sequence[str],
    pairs: Sequence[Mapping[str, str]],
    link_path: Path,
    link_fields: Sequence[str],
    links: Sequence[Mapping[str, str]],
    five_only_gene_index: Mapping[str, Mapping[str, str]],
    metrics: Mapping[str, str],
) -> None:
    require_columns(candidate_path, candidate_fields, CANDIDATE_REQUIRED)
    require_columns(pair_path, pair_fields, PAIR_REQUIRED)
    require_columns(link_path, link_fields, LINK_REQUIRED)

    if metrics.get("overall_status") != "PASS":
        raise ValueError("upstream Step 8 summary is not PASS")

    expected_counts = {
        "step3_new_gene_candidate_rows": len(candidates),
        "integrated_newgene_pcGene_link_rows": len(links),
        "candidate_pcGene_mechanism_pair_rows": len(pairs),
    }
    for metric, observed_count in expected_counts.items():
        expected_count = to_int(metrics.get(metric, ""), metric)
        if observed_count != expected_count:
            raise ValueError(
                f"{metric} not reconciled: {observed_count} != {expected_count}"
            )

    tier_metric = {
        "A_BOTH_95PCT_CREDIBLE_PLUS_ALLELOME": (
            "tier_A_both_credible_plus_allelome_rows"
        ),
        "B_GWAS_NONCREDIBLE_PLUS_ALLELOME": (
            "tier_B_noncredible_plus_allelome_rows"
        ),
        "C_BOTH_95PCT_CREDIBLE_NO_ALLELOME": (
            "tier_C_both_credible_no_allelome_rows"
        ),
        "D_GWAS_NONCREDIBLE_NO_ALLELOME": (
            "tier_D_noncredible_no_allelome_rows"
        ),
    }
    observed_tiers = Counter(row["evidence_tier"].strip() for row in candidates)
    for tier, metric in tier_metric.items():
        expected_count = to_int(metrics.get(metric, ""), metric)
        if observed_tiers[tier] != expected_count:
            raise ValueError(
                f"{metric} not reconciled: {observed_tiers[tier]} != {expected_count}"
            )

    count_checks = {
        "step3_new_gene_candidate_unique_genes": len(
            {row["new_gene_id"].strip() for row in candidates}
        ),
        "integrated_allelome_link_rows": sum(
            to_int(row["allelome_link_count"], "allelome_link_count")
            for row in candidates
        ),
        "integrated_newgene_pcGene_link_rows": sum(
            to_int(row["pcgene_link_count"], "pcgene_link_count")
            for row in candidates
        ),

        "candidate_same_block_multilayer_rows": sum(
            row["allelome_block_match"].strip() == "yes"
            for row in candidates
        ),
        "shared_snp_equals_gwas_top_pip_rows": sum(
            row["overlap_snp_equals_top_pip"].strip() == "yes"
            for row in candidates
        ),
    }
    for metric, observed_count in count_checks.items():
        expected_count = to_int(metrics.get(metric, ""), metric)
        if observed_count != expected_count:
            raise ValueError(
                f"{metric} not reconciled: {observed_count} != {expected_count}"
            )

    candidate_ids = [row["candidate_id"].strip() for row in candidates]
    if len(candidate_ids) != len(set(candidate_ids)):
        raise ValueError("candidate table candidate_id is not unique key")
    candidate_by_id = {row["candidate_id"].strip(): row for row in candidates}
    missing_new_genes = sorted(
        {row["new_gene_id"].strip() for row in candidates}
        - set(five_only_gene_index)
    )
    if missing_new_genes:
        raise ValueError(
            "candidate new_gene_id no match five-dog-specific transcript map: "
            + ",".join(missing_new_genes)
        )

    for index, row in enumerate(candidates, start=2):
        tier = row["evidence_tier"].strip()
        gene_id = row["new_gene_id"].strip()
        if not re.fullmatch(r"DOG5_NEWGENE_[0-9]+", gene_id):
            raise ValueError(f"candidate table row {index} is not five-dog-specific new_gene_id")
        if row["eqtl_expression_gene_id"].strip() != gene_id:
            raise ValueError(f"candidate table row {index} expression_gene_id mismatch")
        if row["eqtl_feature_id"].strip() != gene_id:
            raise ValueError(f"candidate table row {index} feature_id mismatch")
        if row["eqtl_feature_primary_id"].strip() != gene_id:
            raise ValueError(f"candidate table row {index} feature_primary_id mismatch")
        if row["eqtl_feature_origin"].strip() != "dog5_manual_new_gene":
            raise ValueError(f"candidate table row {index} is not five-dog-specific feature_origin")
        if row["eqtl_expression_source"].strip() != (
            "manual_new_gene_log2_sizefactor_norm_counts"
        ):
            raise ValueError(f"candidate table row {index} expression_source mismatch")
        if row["eqtl_expression_matrix_scope"].strip() != (
            "reference_annotation_plus_dog5_manual_new_gene"
        ):
            raise ValueError(f"candidate table row {index} has an unexpected expression_matrix_scope")
        if tier not in TIER_LABELS:
            raise ValueError(f"candidate table row {index} evidence_tier unknown: {tier}")
        if row["allele_match_status"].strip() != "allele_set_match":
            raise ValueError(f"candidate table row {index} allele sets do not match")
        if not re.fullmatch(r"chr(?:[1-9]|[12][0-9]|3[0-8])", row["eqtl_phase_scope"].strip()):
            raise ValueError(f"candidate table row {index} is not autosome all chromosome phase")
        if to_float(row["eqtl_ppa"], f"candidate table row {index} eqtl_ppa") < 0.1:
            raise ValueError(f"candidate table row {index} eqtl_ppa below 0.1")
        if row["gwas_branch"].strip() != "main":
            raise ValueError(f"candidate table row {index} is not GWAS main branch")
        expected_layer = (
            "priority_both" if tier.startswith(("A_", "C_")) else "noncredible"
        )
        if row["result_layer"].strip() != expected_layer:
            raise ValueError(f"candidate table row {index} result_layer with evidence layer mismatch")
        expected_allelome = "yes" if tier.startswith(("A_", "B_")) else "no"
        if row["allelome_block_match"].strip() != expected_allelome:
            raise ValueError(f"candidate table row {index} Allelome status and evidence layer mismatch")
        expected_main = "yes" if tier.startswith("A_") else "no"
        if row["main_result"].strip() != expected_main:
            raise ValueError(f"candidate table row {index} main_result with evidence layer mismatch")

    endpoint_pattern = re.compile(r"DOG5_NEWGENE_[0-9]+")
    for index, row in enumerate(pairs, start=2):
        candidate_id = row["candidate_id"].strip()
        if candidate_id not in candidate_by_id:
            raise ValueError(f"pair table row {index} candidate_id no match candidate table")
        if row["evidence_tier"].strip() != candidate_by_id[candidate_id]["evidence_tier"].strip():
            raise ValueError(f"pair table row {index} evidence_tier with candidate table mismatch")
        endpoints = [
            value for value in row["allelome_endpoints"].split("|") if value
        ]
        if len(endpoints) != to_int(
            row["allelome_endpoint_count"],
            f"pair table row {index} allelome_endpoint_count",
        ):
            raise ValueError(f"pair table row {index} endpoint count not reconciled")
        if any(not endpoint_pattern.fullmatch(value) for value in endpoints):
            raise ValueError(f"Row {index} of the pair table contains a nonstandard new-gene endpoint ID")

    for index, row in enumerate(links, start=2):
        candidate_id = row["candidate_id"].strip()
        if candidate_id not in candidate_by_id:
            raise ValueError(f"expand table row {index} candidate_id no match candidate table")
        if row["partner_endpoint_type"].strip() != "known_protein_coding":
            raise ValueError(f"expand table row {index} partner is not pcGene")
        if row["eqtl_phase_scope"].strip() != candidate_by_id[candidate_id]["eqtl_phase_scope"].strip():
            raise ValueError(f"expand table row {index} phase scope with candidate table mismatch")


def build_block_rows(
    candidates: Sequence[Mapping[str, str]],
    five_only_gene_index: Mapping[str, Mapping[str, str]],
) -> List[Dict[str, str]]:
    rows: List[Dict[str, str]] = []
    for source in candidates:
        annotation = five_only_annotation(
            five_only_gene_index[source["new_gene_id"]]
        )
        row = {
            **{field: source[field] for field in (
                "allelome_evidence_strata", "allelome_multigene_link_rows", "allelome_overlap_link_rows"
            )},
            "evidence_tier": source["evidence_tier"],
            "tier_meaning": TIER_LABELS[source["evidence_tier"]],
            "main_result": source["main_result"],
            "candidate_id": source["candidate_id"],
            "new_gene_id": source["new_gene_id"],
            **annotation,
            "eqtl_block_id": source["eqtl_block_id"],
            "eqtl_tissue": source["eqtl_tissue"],
            "eqtl_finemap_variant": source["normalized_variant_key"],
            "eqtl_ppa": source["eqtl_ppa"],
            "gwas_trait": source["gwas_trait"],
            "gwas_locus_id": source["gwas_locus_id"],
            "gwas_overlap_snp": shared_finemap_snp_text(source),
            "gwas_overlap_snp_p_lrt": source["gwas_overlap_snp_p_lrt"],
            "gwas_overlap_snp_pip": source["gwas_overlap_snp_pip"],
            "gwas_overlap_snp_pip_rank": source["gwas_overlap_snp_pip_rank"],
            "gwas_overlap_snp_in_credible_set": (
                source["gwas_overlap_snp_in_credible_set"]
            ),
            "gwas_credible_sets": source["gwas_credible_sets"],
            "gwas_variant_count": source["gwas_variant_count"],
            "gwas_locus_lead_snp": lead_snp_text(source),
            "gwas_locus_lead_alleles": lead_alleles(source),
            "gwas_locus_lead_p_lrt": source["gwas_locus_lead_p_lrt"],
            "gwas_locus_lead_pip": source["gwas_locus_lead_pip"],
            "gwas_locus_lead_pip_rank": source["gwas_locus_lead_pip_rank"],
            "gwas_locus_lead_credible_sets": (
                source["gwas_locus_lead_credible_sets"]
            ),
            "gwas_locus_lead_in_credible_set": (
                source["gwas_locus_lead_in_credible_set"]
            ),
            "overlap_snp_equals_lead": source["overlap_snp_equals_lead"],
            "gwas_locus_top_pip_snp": top_pip_snp_text(source),
            "gwas_locus_top_pip_alleles": top_pip_alleles(source),
            "gwas_locus_top_pip_p_lrt": source[
                "gwas_locus_top_pip_p_lrt"
            ],
            "gwas_locus_top_pip": source["gwas_locus_top_pip"],
            "gwas_locus_top_pip_rank": source[
                "gwas_locus_top_pip_rank"
            ],
            "gwas_locus_top_pip_credible_sets": source[
                "gwas_locus_top_pip_credible_sets"
            ],
            "gwas_locus_top_pip_in_credible_set": source[
                "gwas_locus_top_pip_in_credible_set"
            ],
            "overlap_snp_equals_top_pip": source[
                "overlap_snp_equals_top_pip"
            ],
            "allelome_block_match": source["allelome_block_match"],
            "multilayer_integration_scope": (
                source["multilayer_integration_scope"]
            ),
            "shared_snp_equals_any_allelome_endpoint": (
                source["shared_snp_equals_any_allelome_endpoint"]
            ),
            "shared_snp_min_distance_to_allelome_endpoint_bp": (
                source["shared_snp_min_distance_to_allelome_endpoint_bp"]
            ),
            "supporting_dog_count": source["supporting_dog_count"],
            "supporting_dogs": source["supporting_dogs"],
            "allelome_link_count": source["allelome_link_count"],
            "pcgene_link_count": source["pcgene_link_count"],
            "partner_gene_count": source["partner_gene_count"],
            "interpretation": source["evidence_interpretation"],
        }
        rows.append(row)
    return rows


def build_pair_rows(
    pairs: Sequence[Mapping[str, str]],
    candidate_by_id: Mapping[str, Mapping[str, str]],
    five_only_gene_index: Mapping[str, Mapping[str, str]],
    link_by_pair: Mapping[Tuple[str, str, str], Mapping[str, str]],
) -> List[Dict[str, str]]:
    rows: List[Dict[str, str]] = []
    for source in pairs:
        candidate = candidate_by_id[source["candidate_id"]]
        annotation = five_only_annotation(
            five_only_gene_index[source["new_gene_id"]]
        )
        link_key = (
            source["candidate_id"],
            source["partner_gene_id"],
            source["mechanism"],
        )
        if link_key not in link_by_pair:
            raise ValueError(f"unique pair no match expand evidence: {link_key}")
        representative = link_by_pair[link_key]
        overlap = overlap_fields(annotation, representative)
        row = {
            **{field: source[field] for field in (
                "allelome_evidence_strata", "allelome_multigene_link_rows", "allelome_overlap_link_rows"
            )},
            "evidence_tier": source["evidence_tier"],
            "tier_meaning": TIER_LABELS[source["evidence_tier"]],
            "main_result": source["main_result"],
            "candidate_id": source["candidate_id"],
            "new_gene_id": source["new_gene_id"],
            "new_gene_class_code": annotation["new_gene_class_code"],
            "new_gene_novelty_category": (
                annotation["new_gene_novelty_category"]
            ),
            "new_gene_candidate_category": (
                annotation["new_gene_candidate_category"]
            ),
            "new_gene_feelnc_biotype": (
                annotation["new_gene_feelnc_biotype"]
            ),
            "new_gene_coding_class": annotation["new_gene_coding_class"],
            "new_gene_interval": annotation["new_gene_interval"],
            "new_gene_strand": annotation["new_gene_strand"],
            "eqtl_tissue": source["eqtl_tissue"],
            "eqtl_finemap_variant": source["normalized_variant_key"],
            "eqtl_ppa": candidate["eqtl_ppa"],
            "gwas_trait": source["gwas_trait"],
            "gwas_overlap_snp": shared_finemap_snp_text(source),
            "gwas_overlap_snp_p_lrt": source["gwas_overlap_snp_p_lrt"],
            "gwas_overlap_snp_pip": source["gwas_overlap_snp_pip"],
            "gwas_overlap_snp_pip_rank": source["gwas_overlap_snp_pip_rank"],
            "gwas_overlap_snp_in_credible_set": (
                source["gwas_overlap_snp_in_credible_set"]
            ),
            "gwas_locus_lead_snp": lead_snp_text(source),
            "gwas_locus_lead_p_lrt": source["gwas_locus_lead_p_lrt"],
            "gwas_locus_lead_pip": source["gwas_locus_lead_pip"],
            "gwas_locus_lead_pip_rank": source["gwas_locus_lead_pip_rank"],
            "gwas_locus_lead_credible_sets": (
                source["gwas_locus_lead_credible_sets"]
            ),
            "gwas_locus_lead_in_credible_set": (
                source["gwas_locus_lead_in_credible_set"]
            ),
            "overlap_snp_equals_lead": source["overlap_snp_equals_lead"],
            "gwas_locus_top_pip_snp": top_pip_snp_text(source),
            "gwas_locus_top_pip_alleles": top_pip_alleles(source),
            "gwas_locus_top_pip_p_lrt": source[
                "gwas_locus_top_pip_p_lrt"
            ],
            "gwas_locus_top_pip": source["gwas_locus_top_pip"],
            "gwas_locus_top_pip_rank": source[
                "gwas_locus_top_pip_rank"
            ],
            "gwas_locus_top_pip_credible_sets": source[
                "gwas_locus_top_pip_credible_sets"
            ],
            "gwas_locus_top_pip_in_credible_set": source[
                "gwas_locus_top_pip_in_credible_set"
            ],
            "overlap_snp_equals_top_pip": source[
                "overlap_snp_equals_top_pip"
            ],
            "partner_gene_id": source["partner_gene_id"],
            "partner_gene_symbol": source["partner_gene_symbol"],
            "partner_biotype": representative["partner_biotype"],
            "partner_gene_interval": gene_interval(representative),
            **overlap,
            "mechanism": source["mechanism"],
            "expanded_link_rows": source["expanded_link_rows"],
            "supporting_dog_count": source["supporting_dog_count"],
            "supporting_dogs": source["supporting_dogs"],
            "allelome_endpoint_count": source["allelome_endpoint_count"],
            "allelome_endpoints": source["allelome_endpoints"],
            "shared_snp_equals_any_allelome_endpoint": (
                source["shared_snp_equals_any_allelome_endpoint"]
            ),
            "shared_snp_min_distance_to_allelome_endpoint_bp": (
                source["shared_snp_min_distance_to_allelome_endpoint_bp"]
            ),
            "linkage_score_min": source["linkage_score_min"],
            "linkage_score_max": source["linkage_score_max"],
            "interval_gap_bp_min": source["interval_gap_bp_min"],
            "interval_gap_bp_max": source["interval_gap_bp_max"],
            "candidate_orientation": source["candidate_orientation"],
            "direction_interpretation": source["direction_interpretation"],
        }
        rows.append(row)
    return rows


def build_detail_rows(
    links: Sequence[Mapping[str, str]],
    five_only_gene_index: Mapping[str, Mapping[str, str]],
) -> List[Dict[str, str]]:
    rows: List[Dict[str, str]] = []
    for source in links:
        annotation = five_only_annotation(
            five_only_gene_index[source["new_gene_id"]]
        )
        overlap = overlap_fields(annotation, source)
        rows.append(
            {
                "evidence_tier": source["evidence_tier"],
                "main_result": source["main_result"],
                "candidate_id": source["candidate_id"],
                "new_gene_id": source["new_gene_id"],
                "new_gene_class_code": annotation["new_gene_class_code"],
                "new_gene_novelty_category": (
                    annotation["new_gene_novelty_category"]
                ),
                "new_gene_candidate_category": (
                    annotation["new_gene_candidate_category"]
                ),
                "new_gene_feelnc_biotype": (
                    annotation["new_gene_feelnc_biotype"]
                ),
                "new_gene_coding_class": (
                    annotation["new_gene_coding_class"]
                ),
                "new_gene_interval": annotation["new_gene_interval"],
                "new_gene_strand": annotation["new_gene_strand"],
                "eqtl_block_id": source["eqtl_block_id"],
                "eqtl_tissue": source["eqtl_tissue"],
                "eqtl_finemap_variant": source["normalized_variant_key"],
                "eqtl_ppa": source["eqtl_ppa"],
                "gwas_trait": source["gwas_trait"],
                "gwas_locus_id": source["gwas_locus_id"],
                "gwas_overlap_snp": shared_finemap_snp_text(source),
                "gwas_overlap_snp_p_lrt": (
                    source["gwas_overlap_snp_p_lrt"]
                ),
                "gwas_overlap_snp_pip": source["gwas_overlap_snp_pip"],
                "gwas_overlap_snp_pip_rank": (
                    source["gwas_overlap_snp_pip_rank"]
                ),
                "gwas_overlap_snp_in_credible_set": (
                    source["gwas_overlap_snp_in_credible_set"]
                ),
                "gwas_credible_sets": source["gwas_credible_sets"],
                "gwas_variant_count": source["gwas_variant_count"],
                "gwas_locus_lead_snp": lead_snp_text(source),
                "gwas_locus_lead_p_lrt": (
                    source["gwas_locus_lead_p_lrt"]
                ),
                "gwas_locus_lead_pip": source["gwas_locus_lead_pip"],
                "gwas_locus_lead_pip_rank": (
                    source["gwas_locus_lead_pip_rank"]
                ),
                "gwas_locus_lead_credible_sets": (
                    source["gwas_locus_lead_credible_sets"]
                ),
                "gwas_locus_lead_in_credible_set": (
                    source["gwas_locus_lead_in_credible_set"]
                ),
                "overlap_snp_equals_lead": source["overlap_snp_equals_lead"],
                "gwas_locus_top_pip_snp": top_pip_snp_text(source),
                "gwas_locus_top_pip_alleles": top_pip_alleles(source),
                "gwas_locus_top_pip_p_lrt": source[
                    "gwas_locus_top_pip_p_lrt"
                ],
                "gwas_locus_top_pip": source["gwas_locus_top_pip"],
                "gwas_locus_top_pip_rank": source[
                    "gwas_locus_top_pip_rank"
                ],
                "gwas_locus_top_pip_credible_sets": source[
                    "gwas_locus_top_pip_credible_sets"
                ],
                "gwas_locus_top_pip_in_credible_set": source[
                    "gwas_locus_top_pip_in_credible_set"
                ],
                "overlap_snp_equals_top_pip": source[
                    "overlap_snp_equals_top_pip"
                ],
                **{field: source[field] for field in (
                    "sample_pair_id", "evidence_stratum", "either_endpoint_multigene",
                    "shared_variant_identity_checked", "source_link_file", "source_link_line"
                )},
                "animal_id": source["animal_id"],
                "sample_tissue": source["sample_tissue"],
                "canonical_pair_id": source["canonical_pair_id"],
                "allelome_endpoint_gene_ids": (
                    source["new_gene_endpoint_manual_ids"]
                ),
                "shared_snp_equals_allelome_endpoint": (
                    source["shared_finemap_snp_equals_allelome_endpoint"]
                ),
                "shared_snp_min_distance_to_allelome_endpoint_bp": (
                    source[
                        "shared_finemap_snp_min_distance_to_allelome_endpoint_bp"
                    ]
                ),
                "partner_gene_id": source["partner_gene_id"],
                "partner_gene_symbol": source["partner_gene_symbol"],
                "partner_gene_strand": source["partner_gene_strand"],
                "partner_gene_interval": gene_interval(source),
                "partner_biotype": source["partner_biotype"],
                **overlap,
                "mechanism": source["mechanism"],
                "linkage_score": source["linkage_score"],
                "interval_gap_bp": source["interval_gap_bp"],
                "candidate_orientation": source["candidate_orientation"],
                "direction_interpretation": (
                    "new gene-pcGene specifies only the display orientation of endpoint categories; "
                    "Allelome enhancing or repressing relationships do not establish causal regulatory direction."
                ),
            }
        )
    return rows


def write_tsv(
    path: Path,
    fields: Sequence[str],
    rows: Sequence[Mapping[str, str]],
) -> None:
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(
            handle, fieldnames=list(fields), delimiter="\t", lineterminator="\n"
        )
        writer.writeheader()
        writer.writerows(
            {field: row.get(field, "") for field in fields} for row in rows
        )


def append_sheet(
    workbook: Workbook,
    title: str,
    fields: Sequence[str],
    rows: Sequence[Mapping[str, str]],
) -> None:
    worksheet = workbook.create_sheet(title)
    worksheet.append(list(fields))
    for row in rows:
        worksheet.append([row.get(field, "") for field in fields])
    header_fill = PatternFill("solid", fgColor="DDE9DF")
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = Font(bold=True)
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    for column_index, field in enumerate(fields, start=1):
        values = [str(row.get(field, "")) for row in rows[:200]]
        width = min(44, max(12, len(field) + 2, *(len(v) + 2 for v in values)))
        worksheet.column_dimensions[get_column_letter(column_index)].width = width
    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def append_focus_sheet(
    workbook: Workbook,
    rows: Sequence[Mapping[str, str]],
) -> None:
    worksheet = workbook.create_sheet("focus_results")
    start_column = 1
    for label, column_count, color in DISPLAY_GROUPS:
        end_column = start_column + column_count - 1
        worksheet.merge_cells(
            start_row=1,
            start_column=start_column,
            end_row=1,
            end_column=end_column,
        )
        cell = worksheet.cell(row=1, column=start_column, value=label)
        cell.font = Font(bold=True, color="222222")
        cell.fill = PatternFill("solid", fgColor=color)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        for column in range(start_column + 1, end_column + 1):
            worksheet.cell(row=1, column=column).fill = PatternFill(
                "solid", fgColor=color
            )
        start_column = end_column + 1

    worksheet.append([DISPLAY_LABELS[field] for field in DISPLAY_COLUMNS])
    for row in rows:
        worksheet.append([row.get(field, "") for field in DISPLAY_COLUMNS])

    for cell in worksheet[2]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAD3")
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
    worksheet.freeze_panes = "A3"
    worksheet.auto_filter.ref = (
        f"A2:{get_column_letter(len(DISPLAY_COLUMNS))}"
        f"{worksheet.max_row}"
    )
    for column_index, (_field, pixel_width, _wrap_width) in enumerate(
        DISPLAY_COLUMN_SPECS, start=1
    ):
        worksheet.column_dimensions[
            get_column_letter(column_index)
        ].width = max(12, min(35, round(pixel_width / 7)))
    for row in worksheet.iter_rows(min_row=3):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def write_workbook(
    path: Path,
    summary_rows: Sequence[Mapping[str, str]],
    block_rows: Sequence[Mapping[str, str]],
    pair_rows: Sequence[Mapping[str, str]],
    detail_rows: Sequence[Mapping[str, str]],
    focus_detail_rows: Sequence[Mapping[str, str]],
) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    append_sheet(workbook, "overview", ["metric", "value"], summary_rows)
    append_focus_sheet(workbook, focus_detail_rows)
    append_sheet(workbook, "full_results", DETAIL_COLUMNS, detail_rows)
    append_sheet(workbook, "candidate_records", BLOCK_COLUMNS, block_rows)
    append_sheet(workbook, "candidate_pcgene_pairs", PAIR_COLUMNS, pair_rows)
    append_sheet(
        workbook,
        "field_notes",
        ["field_group", "meaning"],
        [
            {
                "field_group": "Shared fine-map SNP",
                "meaning": (
                    "SNP exactly matched between eQTL fine-map and GWAS fine-map results; "
                    "The focus view follows the reference summary layout to display ID, chr:pos:alleles, "
                    "p_lrt, PIP, rank/total variants in the locus, and credible-set ID together."
                ),
            },
            {
                "field_group": "GWAS credible-set minimum-p SNP",
                "meaning": (
                    "From the original credible-set member table referenced by the GWAS main locus manifest, "
                    "select the member with the lowest p_lrt in the credible set containing the matched SNP; "
                    "This is not approximated using the integrated-candidate subset."
                ),
            },
            {
                "field_group": "GWAS locus association lead SNP",
                "meaning": (
                    "gwas_lead_variant_id explicitly recorded in the GWAS main manifest; "
                    "It is the association lead, not necessarily the top-PIP SNP."
                ),
            },
            {
                "field_group": "GWAS locus top-PIP SNP",
                "meaning": (
                    "top_variant_id and max_pip explicitly recorded in the GWAS main manifest; "
                    "Its p_lrt/PIP are reported separately from those of the overlapping SNP and association lead."
                ),
            },
            {
                "field_group": "Allelome",
                "meaning": (
                    "Gene-interval endpoint relationships within the same new-gene/eQTL block; NA exact identity means not evaluable. Review flags are retained; Tier A does not imply absence of assignment issues; "
                    "enhancing/repressing does not establish causal regulatory direction."
                ),
            },
            {
                "field_group": "Manual new gene",
                "meaning": (
                    "The five-dog transcript map aggregates gene span, "
                    "class code, novelty, and FEELnc biotype by final_gene_id; coding class is explicitly derived from "
                    "FEELnc values."
                ),
            },
            {
                "field_group": "New gene / pcGene spatial relationship",
                "meaning": (
                    "Both five-dog gene spans and Allelome known_gene endpoints use "
                    "0-based half-open coordinates to calculate overlap in bp and proportions relative to both intervals directly; "
                    "Human-readable intervals are displayed separately as 1-based inclusive."
                ),
            },
        ],
    )
    workbook.save(path)


def write_html(
    path: Path,
    detail_rows: Sequence[Mapping[str, str]],
    summary: Mapping[str, str],
) -> None:
    group_header = "".join(
        f'<th colspan="{count}" style="background:#{color}">'
        f"{html.escape(label)}</th>"
        for label, count, color in DISPLAY_GROUPS
    )
    column_header = "".join(
        f"<th>{html.escape(DISPLAY_LABELS[field])}</th>"
        for field in DISPLAY_COLUMNS
    )
    body: List[str] = []
    previous_candidate = None
    visual_group_index = -1
    for row in detail_rows:
        candidate_id = row["candidate_id"]
        block_start = candidate_id != previous_candidate
        if block_start:
            visual_group_index += 1
            previous_candidate = candidate_id
        cells = "".join(
            f"<td>{html.escape(str(row.get(field, '') or 'NA'))}</td>"
            for field in DISPLAY_COLUMNS
        )
        classes = []
        if block_start:
            classes.append("block-start")
        if visual_group_index % 2:
            classes.append("alternate-block")
        class_text = f' class="{" ".join(classes)}"' if classes else ""
        body.append(f"<tr{class_text}>{cells}</tr>")
    document = f"""<!doctype html>
<html lang="en">
<head>
<meta charset="utf-8">
<title>BarkBase step 8 Tier A shared credible-set results</title>
<style>
body {{ font-family: Arial, sans-serif; margin: 24px; background: #f7f7f4; color: #222; }}
h1 {{ font-size: 22px; margin-bottom: 8px; }}
.meta {{ max-width: 1250px; color: #555; font-size: 13px; line-height: 1.55; }}
.table-wrap {{ overflow-x: auto; border: 1px solid #ccc; background: #fff; }}
table {{ border-collapse: collapse; min-width: 3600px; width: 100%; font-size: 12px; }}
th, td {{ border: 1px solid #ccc; padding: 6px 8px; vertical-align: top; }}
th {{ background: #e6efe8; position: sticky; top: 29px; z-index: 1; }}
.group th {{ text-align: center; top: 0; z-index: 2; }}
tr.alternate-block {{ background: #fafafa; }}
tr.block-start td {{ border-top: 2px solid #777; }}
</style>
</head>
<body>
<h1>BarkBase: Tier A new gene × Allelome × eQTL/GWAS fine-map</h1>
<p class="meta">Candidate records={html.escape(summary["candidate_rows"])};
Tier A={html.escape(summary["tier_A_candidate_rows"])}；
Within-candidate aggregated new gene-pcGene combinations={html.escape(summary["unique_pair_rows"])};
All expanded Allelome relationships={html.escape(summary["all_allelome_expanded_link_rows"])};
Expanded pcGene relationships among these={html.escape(summary["pcgene_expanded_link_rows"])};
lncRNA new genes={html.escape(summary["lncrna_new_gene_count"])};
Evidence records with interval overlap={html.escape(summary["overlapping_expanded_rows"])};
Unique GWAS main credible-set SNPs={html.escape(summary["gwas_main_unique_credible_set_snps"])};
Unique SNPs among these with p_lrt&lt;1e-5={html.escape(summary["gwas_main_p_lrt_lt_1e-5_unique_snps"])}.
The GWAS section follows the Allelome-GWAS main summary layout, displaying the matched SNP ID/coordinates/alleles,
p_lrt, PIP, within-locus rank/total, within-credible-set rank/size, and the credible-set SNP with the lowest p_lrt;
association lead and top-PIP SNP information is retained in full in full_results and candidate audit tables;
coding classes come from FEELnc results in the five-dog transcript map; mixed types are not forced into coding or noncoding categories.
This view retains only Tier A candidates shared by both 95% credible sets with Allelome in the same block and their pcGene
subset relationships. This remains block-level convergence of multiple layers of candidate evidence, not formal colocalization or proof of causality. Allelome endpoints are gene intervals, so three-layer overlap at the same SNP cannot be established; Review Stratum flags must be interpreted together with Tier A.</p>
<div class="table-wrap"><table><thead>
<tr class="group">{group_header}</tr>
<tr>{column_header}</tr></thead>
<tbody>{''.join(body)}</tbody></table></div>
</body>
</html>
"""
    path.write_text(document, encoding="utf-8")


def wrap_text(value: str, width: int = 20) -> List[str]:
    text = str(value).strip() or "NA"
    return [text[index : index + width] for index in range(0, len(text), width)]


def wrap_svg_value(field: str, value: str, width: int) -> List[str]:

    text = str(value).strip() or "NA"
    if field in NUMERIC_DISPLAY_COLUMNS:
        return [text]
    return wrap_text(text, width)


def write_svg(
    path: Path,
    detail_rows: Sequence[Mapping[str, str]],
    summary: Mapping[str, str],
) -> None:
    column_widths = [
        pixel_width
        for _field, pixel_width, _wrap_width in DISPLAY_COLUMN_SPECS
    ]
    left = 36
    top = 34
    title_height = 86
    group_height = 24
    header_height = 46
    line_height = 11
    width = left * 2 + sum(column_widths)
    wrapped_rows: List[List[List[str]]] = []
    row_heights: List[int] = []
    for row in detail_rows:
        cells = [
            wrap_svg_value(
                field,
                row.get(field, "NA"),
                DISPLAY_COLUMN_SPECS[index][2],
            )
            for index, field in enumerate(DISPLAY_COLUMNS)
        ]
        wrapped_rows.append(cells)
        row_heights.append(max(34, 10 + max(map(len, cells)) * line_height))
    height = (
        top + title_height + group_height + header_height
        + sum(row_heights) + 24
    )
    svg = [
        f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}">',
        '<rect width="100%" height="100%" fill="#ffffff"/>',
        f'<text x="{width / 2:.1f}" y="{top}" font-size="16" '
        'font-family="Arial" text-anchor="middle">BarkBase multilayer '
        "Tier A fine-mapping summary</text>",
        f'<text x="{width / 2:.1f}" y="{top + 22}" font-size="11" '
        'font-family="Arial" text-anchor="middle">'
        f"candidate records={escape(summary['candidate_rows'])}; "
        f"Tier A={escape(summary['tier_A_candidate_rows'])}; "
        f"candidate-specific pcGene pairs={escape(summary['unique_pair_rows'])}; "
        f"all Allelome links={escape(summary['all_allelome_expanded_link_rows'])}; "
        f"pcGene links={escape(summary['pcgene_expanded_link_rows'])}; "
        f"overlap rows={escape(summary['overlapping_expanded_rows'])}</text>",
        f'<text x="{width / 2:.1f}" y="{top + 40}" font-size="11" '
        'font-family="Arial" text-anchor="middle">'
        f"all-main credible-set SNPs="
        f"{escape(summary['gwas_main_unique_credible_set_snps'])}; "
        f"p_lrt&lt;1e-5 SNPs="
        f"{escape(summary['gwas_main_p_lrt_lt_1e-5_unique_snps'])}</text>",
        f'<text x="{width / 2:.1f}" y="{top + 58}" font-size="11" '
        'font-family="Arial" text-anchor="middle">'
        "GWAS display reports locus and credible-set PIP ranks; lead and "
        "top-PIP fields remain in full_results.</text>",
    ]
    x_positions: List[int] = []
    current_x = left
    for column_width in column_widths:
        x_positions.append(current_x)
        current_x += column_width
    header_y = top + title_height
    group_start = 0
    for label, count, color in DISPLAY_GROUPS:
        group_width = sum(column_widths[group_start : group_start + count])
        group_x = x_positions[group_start]
        svg.append(
            f'<rect x="{group_x}" y="{header_y}" width="{group_width}" '
            f'height="{group_height}" fill="#{color}"/>'
        )
        svg.append(
            f'<text x="{group_x + group_width / 2:.1f}" '
            f'y="{header_y + 16}" font-size="10" font-family="Arial" '
            f'font-weight="700" text-anchor="middle">{escape(label)}</text>'
        )
        group_start += count
    column_y = header_y + group_height
    for index, field in enumerate(DISPLAY_COLUMNS):
        x = x_positions[index]
        column_width = column_widths[index]
        for line_index, line in enumerate(
            wrap_text(DISPLAY_LABELS[field], 22)[:3]
        ):
            svg.append(
                f'<text x="{x + column_width / 2:.1f}" '
                f'y="{column_y + 14 + line_index * 11}" font-size="9.5" '
                'font-family="Arial" font-weight="700" text-anchor="middle">'
                f"{escape(line)}</text>"
            )
    current_y = column_y + header_height
    previous_candidate = None
    visual_group_index = -1
    for row_index, (source_row, row_cells, row_height) in enumerate(
        zip(detail_rows, wrapped_rows, row_heights)
    ):
        candidate_id = source_row["candidate_id"]
        if candidate_id != previous_candidate:
            visual_group_index += 1
            previous_candidate = candidate_id
            if row_index > 0:
                svg.append(
                    f'<line x1="{left}" y1="{current_y}" '
                    f'x2="{left + sum(column_widths)}" y2="{current_y}" '
                    'stroke="#777777" stroke-width="1.1"/>'
                )
        if visual_group_index % 2:
            svg.append(
                f'<rect x="{left}" y="{current_y}" '
                f'width="{sum(column_widths)}" height="{row_height}" '
                'fill="#fafafa"/>'
            )
        for column_index, lines in enumerate(row_cells):
            field = DISPLAY_COLUMNS[column_index]
            numeric = field in NUMERIC_DISPLAY_COLUMNS
            if numeric:
                x = x_positions[column_index] + column_widths[column_index] / 2
                text_anchor = "middle"
                font_family = "Menlo,Consolas,monospace"
            else:
                x = x_positions[column_index] + 6
                text_anchor = "start"
                font_family = "Arial"
            start_y = current_y + 14
            for line_index, line in enumerate(lines):
                svg.append(
                    f'<text x="{x}" y="{start_y + line_index * line_height}" '
                    f'font-size="9" font-family="{font_family}" '
                    f'fill="#111111" text-anchor="{text_anchor}">'
                    f"{escape(line)}</text>"
                )
        current_y += row_height
        svg.append(
            f'<line x1="{left}" y1="{current_y}" '
            f'x2="{left + sum(column_widths)}" y2="{current_y}" '
            'stroke="#dddddd" stroke-width="0.7"/>'
        )
    svg.append("</svg>")
    path.write_text("\n".join(svg), encoding="utf-8")


def main() -> int:
    parser = argparse.ArgumentParser(
        description=(
            "Summarize five-dog new-gene-Allelome-eQTL-fine-mapping-GWAS-fine-mapping results."
        )
    )
    parser.add_argument("--candidates", type=Path, default=DEFAULT_CANDIDATES)
    parser.add_argument("--pcgene-links", type=Path, default=DEFAULT_PCG_LINKS)
    parser.add_argument("--candidate-pairs", type=Path, default=DEFAULT_PAIRS)
    parser.add_argument(
        "--source-summary", type=Path, default=DEFAULT_SOURCE_SUMMARY
    )
    parser.add_argument(
        "--source-manifest", type=Path, default=DEFAULT_SOURCE_MANIFEST
    )
    parser.add_argument(
        "--five-only-new-gene-map",
        type=Path,
        default=DEFAULT_FIVE_ONLY_NEW_GENE_MAP,
    )
    parser.add_argument("--output-root", type=Path, default=DEFAULT_OUTPUT_ROOT)
    parser.add_argument(
        "--run-id", default=datetime.now().strftime("%Y%m%d_%H%M%S")
    )
    args = parser.parse_args()

    log("[stage 1/5] check Step 8 main input")
    input_paths = [
        (args.candidates, "All candidate records"),
        (args.pcgene_links, "Individual new gene-pcGene evidence records"),
        (args.candidate_pairs, "Within-candidate new gene-pcGene mechanism combinations"),
        (args.source_summary, "Upstream summary"),
        (args.source_manifest, "Upstream manifest"),
        (args.five_only_new_gene_map, "Five-dog new-gene transcript map"),
    ]
    for path, label in input_paths:
        if not path.is_absolute():
            raise ValueError(f"{label} is not an absolute path: {path}")
        require_file(path, label)
        log(f"{label}: {path}")
    if not args.output_root.is_absolute():
        raise ValueError(f"output_root is not an absolute path: {args.output_root}")
    if not re.fullmatch(r"[A-Za-z0-9_.-]+", args.run_id):
        raise ValueError(f"run_id contains not allowed character: {args.run_id!r}")
    gwas_manifest = source_manifest_input_path(
        args.source_manifest, "gwas_manifest"
    )
    log(f"GWAS main locus manifest: {gwas_manifest}")

    candidate_fields, candidates = read_tsv(args.candidates)
    link_fields, links = read_tsv(args.pcgene_links)
    pair_fields, pairs = read_tsv(args.candidate_pairs)
    five_only_fields, five_only_rows = read_tsv(
        args.five_only_new_gene_map
    )
    five_only_gene_index = build_five_only_gene_index(
        args.five_only_new_gene_map,
        five_only_fields,
        five_only_rows,
    )
    source_metrics = read_metrics(args.source_summary)
    validate_manifest(
        args.source_manifest,
        {
            "all_new_gene_candidates": args.candidates,
            "integrated_newgene_pcGene_links": args.pcgene_links,
            "candidate_newgene_pcGene_mechanism_pairs": args.candidate_pairs,
            "integration_summary": args.source_summary,
        },
    )
    validate_inputs(
        args.candidates,
        candidate_fields,
        candidates,
        args.candidate_pairs,
        pair_fields,
        pairs,
        args.pcgene_links,
        link_fields,
        links,
        five_only_gene_index,
        source_metrics,
    )
    log(
        f"candidate record={len(candidates)}; within candidate aggregation pcGene relationship={len(pairs)}; "
        f"pcGene expand relationship={len(links)}"
    )

    log("[stage 2/5] build three levels of detail and supplement GWAS credible set display context")
    gwas_main_loci, gwas_credible_members = read_gwas_credible_members(
        gwas_manifest
    )
    gwas_credible_summary = summarize_gwas_credible_members(
        gwas_main_loci, gwas_credible_members
    )
    log(
        "GWAS main credible set: "
        f"loci={gwas_credible_summary['gwas_main_locus_count']}; "
        f"credible sets={gwas_credible_summary['gwas_main_credible_set_count']}; "
        f"unique SNPs={gwas_credible_summary['gwas_main_unique_credible_set_snps']}; "
        f"p_lrt<1e-5 SNPs="
        f"{gwas_credible_summary['gwas_main_p_lrt_lt_1e-5_unique_snps']}; "
        f"minimum p_lrt={gwas_credible_summary['gwas_main_min_p_lrt']}"
    )
    candidate_by_id = {row["candidate_id"]: row for row in candidates}
    link_by_pair: Dict[Tuple[str, str, str], Mapping[str, str]] = {}
    for link in links:
        key = (
            link["candidate_id"],
            link["partner_gene_id"],
            link["mechanism"],
        )
        if key in link_by_pair:
            previous = link_by_pair[key]
            coordinate_fields = (
                "partner_chrom",
                "partner_start",
                "partner_end",
                "partner_biotype",
            )
            if any(previous[field] != link[field] for field in coordinate_fields):
                raise ValueError(f"the same pair of pcGene coordinates or biotype conflict: {key}")
        else:
            link_by_pair[key] = link
    block_rows = build_block_rows(candidates, five_only_gene_index)
    pair_rows = build_pair_rows(
        pairs, candidate_by_id, five_only_gene_index, link_by_pair
    )
    detail_rows = build_detail_rows(links, five_only_gene_index)
    focus_block_rows = [
        row for row in block_rows if row["evidence_tier"] == FOCUS_TIER
    ]
    focus_pair_rows = [
        row for row in pair_rows if row["evidence_tier"] == FOCUS_TIER
    ]
    focus_detail_rows = [
        row for row in detail_rows if row["evidence_tier"] == FOCUS_TIER
    ]
    focus_detail_rows = enrich_focus_gwas_layout(
        focus_detail_rows, gwas_credible_members
    )
    focus_candidate_ids = {
        row["candidate_id"] for row in focus_detail_rows
    }


    tier_counts = Counter(row["evidence_tier"] for row in candidates)
    summary_map = {
        "source_run_id": source_metrics["run_id"],
        "overall_status": "PASS",
        "candidate_rows": str(len(block_rows)),
        "unique_newgene_pcgene_pairs": source_metrics["unique_newgene_pcGene_pairs"],
        "focus_candidates_with_pcgene": str(len(focus_candidate_ids)),
        "focus_candidates_without_pcgene": str(len(focus_block_rows) - len(focus_candidate_ids)),
        "unique_new_gene_count": str(
            len({row["new_gene_id"] for row in block_rows})
        ),
        "unique_eqtl_block_count": str(
            len({row["eqtl_block_id"] for row in block_rows})
        ),
        "tier_A_candidate_rows": str(
            tier_counts["A_BOTH_95PCT_CREDIBLE_PLUS_ALLELOME"]
        ),
        "tier_B_candidate_rows": str(
            tier_counts["B_GWAS_NONCREDIBLE_PLUS_ALLELOME"]
        ),
        "tier_C_candidate_rows": str(
            tier_counts["C_BOTH_95PCT_CREDIBLE_NO_ALLELOME"]
        ),
        "tier_D_candidate_rows": str(
            tier_counts["D_GWAS_NONCREDIBLE_NO_ALLELOME"]
        ),
        "all_allelome_expanded_link_rows": source_metrics[
            "integrated_allelome_link_rows"
        ],
        "pcgene_expanded_link_rows": str(len(detail_rows)),
        "candidate_specific_unique_pcgene_pair_rows": str(len(pair_rows)),

        "unique_pair_rows": str(len(pair_rows)),
        "expanded_link_rows": str(len(detail_rows)),
        "lncrna_new_gene_count": str(
            len(
                {
                    row["new_gene_id"]
                    for row in block_rows
                    if row["new_gene_coding_class"] == "noncoding_lncRNA"
                }
            )
        ),
        "coding_mrna_new_gene_count": str(
            len(
                {
                    row["new_gene_id"]
                    for row in block_rows
                    if row["new_gene_coding_class"] == "coding_mRNA"
                }
            )
        ),
        "mixed_new_gene_count": str(
            len(
                {
                    row["new_gene_id"]
                    for row in block_rows
                    if row["new_gene_coding_class"]
                    == "mixed_lncRNA_mRNA"
                }
            )
        ),
        "overlapping_expanded_rows": str(
            sum(
                row["new_gene_pcgene_overlap"] == "yes"
                for row in detail_rows
            )
        ),
        "overlapping_unique_pair_rows": str(
            sum(
                row["new_gene_pcgene_overlap"] == "yes"
                for row in pair_rows
            )
        ),
        "focus_evidence_tier": FOCUS_TIER,
        "focus_candidate_rows": str(len(focus_block_rows)),
        "focus_unique_pair_rows": str(len(focus_pair_rows)),
        "focus_expanded_link_rows": str(len(focus_detail_rows)),
        "focus_all_allelome_expanded_link_rows": str(
            sum(
                to_int(row["allelome_link_count"], "allelome_link_count")
                for row in focus_block_rows
            )
        ),
        "focus_pcgene_expanded_link_rows": str(len(focus_detail_rows)),
        **gwas_credible_summary,
        "focus_candidate_specific_unique_pcgene_pair_rows": str(
            len(focus_pair_rows)
        ),
        "candidate_exact_three_layer_variant_rows": source_metrics[
            "candidate_exact_three_layer_variant_rows"
        ],
        "candidate_same_block_multilayer_rows": source_metrics[
            "candidate_same_block_multilayer_rows"
        ],
        "shared_snp_equals_gwas_lead_candidate_rows": str(
            sum(row["overlap_snp_equals_lead"] == "yes" for row in block_rows)
        ),
        "shared_snp_equals_gwas_top_pip_candidate_rows": str(
            sum(
                row["overlap_snp_equals_top_pip"] == "yes"
                for row in block_rows
            )
        ),
        "shared_snp_equals_allelome_endpoint_candidate_rows": "NA",
        "integration_rule": source_metrics["integration_rule"],
        "evidence_boundary": source_metrics["evidence_boundary"],
        "direction_interpretation": source_metrics[
            "direction_interpretation"
        ],
    }
    summary_rows = [
        {"metric": key, "value": value} for key, value in summary_map.items()
    ]
    focus_summary_map = dict(summary_map)
    focus_summary_map.update(
        {
            "candidate_rows": str(len(focus_block_rows)),
            "unique_pair_rows": str(len(focus_pair_rows)),
            "expanded_link_rows": str(len(focus_detail_rows)),
            "all_allelome_expanded_link_rows": str(
                sum(
                    to_int(
                        row["allelome_link_count"], "allelome_link_count"
                    )
                    for row in focus_block_rows
                )
            ),
            "pcgene_expanded_link_rows": str(len(focus_detail_rows)),
            "candidate_specific_unique_pcgene_pair_rows": str(
                len(focus_pair_rows)
            ),
            "lncrna_new_gene_count": str(
                len(
                    {
                        row["new_gene_id"]
                        for row in focus_block_rows
                        if row["new_gene_coding_class"]
                        == "noncoding_lncRNA"
                    }
                )
            ),
            "overlapping_expanded_rows": str(
                sum(
                    row["new_gene_pcgene_overlap"] == "yes"
                    for row in focus_detail_rows
                )
            ),
        }
    )

    log("[stage 3/5] create timestamp result directory")
    run_dir = args.output_root / args.run_id
    if run_dir.exists():
        raise FileExistsError(f"refuse to overwrite existing result directory: {run_dir}")
    run_dir.mkdir(parents=True, exist_ok=False)

    block_path = run_dir / "dog5_candidate_records.tsv"
    pair_path = run_dir / "dog5_candidate_newgene_pcGene_mechanism_pairs.tsv"
    detail_path = run_dir / "dog5_expanded_newgene_pcGene_evidence.tsv"
    summary_path = run_dir / "dog5_summary_metrics.tsv"
    manifest_path = run_dir / "dog5_summary_manifest.tsv"
    workbook_path = (
        run_dir / "dog5_newgene_allelome_eqtl_gwas_finemap_summary.xlsx"
    )
    html_path = (
        run_dir / "dog5_tierA_both_credible_allelome_results.html"
    )
    svg_path = (
        run_dir / "dog5_tierA_both_credible_allelome_results.svg"
    )

    log("[stage 4/5] write TSV, Excel, HTML and SVG")
    log(
        "Tier A main display: "
        f"candidate records={len(focus_block_rows)}; "
        f"candidate-specific pcGene pairs={len(focus_pair_rows)}; "
        f"pcGene expanded evidence={len(focus_detail_rows)}"
    )
    write_tsv(block_path, BLOCK_COLUMNS, block_rows)
    write_tsv(pair_path, PAIR_COLUMNS, pair_rows)
    write_tsv(detail_path, DETAIL_COLUMNS, detail_rows)
    write_tsv(summary_path, ["metric", "value"], summary_rows)
    write_workbook(
        workbook_path,
        summary_rows,
        block_rows,
        pair_rows,
        detail_rows,
        focus_detail_rows,
    )
    write_html(html_path, focus_detail_rows, focus_summary_map)
    write_svg(svg_path, focus_detail_rows, focus_summary_map)

    manifest_rows = [
        {
            "kind": "input",
            "label": label,
            "path": str(path),
            "purpose": purpose,
        }
        for label, path, purpose in [
            (
                "candidate_records",
                args.candidates,
                f"All {len(candidates)} candidate records from step 8",
            ),
            (
                "pcgene_links",
                args.pcgene_links,
                f"{len(links)} expanded per-dog new gene-pcGene relationships from step 8",
            ),
            (
                "candidate_pcgene_pairs",
                args.candidate_pairs,
                f"{len(pairs)} within-candidate aggregated new gene-pcGene relationships from step 8",
            ),
            ("source_summary", args.source_summary, "Step 8 PASS status and count requirements"),
            ("source_manifest", args.source_manifest, "Step 8 input/output provenance"),
            (
                "gwas_locus_manifest",
                gwas_manifest,
                "Absolute paths and run statuses of per-locus credible-set member tables in GWAS main",
            ),
            (
                "five_only_new_gene_map",
                args.five_only_new_gene_map,
                "Five-dog new_gene_id, FEELnc type, coding class, and gene span",
            ),
        ]
    ]
    for label, path, purpose in [
        ("candidate_records", block_path, "Candidate records and A/B/C/D evidence tiers"),
        ("candidate_pcgene_pairs", pair_path, "Within-candidate aggregated new gene-pcGene relationships"),
        ("expanded_evidence", detail_path, "Per-dog, per-pcGene relationship evidence"),
        ("summary_metrics", summary_path, "Counts and interpretation limits"),
        (
            "workbook",
            workbook_path,
            "focus_results contains only Tier A pcGene relationships; full_results retains Tier A/B pcGene evidence.",
        ),
        (
            "html",
            html_path,
            "Interactive table of per-dog pcGene evidence restricted to Tier A with shared credible-set membership and Allelome support",
        ),
        (
            "svg",
            svg_path,
            "Static plot of per-dog pcGene evidence restricted to Tier A with shared credible-set membership and Allelome support",
        ),
    ]:
        manifest_rows.append(
            {
                "kind": "output",
                "label": label,
                "path": str(path),
                "purpose": purpose,
            }
        )
    write_tsv(
        manifest_path,
        ["kind", "label", "path", "purpose"],
        manifest_rows,
    )

    log("[stage 5/5] completed count check")
    log(
        f"candidate records={len(block_rows)}; "
        f"candidate-specific pcGene pairs={len(pair_rows)}; "
        f"pcGene expanded evidence={len(detail_rows)}"
    )
    log(f"task completed / done: {run_dir}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
